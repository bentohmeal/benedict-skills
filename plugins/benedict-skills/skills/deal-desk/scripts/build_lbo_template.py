"""Build a blank standard PE LBO template (house style v2), or fill it with a test case and verify it.

    python build_lbo_template.py out.xlsx                          # blank template, recalc, lint
    python build_lbo_template.py out.xlsx --fill case.json         # test-filled copy, cross-checked against Python
    options: --years N (default 5) · --units "$m" · --no-verify

Structure: Cover · Inputs · Deal · Model · Checks. Same grid on every sheet: A markers · B labels · C units ·
D/E constants · F-H = Year -2 .. Year 0 (LTM / close) · then Years 1..n.
Debt: revolver (RCF) + Term Loan A + Term Loan B (mandatory amortisation, cash sweep in that order) + senior notes
(bullet). Interest on opening balances, so no circular reference. Scenario toggle (Base / Upside / Downside),
credit stats and covenants, returns attribution, three 5x5 data tables with formula axes (base ± step).
Blank-safe: every ratio is IF(denominator=0,0,...), so an empty template shows "–", never #DIV/0!.
Test case fields: see ../assets/examples/lbo-template-test.json.
"""
import json
import os
import re
import subprocess
import sys

from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlstyle import (FIRST_PERIOD, MULT, NUM, NUM1, PCT, STYLE_VERSION, TEXT, TOTAL, GRAND, YEAR_FMT, check_cell, col,  # noqa: E402
                     colour_key, comment, output_box, sensitivity,
                     end_marker, header_row, input_cell, label, na_cell, note, put, row_of, section, setup,
                     sheet_header, sources_uses, subsection)

HERE = os.path.dirname(os.path.abspath(__file__))
HIST = 3                          # Year -2, Year -1, Year 0
Y0C = FIRST_PERIOD + HIST - 1     # column index of Year 0 (H)
CASES = ("Base", "Upside", "Downside")
DRIVERS = [("growth", "Revenue growth"), ("gm", "Gross margin"), ("opex", "Operating expenses (% of revenue)"),
           ("da", "D&A (% of revenue)"), ("capex", "Capex (% of revenue)"), ("nwc", "Net working capital (% of revenue)")]
TERM = [("tla", "Term Loan A"), ("tlb", "Term Loan B"), ("notes", "Senior notes")]
FLAG = "0"


def yc(t):
    """Column letter of Year t (t = -2 .. n)."""
    return col(Y0C + t)


def div(num, den):
    """Blank-safe ratio (no leading '=')."""
    return f"IF({den}=0,0,{num}/{den})"


def get(c, *keys, default=None):
    """Read a value from the test case; None (blank input) when building the empty template."""
    v = c
    for k in keys:
        if v is None:
            return default
        try:
            v = v[k]
        except (KeyError, IndexError, TypeError):
            return default
    return v


# ----------------------------------------------------------------------------------------------------------
def build(c, path, n=5, U="$m"):
    """c = test case dict, or None for the blank template. Returns the reference maps used by verify()."""
    LASTC = Y0C + n
    last = yc(n)
    wb = Workbook()
    from xlstyle import workbook_defaults
    workbook_defaults(wb)
    cover = wb.active
    cover.title = "Cover"
    inp, deal, m, chk = (wb.create_sheet(s) for s in ("Inputs", "Deal", "Model", "Checks"))
    for ws in (cover, inp, deal, m, chk):
        setup(ws)
    case = get(c, "case_name", default="Project [Name] – LBO")
    fy0 = get(c, "first_year")   # fiscal year of Year 0 (e.g. 2025) -> FY2023 .. FY2030 labels (lessons #30)
    timeline = [("Financial year", [(fy0 or 0) - HIST + 1] + [f"={col(FIRST_PERIOD + i - 1)}5+1" for i in range(1, HIST + n)],
                 YEAR_FMT, False),
                ("Period", ["Actual"] * (HIST - 1) + ["LTM / close"] + ["Forecast"] * n, "@", True)]
    I, D, R = {}, {}, {}   # Inputs refs · Deal refs · Model rows

    # ================================================================ Inputs
    ws = inp
    r = sheet_header(ws, case, f"Inputs · {U} unless stated · type only in the yellow cells", LASTC, timeline)

    def row_inputs(key, lab, unit, fmt, values, cols):
        """values: list aligned with cols (t indices); other period columns are grey N/A."""
        nonlocal r
        label(ws, r, lab, unit)
        for t in range(-HIST + 1, n + 1):
            if t in cols:
                input_cell(ws, f"{yc(t)}{r}", values[cols.index(t)] if values else None, fmt)
            else:
                na_cell(ws, f"{yc(t)}{r}")
        I[key] = r
        r += 1

    def const_input(key, lab, unit, value, fmt, note_txt=None):
        nonlocal r
        label(ws, r, lab, unit)
        input_cell(ws, f"D{r}", value, fmt)
        if note_txt:
            comment(ws, f"E{r}", note_txt)
        I[key] = f"Inputs!$D${r}"
        r += 1

    hist_cols = list(range(-HIST + 1, 1))
    fc_cols = list(range(1, n + 1))
    section(ws, r, "Historical financials (costs entered as positive numbers)", last_col=LASTC)
    r += 1
    for k, lab in [("revenue", "Revenue"), ("cogs", "Cost of sales"), ("opex", "Operating expenses (excl. D&A)"),
                   ("da", "Depreciation & amortisation"), ("capex", "Capex"), ("nwc", "Net working capital (balance)")]:
        row_inputs("h_" + k, lab, U, NUM1, get(c, "historical", k), hist_cols)
    r += 1

    section(ws, r, "Scenario selector", last_col=LASTC)
    r += 1
    const_input("case", "Case in use (1 = Base, 2 = Upside, 3 = Downside)", "#", get(c, "case", default=1), FLAG)
    put(ws, f"E{r - 1}", f'=CHOOSE(D{r - 1},"{CASES[0]}","{CASES[1]}","{CASES[2]}")', TEXT, bold=True)
    r += 1

    section(ws, r, "Operating drivers by case (forecast years)", last_col=LASTC)
    r += 1
    for k, lab in DRIVERS:
        subsection(ws, r, lab, LASTC)
        r += 1
        for i, cs in enumerate(CASES):
            row_inputs(f"d_{k}_{i}", f"{cs} case", "%", PCT, get(c, "drivers", cs.lower(), k), fc_cols)
    r += 1

    section(ws, r, "Tax, cash and base rate", last_col=LASTC)
    r += 1
    const_input("tax", "Tax rate", "%", get(c, "tax_rate"), PCT)
    const_input("min_cash", "Minimum cash balance", U, get(c, "min_cash"), NUM1, "also the cash put on the balance sheet at close")
    const_input("cash_rate", "Interest earned on cash", "%", get(c, "cash_rate"), PCT, "on the opening cash balance")
    row_inputs("base_rate", "Base rate (e.g. SOFR)", "%", PCT, get(c, "base_rate"), fc_cols)
    r += 1

    section(ws, r, "Debt terms", last_col=LASTC)
    r += 1
    subsection(ws, r, "Revolving credit facility (RCF)", LASTC)
    r += 1
    rc = lambda k: get(c, "rcf", k)  # noqa: E731
    const_input("rcf_commit", "Commitment", U, rc("commitment"), NUM1)
    const_input("rcf_drawn", "Drawn at close", U, rc("drawn"), NUM1)
    const_input("rcf_float", "Floating rate (1) or fixed (0)", "(1,0)", get(c, "rcf", "float", default=1), FLAG)
    const_input("rcf_margin", "Margin over base rate (or fixed coupon)", "%", rc("margin"), PCT)
    const_input("rcf_floor", "Base rate floor", "%", rc("floor"), PCT)
    const_input("rcf_undrawn", "Undrawn commitment fee", "%", rc("undrawn_fee"), PCT, "on commitment less opening drawn balance")
    const_input("rcf_fee", "Upfront fee (% of commitment)", "%", rc("fee_pct"), PCT)
    const_input("rcf_years", "Fee amortisation period", "years", rc("fee_years"), FLAG)
    for key, name in TERM:
        subsection(ws, r, name, LASTC)
        r += 1
        t_ = lambda k, d=None: get(c, key, k, default=d)  # noqa: E731
        const_input(f"{key}_x", "Amount (x Year 0 EBITDA)", "x", t_("x_ebitda"), MULT)
        const_input(f"{key}_float", "Floating rate (1) or fixed (0)", "(1,0)", t_("float", 0 if key == "notes" else 1), FLAG)
        const_input(f"{key}_margin", "Margin over base rate (or fixed coupon)", "%", t_("margin"), PCT)
        const_input(f"{key}_floor", "Base rate floor", "%", t_("floor"), PCT)
        const_input(f"{key}_fee", "Upfront fee / OID (% of amount)", "%", t_("fee_pct"), PCT)
        const_input(f"{key}_years", "Fee amortisation period", "years", t_("fee_years"), FLAG)
        if key == "notes":
            const_input("notes_mat", "Bullet maturity (forecast year number)", "year", t_("maturity"), FLAG,
                        "blank or beyond the forecast = still outstanding at exit")
        else:
            row_inputs(f"{key}_amort", "Mandatory amortisation (% of original amount)", "%", PCT, t_("amort"), fc_cols)
            const_input(f"{key}_sweep", "Cash sweep (% of excess cash)", "%", t_("sweep", 1), PCT,
                        "excess cash goes to TLA first, then TLB")
    r += 1

    section(ws, r, "Maintenance covenants", last_col=LASTC)
    r += 1
    row_inputs("cov_lev", "Maximum net debt / EBITDA", "x", MULT, get(c, "covenants", "max_net_leverage"), fc_cols)
    row_inputs("cov_icr", "Minimum EBITDA / cash interest", "x", MULT, get(c, "covenants", "min_interest_cover"), fc_cols)
    r += 1

    section(ws, r, "Returns table steps (axis spacing, Deal sheet)", last_col=LASTC)
    r += 1
    st = lambda k, d: get(c, "steps", k, default=d)  # noqa: E731
    const_input("step_entry", "Entry multiple step", "x", st("entry", 0.5), MULT)
    const_input("step_exit", "Exit multiple step", "x", st("exit", 0.5), MULT)
    const_input("step_growth", "Revenue growth step", "%", st("growth", 0.01), PCT)
    const_input("step_margin", "EBITDA margin step", "%", st("margin", 0.01), PCT)
    end_marker(ws, r + 1, LASTC)

    def iv(key, t):
        """Inputs cell of a by-year row for Year t (row-absolute)."""
        return f"Inputs!{yc(t)}${I[key]}"

    # ================================================================ Deal (top: assumptions, debt sizing, S&U)
    ws = deal
    r = sheet_header(ws, case, f"Deal: transaction assumptions, sources & uses, returns, sensitivities · {U}", LASTC)
    note(ws, f"B{r}", "Override cells stay blank: the sensitivity tables flex them (Excel needs them on this sheet).")
    r += 1
    section(ws, r, "Transaction assumptions", last_col=LASTC)
    r += 1

    def drow(key, lab, unit, value, fmt, kind="input", note_txt=None, bold=False):
        nonlocal r
        label(ws, r, lab, unit, bold=bold)
        if kind == "input":
            input_cell(ws, f"D{r}", value, fmt)
        else:
            put(ws, f"D{r}", value, fmt, bold=bold, kind="link" if "@@" in str(value) else None)
        if note_txt:
            comment(ws, f"E{r}", note_txt)
        D[key] = f"$D${r}"
        r += 1

    dg = lambda k: get(c, "deal", k)  # noqa: E731
    drow("ltm", "Year 0 (LTM) EBITDA", U, "=Model!$H$@@ebitda@@", NUM1, "calc")
    drow("entry_x", "Entry EV / LTM EBITDA", "x", dg("entry_multiple"), MULT)
    drow("entry_ovr", "Entry multiple – sensitivity override", "x", None, MULT, note_txt="leave blank")
    drow("entry_used", "Entry multiple used", "x", "=IF(ISNUMBER({entry_ovr}),{entry_ovr},{entry_x})", MULT, "calc")
    drow("ev", "Enterprise value (cash-free / debt-free)", U, "={ltm}*{entry_used}", NUM1, "calc", bold=True)
    drow("nd_old", "Existing net debt refinanced at close", U, dg("existing_net_debt"), NUM1)
    drow("eq_price", "Equity purchase price", U, "={ev}-{nd_old}", NUM1, "calc")
    drow("txn_pct", "Transaction fees (% of EV)", "%", dg("transaction_fee_pct"), PCT, note_txt="advisory, legal, diligence")
    drow("exit_x", "Exit EV / EBITDA", "x", dg("exit_multiple"), MULT)
    drow("exit_ovr", "Exit multiple – sensitivity override", "x", None, MULT, note_txt="leave blank")
    drow("exit_used", "Exit multiple used", "x", "=IF(ISNUMBER({exit_ovr}),{exit_ovr},{exit_x})", MULT, "calc")
    drow("exit_cost", "Exit costs (% of exit EV)", "%", dg("exit_cost_pct"), PCT)
    drow("g_adj", "Revenue growth adjustment – sensitivity", "%", None, PCT, note_txt="leave blank; added to every forecast year")
    drow("m_adj", "EBITDA margin adjustment – sensitivity", "%", None, PCT, note_txt="leave blank; cuts opex % of revenue")
    # resolve {key} placeholders inside this sheet's formulas
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=") and "{" in cell.value:
                cell.value = cell.value.format(**D)
    r += 1

    section(ws, r, "Debt sizing and financing fees at close", last_col=LASTC)
    r += 1
    header_row(ws, r, [U, "x EBITDA", "Fee %", "Fee"], 4)
    r += 1
    L = D["ltm"]

    def debt_row(key, lab, amount, fee_pct=None, bold=False):
        nonlocal r
        label(ws, r, lab, U, bold=bold)
        put(ws, f"D{r}", amount, NUM1, bold=bold)
        put(ws, f"E{r}", "=" + div(f"D{r}", L), MULT, bold=bold)
        if fee_pct:
            put(ws, f"F{r}", fee_pct, PCT)
            put(ws, f"G{r}", f"=D{r}*F{r}", NUM1)
        D[key] = f"$D${r}"
        D["fee_" + key] = f"$G${r}"
        r += 1

    debt_row("rcf_commit", "RCF commitment", f"={I['rcf_commit']}", f"={I['rcf_fee']}")
    first_debt = r
    debt_row("rcf", "RCF drawn at close", f"={I['rcf_drawn']}")
    for key, name in TERM:
        debt_row(key, name, f"={I[key + '_x']}*{L}", f"={I[key + '_fee']}")
    label(ws, r, "Total debt at close", U, bold=True)
    put(ws, f"D{r}", f"=SUM(D{first_debt}:D{r - 1})", NUM1, bold=True).border = TOTAL
    put(ws, f"E{r}", "=" + div(f"D{r}", L), MULT, bold=True).border = TOTAL
    put(ws, f"G{r}", f"=SUM(G{first_debt - 1}:G{r - 1})", NUM1, bold=True).border = TOTAL
    D["debt"], D["fees"] = f"$D${r}", f"$G${r}"
    comment(ws, f"H{r}", "fees are capitalised and amortised through the P&L (non-cash)")
    r += 2

    section(ws, r, "Sources and uses at close", last_col=LASTC)
    r += 2
    su, tot_s, tot_u, r = sources_uses(
        ws, r, U,
        sources=[("Revolving credit facility (drawn)", f"={D['rcf']}")] + [(nm, f"={D[k]}") for k, nm in TERM]
        + [("Sponsor equity", "PLUG")],
        uses=[("Equity purchase price", f"={D['eq_price']}"), ("Refinance existing net debt", f"={D['nd_old']}"),
              ("Transaction fees", f"={D['ev']}*{D['txn_pct']}"), ("Financing fees", f"={D['fees']}"),
              ("Cash to balance sheet", f"={I['min_cash']}")],
        ebitda_ref=L, guard=True)
    eq = su["Sponsor equity"]
    comment(ws, f"G{eq[1:]}", "the plug: total uses less debt")
    D["equity"] = f"${eq[0]}${eq[1:]}"
    D["sources"], D["uses"] = tot_s, tot_u
    D["txn_fees"] = "$" + su["Transaction fees"][0] + "$" + su["Transaction fees"][1:]
    cb = su["Cash to balance sheet"]
    D["cash_close"] = f"${cb[0]}${cb[1:]}"

    section(ws, r, "Pro forma capitalisation at close", last_col=LASTC)
    r += 1
    header_row(ws, r, [U, "x EBITDA", "Cumul. x", "% of total"], 4)
    r += 1
    cap_first = r
    cap = {}

    def cap_row(key, lab, f, bold=False, cumul=True, grand=False):
        nonlocal r
        label(ws, r, lab, U, bold=bold)
        cells = [put(ws, f"D{r}", f, NUM1, bold=bold), put(ws, f"E{r}", "=" + div(f"D{r}", L), MULT, bold=bold)]
        if cumul:
            cells.append(put(ws, f"F{r}", f"=E{r}" if key == "rcf" else f"=F{r - 1}+E{r}" if key in ("tla", "tlb", "notes")
                             else f"=E{r}", MULT, bold=bold))
        cap[key] = r
        r += 1
        for cl in cells:
            if grand:
                cl.border = GRAND
            elif bold:
                cl.border = TOTAL

    cap_row("rcf", "RCF (drawn)", f"={D['rcf']}")
    cap_row("tla", "Term Loan A", f"={D['tla']}")
    cap_row("tlb", "Term Loan B", f"={D['tlb']}")
    cap_row("senior", "Total senior secured debt", f"=SUM(D{cap['rcf']}:D{cap['tlb']})", bold=True)
    # the notes' cumulative multiple builds on the senior total
    label(ws, r, "Senior notes", U)
    put(ws, f"D{r}", f"={D['notes']}", NUM1)
    put(ws, f"E{r}", "=" + div(f"D{r}", L), MULT)
    put(ws, f"F{r}", f"=F{cap['senior']}+E{r}", MULT)
    cap["notes"] = r
    r += 1
    cap_row("debt", "Total debt", f"=D{cap['senior']}+D{cap['notes']}", bold=True)
    cap_row("cash", "Less: cash at close", f"=-{D['cash_close']}", cumul=False)
    cap_row("nd", "Net debt", f"=D{cap['debt']}+D{cap['cash']}", bold=True, cumul=False)
    cap_row("eq", "Sponsor equity", f"={D['equity']}", cumul=False)
    cap_row("tot", "Total capitalisation (total debt + equity)", f"=D{cap['debt']}+D{cap['eq']}", bold=True, cumul=False,
            grand=True)
    for key in ("rcf", "tla", "tlb", "senior", "notes", "debt", "eq", "tot"):
        put(ws, f"G{cap[key]}", "=" + div(f"D{cap[key]}", f"$D${cap['tot']}"), PCT, bold=key in ("senior", "debt", "tot"))
    r += 1
    D["_row"] = r   # next free row on Deal (returns, attribution and sensitivities go here)

    # ================================================================ Model
    ws = m
    r = sheet_header(ws, case, f"Model: operating model, debt schedule, cash flow, credit stats, returns · {U}", LASTC,
                     timeline, const_labels=("Value", ""))

    def sec(text):
        nonlocal r
        section(ws, r, text, last_col=LASTC)
        r += 1

    def line(key, lab, unit, fn, fmt=NUM1, hist=None, bold=False, total=False, italic=False, grand=False, start=1):
        """fn(t, L) -> formula for forecast years start..n; hist(t, L) -> formula for Year -2..0 (or None)."""
        nonlocal r
        label(ws, r, lab, unit, bold=bold, italic=italic, indent=1 if italic else 0)
        R[key] = r
        vals = []
        for t in range(-HIST + 1, n + 1):
            if t < start:
                vals.append(hist(t, yc(t)) if hist else None)
            else:
                vals.append(fn(t, yc(t)))
        row_of(ws, r, FIRST_PERIOD, vals, fmt, bold=bold, total=total, italic=italic, grand=grand)
        r += 1

    pc = lambda t: yc(t - 1)  # noqa: E731
    CASE_ = I["case"]

    def live(k, t):
        return f"CHOOSE({CASE_},{iv(f'd_{k}_0', t)},{iv(f'd_{k}_1', t)},{iv(f'd_{k}_2', t)})"

    sec("Operating drivers – live case (chosen on Inputs)")
    line("g", "Revenue growth", "%", lambda t, X: f"={live('growth', t)}+Deal!{D['g_adj']}", PCT)
    line("gm", "Gross margin", "%", lambda t, X: f"={live('gm', t)}", PCT)
    line("opex_pct", "Operating expenses (% of revenue)", "%", lambda t, X: f"={live('opex', t)}-Deal!{D['m_adj']}", PCT)
    line("da_pct", "D&A (% of revenue)", "%", lambda t, X: f"={live('da', t)}", PCT)
    line("capex_pct", "Capex (% of revenue)", "%", lambda t, X: f"={live('capex', t)}", PCT)
    line("nwc_pct", "Net working capital (% of revenue)", "%", lambda t, X: f"={live('nwc', t)}", PCT)
    r += 1

    hi = lambda key, sign="": (lambda t, X: f"={sign}Inputs!{X}${I[key]}")  # noqa: E731
    sec("Income statement")
    line("rev", "Revenue", U, lambda t, X: f"={pc(t)}{R['rev']}*(1+{X}{R['g']})", hist=hi("h_revenue"), bold=True)
    line("rev_g", "Growth", "%", lambda t, X: f"=IF({pc(t)}{R['rev']}=0,0,{X}{R['rev']}/{pc(t)}{R['rev']}-1)", PCT,
         italic=True, start=-HIST + 2)
    line("cogs", "Cost of sales", U, lambda t, X: f"=-{X}{R['rev']}*(1-{X}{R['gm']})", hist=hi("h_cogs", "-"))
    line("gp", "Gross profit", U, lambda t, X: f"={X}{R['rev']}+{X}{R['cogs']}", bold=True, total=True, start=-HIST + 1)
    line("gp_m", "Gross margin", "%", lambda t, X: "=" + div(f"{X}{R['gp']}", f"{X}{R['rev']}"), PCT, italic=True,
         start=-HIST + 1)
    line("opex", "Operating expenses", U, lambda t, X: f"=-{X}{R['rev']}*{X}{R['opex_pct']}", hist=hi("h_opex", "-"))
    line("ebitda", "EBITDA", U, lambda t, X: f"={X}{R['gp']}+{X}{R['opex']}", bold=True, total=True, start=-HIST + 1)
    line("ebitda_m", "EBITDA margin", "%", lambda t, X: "=" + div(f"{X}{R['ebitda']}", f"{X}{R['rev']}"), PCT,
         italic=True, start=-HIST + 1)
    line("da", "Depreciation & amortisation", U, lambda t, X: f"=-{X}{R['rev']}*{X}{R['da_pct']}", hist=hi("h_da", "-"))
    line("ebit", "EBIT", U, lambda t, X: f"={X}{R['ebitda']}+{X}{R['da']}", bold=True, total=True, start=-HIST + 1)
    # interest lines point forward to the debt schedule (@@row@@ tokens are resolved at the end)
    line("pl_rcf", "Interest – RCF", U, lambda t, X: f"=-{X}@@rcf_int@@")
    line("pl_cfee", "Commitment fee – RCF (undrawn)", U, lambda t, X: f"=-{X}@@rcf_cfee@@")
    for key, name in TERM:
        line(f"pl_{key}", f"Interest – {name}", U, lambda t, X, k=key: f"=-{X}@@{k}_int@@")
    line("pl_fees", "Amortisation of financing fees (non-cash)", U, lambda t, X: f"=-{X}@@fa_tot@@")
    line("pl_cash", "Interest income on cash", U, lambda t, X: f"={X}@@cash_int@@")
    line("net_int", "Net interest expense", U, lambda t, X: f"=SUM({X}{R['pl_rcf']}:{X}{R['pl_cash']})", bold=True,
         total=True)
    line("pbt", "Profit before tax", U, lambda t, X: f"={X}{R['ebit']}+{X}{R['net_int']}", bold=True)
    line("tax", "Income tax", U, lambda t, X: f"=-MAX(0,{X}{R['pbt']})*{I['tax']}")
    comment(ws, f"{col(LASTC + 1)}{R['tax']}", "no tax credit on losses (conservative)")
    line("ni", "Net income", U, lambda t, X: f"={X}{R['pbt']}+{X}{R['tax']}", bold=True, grand=True)
    r += 1

    sec("Capex and net working capital")
    line("capex", "Capex", U, lambda t, X: f"=-{X}{R['rev']}*{X}{R['capex_pct']}", hist=hi("h_capex", "-"))
    line("nwc", "Net working capital (balance)", U, lambda t, X: f"={X}{R['rev']}*{X}{R['nwc_pct']}", hist=hi("h_nwc"))
    line("dnwc", "(Increase) / decrease in net working capital", U,
         lambda t, X: f"=-({X}{R['nwc']}-{pc(t)}{R['nwc']})", start=-HIST + 2)
    r += 1

    sec("Financing fees – amortisation (non-cash)")
    for key, name, fee in [("rcf", "RCF", "fee_rcf_commit")] + [(k, nm, "fee_" + k) for k, nm in TERM]:
        yrs = I[f"{key}_years"]
        line(f"fa_{key}", name, U, lambda t, X, y=yrs, f=fee: f"=IF({y}=0,0,IF({X}$5-${col(Y0C)}$5<={y},Deal!{D[f]}/{y},0))")
    line("fa_tot", "Total amortisation of financing fees", U, lambda t, X: f"=SUM({X}{R['fa_rcf']}:{X}{R['fa_notes']})",
         bold=True, total=True)
    r += 1

    sec("Cash flow, revolver and cash sweep")
    line("cf_ni", "Net income", U, lambda t, X: f"={X}{R['ni']}")
    line("cf_da", "Add: depreciation & amortisation", U, lambda t, X: f"=-{X}{R['da']}")
    line("cf_fees", "Add: amortisation of financing fees", U, lambda t, X: f"=-{X}{R['pl_fees']}")
    line("cf_capex", "Less: capex", U, lambda t, X: f"={X}{R['capex']}")
    line("cf_nwc", "Less: increase in net working capital", U, lambda t, X: f"={X}{R['dnwc']}")
    line("fcf", "Free cash flow before debt repayment", U, lambda t, X: f"=SUM({X}{R['cf_ni']}:{X}{R['cf_nwc']})",
         bold=True, grand=True)
    line("cf_mand", "Less: mandatory debt repayment", U,
         lambda t, X: f"={X}@@tla_mand@@+{X}@@tlb_mand@@+{X}@@notes_mand@@")
    line("cf_open", "Add: opening cash", U, lambda t, X: f"={pc(t)}@@cash@@")
    line("cf_min", "Less: minimum cash balance", U, lambda t, X: f"=-{I['min_cash']}")
    line("avail", "Cash available after mandatory repayment", U,
         lambda t, X: f"=SUM({X}{R['fcf']}:{X}{R['cf_min']})", bold=True, total=True)
    line("rcf_move", "RCF drawdown / (repayment)", U,
         lambda t, X: f"=MIN(MAX(0,-{X}{R['avail']}),{I['rcf_commit']}-{X}@@rcf_open@@)"
                      f"-MIN(MAX(0,{X}{R['avail']}),{X}@@rcf_open@@)")
    comment(ws, f"{col(LASTC + 1)}{R['rcf_move']}", "draws to cover a shortfall (up to the commitment); repaid first")
    line("sweep", "Cash available for the sweep", U, lambda t, X: f"=MAX(0,{X}{R['avail']}+{X}{R['rcf_move']})")
    line("tla_sw", "Optional repayment – Term Loan A", U,
         lambda t, X: f"=-MIN({X}@@tla_open@@+{X}@@tla_mand@@,{X}{R['sweep']}*{I['tla_sweep']})")
    line("tlb_sw", "Optional repayment – Term Loan B", U,
         lambda t, X: f"=-MIN({X}@@tlb_open@@+{X}@@tlb_mand@@,({X}{R['sweep']}+{X}{R['tla_sw']})*{I['tlb_sweep']})")
    line("cash", "Closing cash", U,
         lambda t, X: f"={X}{R['fcf']}+{X}{R['cf_mand']}+{X}{R['cf_open']}+{X}{R['rcf_move']}+{X}{R['tla_sw']}"
                      f"+{X}{R['tlb_sw']}", hist=lambda t, X: f"=Deal!{D['cash_close']}" if t == 0 else None,
         bold=True, grand=True)
    line("cash_int", "Interest income on opening cash", U, lambda t, X: f"={X}{R['cf_open']}*{I['cash_rate']}",
         italic=True)
    r += 1

    sec("Debt schedule")

    def rate_line(key):
        line(f"{key}_rate", "Interest rate", "%",
             lambda t, X: f"=IF({I[key + '_float']}=1,MAX({iv('base_rate', t)},{I[key + '_floor']})+{I[key + '_margin']},"
                          f"{I[key + '_margin']})", PCT, italic=True)
        line(f"{key}_int", "Interest expense", U, lambda t, X: f"={X}{R[key + '_open']}*{X}{R[key + '_rate']}")

    subsection(ws, r, "Revolving credit facility (RCF)", LASTC)
    r += 1
    line("rcf_open", "Opening balance", U, lambda t, X: f"={pc(t)}{r + 2}")
    line("rcf_mv", "Drawdown / (repayment)", U, lambda t, X: f"={X}{R['rcf_move']}")
    line("rcf_close", "Closing balance", U, lambda t, X: f"={X}{R['rcf_open']}+{X}{R['rcf_mv']}",
         hist=lambda t, X: f"=Deal!{D['rcf']}" if t == 0 else None, bold=True, total=True)
    rate_line("rcf")
    line("rcf_cfee", "Commitment fee (undrawn)", U,
         lambda t, X: f"={I['rcf_undrawn']}*MAX(0,{I['rcf_commit']}-{X}{R['rcf_open']})")
    for key, name in TERM:
        subsection(ws, r, name, LASTC)
        r += 1
        line(f"{key}_open", "Opening balance", U, lambda t, X: f"={pc(t)}{r + (3 if key != 'notes' else 2)}")
        if key == "notes":
            line("notes_mand", "Repayment at maturity (bullet)", U,
                 lambda t, X: f"=-IF({X}$5-${col(Y0C)}$5={I['notes_mat']},{X}{R['notes_open']},0)")
            line("notes_close", "Closing balance", U, lambda t, X: f"={X}{R['notes_open']}+{X}{R['notes_mand']}",
                 hist=lambda t, X: f"=Deal!{D['notes']}" if t == 0 else None, bold=True, total=True)
        else:
            line(f"{key}_mand", "Mandatory amortisation", U,
                 lambda t, X, k=key: f"=-MIN({X}{R[k + '_open']},Deal!{D[k]}*{iv(k + '_amort', t)})")
            line(f"{key}_opt", "Optional repayment (cash sweep)", U, lambda t, X, k=key: f"={X}{R[k + '_sw']}")
            line(f"{key}_close", "Closing balance", U,
                 lambda t, X, k=key: f"=SUM({X}{R[k + '_open']}:{X}{R[k + '_opt']})",
                 hist=lambda t, X, k=key: f"=Deal!{D[k]}" if t == 0 else None, bold=True, total=True)
        rate_line(key)
    r += 1
    closings = ["rcf_close", "tla_close", "tlb_close", "notes_close"]
    y0 = lambda f: (lambda t, X: f(t, X) if t == 0 else None)  # noqa: E731
    tot = lambda keys: (lambda t, X: "=" + "+".join(f"{X}{R[k]}" for k in keys))  # noqa: E731
    line("debt", "Total debt", U, tot(closings), hist=y0(tot(closings)), bold=True, total=True)
    line("senior", "Senior secured debt (RCF + TLA + TLB)", U, tot(closings[:3]), hist=y0(tot(closings[:3])))
    nd = lambda t, X: f"={X}{R['debt']}-{X}{R['cash']}"  # noqa: E731
    line("nd", "Net debt", U, nd, hist=y0(nd), bold=True)
    r += 1

    sec("Credit statistics and covenants")
    lev = lambda k: (lambda t, X: "=" + div(f"{X}{R[k]}", f"{X}{R['ebitda']}"))  # noqa: E731
    line("cs_total", "Total debt / EBITDA", "x", lev("debt"), MULT, hist=y0(lev("debt")))
    line("cs_senior", "Senior secured debt / EBITDA", "x", lev("senior"), MULT, hist=y0(lev("senior")))
    line("cs_net", "Net debt / EBITDA", "x", lev("nd"), MULT, hist=y0(lev("nd")), bold=True)
    line("cs_int", "Cash interest expense", U, tot(("rcf_int", "rcf_cfee", "tla_int", "tlb_int", "notes_int")))
    line("cs_icr", "EBITDA / cash interest", "x", lambda t, X: "=" + div(f"{X}{R['ebitda']}", f"{X}{R['cs_int']}"),
         MULT, bold=True)
    line("cs_icr2", "(EBITDA – capex) / cash interest", "x",
         lambda t, X: "=" + div(f"({X}{R['ebitda']}+{X}{R['capex']})", f"{X}{R['cs_int']}"), MULT)
    line("cs_fcf", "FCF before debt repayment / total debt", "%",
         lambda t, X: "=" + div(f"{X}{R['fcf']}", f"{X}{R['debt']}"), PCT)
    line("cs_pay", "Cumulative debt paydown (% of debt at close)", "%",
         lambda t, X: f"=IF(${yc(0)}{R['debt']}=0,0,1-{X}{R['debt']}/${yc(0)}{R['debt']})", PCT)
    subsection(ws, r, "Maintenance covenants", LASTC)
    r += 1
    line("cov_lev", "Maximum net debt / EBITDA", "x", lambda t, X: f"={iv('cov_lev', t)}", MULT)
    line("cov_lev_h", "Headroom", "x",
         lambda t, X: f"=IF({X}{R['cov_lev']}=0,0,{X}{R['cov_lev']}-{X}{R['cs_net']})", MULT, italic=True)
    line("cov_lev_b", "Breach (1 = yes)", "flag",
         lambda t, X: f"=IF(AND({X}{R['cov_lev']}>0,{X}{R['cs_net']}>{X}{R['cov_lev']}),1,0)", FLAG, italic=True)
    line("cov_icr", "Minimum EBITDA / cash interest", "x", lambda t, X: f"={iv('cov_icr', t)}", MULT)
    line("cov_icr_h", "Headroom", "x",
         lambda t, X: f"=IF({X}{R['cov_icr']}=0,0,{X}{R['cs_icr']}-{X}{R['cov_icr']})", MULT, italic=True)
    line("cov_icr_b", "Breach (1 = yes)", "flag",
         lambda t, X: f"=IF(AND({X}{R['cov_icr']}>0,{X}{R['cs_icr']}<{X}{R['cov_icr']}),1,0)", FLAG, italic=True)
    r += 1

    sec(f"Returns – exit at the end of Year {n}")
    M_ = f"${last}$"
    rr = {}
    for k, lab, f, fmt in [
            ("x_ebitda", f"Year {n} EBITDA", f"={last}{R['ebitda']}", NUM1),
            ("x_mult", "Exit EV / EBITDA", f"=Deal!{D['exit_used']}", MULT),
            ("x_ev", "Exit enterprise value", None, NUM1),
            ("x_cost", "Less: exit costs", None, NUM1),
            ("x_debt", "Less: total debt at exit", f"=-{last}{R['debt']}", NUM1),
            ("x_cash", "Add: cash at exit", f"={last}{R['cash']}", NUM1),
            ("x_eq", "Equity value at exit", None, NUM1)]:
        label(ws, r, lab, "x" if fmt == MULT else U, bold=(k == "x_eq"))
        rr[k] = r
        if f:
            put(ws, f"{last}{r}", f, fmt)
        r += 1
    put(ws, f"{last}{rr['x_ev']}", f"={last}{rr['x_ebitda']}*{last}{rr['x_mult']}", NUM1)
    put(ws, f"{last}{rr['x_cost']}", f"=-{last}{rr['x_ev']}*Deal!{D['exit_cost']}", NUM1)
    put(ws, f"{last}{rr['x_eq']}", f"=SUM({last}{rr['x_ev']}:{last}{rr['x_cash']})", NUM1, bold=True).border = GRAND
    R.update(rr)
    r += 1
    line("exit_flag", "Exit year (1 = yes)", "flag", lambda t, X: f"=IF({X}$5=${last}$5,1,0)", FLAG, italic=True)
    line("eq_cf", "Sponsor cash flows (equity in, exit proceeds)", U,
         lambda t, X: f"={X}{R['exit_flag']}*{M_}{R['x_eq']}",
         hist=lambda t, X: f"=-Deal!{D['equity']}" if t == 0 else None, bold=True)
    H0 = yc(0)
    R["irr"], R["mom"] = r + 1, r + 2
    label(ws, r + 1, "Equity IRR", "%", bold=True)
    output_box(ws, f"D{r + 1}", f"=IF({H0}{R['eq_cf']}=0,0,IRR({H0}{R['eq_cf']}:{last}{R['eq_cf']}))", PCT)
    label(ws, r + 2, "Equity multiple (MoM)", "x", bold=True)
    output_box(ws, f"D{r + 2}",
               f"=IF({H0}{R['eq_cf']}=0,0,SUM({yc(1)}{R['eq_cf']}:{last}{R['eq_cf']})/-{H0}{R['eq_cf']})", MULT)
    r += 4
    R["_row"] = r
    end_marker(ws, r, LASTC)

    # ================================================================ Deal (bottom: returns, attribution, tables)
    ws = deal
    r = D.pop("_row")
    section(ws, r, "Returns summary", last_col=LASTC)
    r += 1
    for k, lab, u, f, fmt, box in [
            ("irr", "Equity IRR", "%", f"=Model!$D${R['irr']}", PCT, True),
            ("mom", "Equity multiple (MoM)", "x", f"=Model!$D${R['mom']}", MULT, True),
            ("eq_in", "Sponsor equity at close", U, f"={D['equity']}", NUM1, False),
            ("eq_out", f"Equity value at exit (end of Year {n})", U, f"=Model!{M_}{R['x_eq']}", NUM1, False)]:
        label(ws, r, lab, u, bold=box)
        (output_box if box else put)(ws, f"D{r}", f, fmt)
        D[k] = f"$D${r}"
        r += 1
    r += 1
    section(ws, r, "Returns attribution (value creation bridge)", last_col=LASTC)
    r += 1
    header_row(ws, r, [U, "% of total"], 4)
    r += 1
    a0 = r
    for k, lab, f in [
            ("a_growth", "EBITDA growth (valued at the entry multiple)",
             f"=(Model!{M_}{R['ebitda']}-{D['ltm']})*{D['entry_used']}"),
            ("a_mult", "Multiple expansion / (contraction)",
             f"=({D['exit_used']}-{D['entry_used']})*Model!{M_}{R['ebitda']}"),
            ("a_delev", "Net debt paydown and cash build", f"=Model!${yc(0)}${R['nd']}-Model!{M_}{R['nd']}"),
            ("a_fees", "Fees and costs (transaction, financing, exit)",
             f"=-{D['txn_fees']}-{D['fees']}+Model!{M_}{R['x_cost']}")]:
        label(ws, r, lab, U)
        put(ws, f"D{r}", f, NUM1)
        D[k] = f"$D${r}"
        r += 1
    label(ws, r, "Value created (exit equity less sponsor equity)", U, bold=True)
    put(ws, f"D{r}", f"=SUM(D{a0}:D{r - 1})", NUM1, bold=True).border = GRAND
    D["a_tot"] = f"$D${r}"
    for rw in range(a0, r + 1):
        put(ws, f"E{rw}", "=" + div(f"D{rw}", D["a_tot"]), PCT, bold=(rw == r))
    ws[f"E{r}"].border = GRAND
    r += 1
    label(ws, r, "Check: bridge less (exit equity – sponsor equity)", U, italic=True)
    put(ws, f"D{r}", f"={D['a_tot']}-({D['eq_out']}-{D['eq_in']})", NUM1, italic=True)
    D["a_chk"] = f"$D${r}"
    r += 2

    T = {"tables": [], "centres": []}
    for key, text, top_title, side_title, top_base, side_base, top_step, side_step, top_in, side_in, fmt_axis, out, body in [
            ("t1", "Equity IRR sensitivity – entry multiple vs exit multiple", "Entry multiple (EV / LTM EBITDA)",
             "Exit multiple", D["entry_x"], D["exit_x"], I["step_entry"], I["step_exit"], D["entry_ovr"], D["exit_ovr"],
             MULT, D["irr"], PCT),
            ("t2", "Equity multiple (MoM) sensitivity – entry multiple vs exit multiple",
             "Entry multiple (EV / LTM EBITDA)", "Exit multiple", D["entry_x"], D["exit_x"], I["step_entry"],
             I["step_exit"], D["entry_ovr"], D["exit_ovr"], MULT, D["mom"], MULT),
            ("t3", "Equity IRR sensitivity – revenue growth vs EBITDA margin (every forecast year)",
             "Revenue growth adjustment", "EBITDA margin adjustment", None, None, I["step_growth"], I["step_margin"],
             D["g_adj"], D["m_adj"], PCT, D["irr"], PCT)]:
        section(ws, r, text, last_col=LASTC)
        top, left = r + 2, 4
        c0 = left + 1
        tc = lambda j: f"{col(c0 + j)}{top + 1}"  # noqa: E731
        sc = lambda i: f"{col(left)}{top + 2 + i}"  # noqa: E731

        def build_axis(cell, base, step):
            """Centre = base (a link, or 0 for an adjustment); neighbours = centre ± step, built outwards."""
            if base is None:
                centre, g = 0, (lambda ref, sign: f"={ref}{sign}{step}")
            else:
                centre, g = f"={base}", (lambda ref, sign: f"=IF({ref}=0,0,{ref}{sign}{step})")
            return [g(cell(1), "-"), g(cell(2), "-"), centre, g(cell(2), "+"), g(cell(3), "+")]

        rng, base_ref = sensitivity(ws, top, left, f"={out}", build_axis(sc, side_base, side_step),
                                    build_axis(tc, top_base, top_step), side_title, top_title,
                                    fmt_axis, fmt_axis, body, better="high", centre=True)
        T["tables"].append((f"Deal!{rng}", f"Deal!{top_in}", f"Deal!{side_in}"))
        T["centres"].append((key, f"Deal!{base_ref}", out))
        T[key] = rng
        r = top + 2 + 5 + 1
        note(ws, f"D{r}", "Excel data table (What-If Analysis): axes = base case ± the steps on Inputs; base case boxed.")
        r += 2
    end_marker(ws, r, LASTC)

    # ================================================================ Checks
    ws = chk
    r = sheet_header(ws, case, "Checks: every line must show OK (0); the total feeds the Cover", LASTC,
                     const_labels=("Value", ""))
    section(ws, r, "Integrity checks", last_col=LASTC)
    r += 1
    rng_ = lambda k, t0=1: f"Model!{yc(t0)}{R[k]}:{last}{R[k]}"  # noqa: E731
    checks = [
        ("Sources = uses", f"=Deal!{D['sources']}-Deal!{D['uses']}"),
        ("Debt in the model at close = debt raised", f"=Model!{yc(0)}{R['debt']}-Deal!{D['debt']}"),
        ("No negative debt balance", "=-MIN(0,MIN(" + ",".join(rng_(k, 0) for k in closings) + "))"),
        ("RCF never above its commitment", f"=MAX(0,MAX({rng_('rcf_close', 0)})-{I['rcf_commit']})"),
        ("Cash never below the minimum balance", f"=-MIN(0,MIN({rng_('cash')})-{I['min_cash']})"),
        ("Returns attribution adds up", f"=Deal!{D['a_chk']}"),
        ("NPV of sponsor cash flows at the IRR = 0",
         f"=IF(Model!$D${R['irr']}=0,0,NPV(Model!$D${R['irr']},{rng_('eq_cf')})+Model!{yc(0)}{R['eq_cf']})"),
        ("Case selector is 1, 2 or 3", f"=IF(AND({I['case']}>=1,{I['case']}<=3,INT({I['case']})={I['case']}),0,1)"),
    ] + [(f"Sensitivity table {k[1]}: base case = model output", f"={ref}-Deal!{o}") for k, ref, o in T["centres"]]
    first = r
    for lab, f in checks:
        label(ws, r, lab)
        check_cell(ws, f"D{r}", f)
        r += 1
    label(ws, r, "Total (sum of absolute values)", bold=True)
    check_cell(ws, f"D{r}", f"=SUMPRODUCT(ABS(D{first}:D{r - 1}))")
    ws[f"D{r}"].border = TOTAL
    T["check_total"] = f"Checks!D{r}"
    tot_ref = f"Checks!$D${r}"
    end_marker(ws, r + 2, LASTC)

    # ================================================================ Cover
    ws = cover
    r = sheet_header(ws, case, f"LBO template · {U} · deal-desk house style v{STYLE_VERSION}", LASTC, units_label=False)
    section(ws, r, "Contents", last_col=LASTC)
    r += 1
    for a, b in [("Inputs", "Historicals, Base / Upside / Downside drivers, case selector, debt terms, covenants"),
                 ("Deal", "Transaction assumptions, debt sizing, sources & uses, capitalisation, returns, bridge, "
                          "3 data tables"),
                 ("Model", "Operating model, P&L, cash flow, revolver + cash sweep, debt schedule, credit stats, exit"),
                 ("Checks", "Integrity checks (all must show OK)")]:
        label(ws, r, a, bold=True)
        ws[f"D{r}"] = b
        r += 1
    r += 1
    section(ws, r, "How to fill this template", last_col=LASTC)
    r += 1
    for step in ["1. Inputs: enter Year -2 to Year 0 historicals (costs as positive numbers).",
                 f"2. Inputs: enter Base / Upside / Downside drivers for Years 1–{n}, then pick the case (1, 2 or 3).",
                 "3. Inputs: enter tax, minimum cash, the base-rate curve, each tranche's terms and the covenant levels.",
                 "4. Deal: enter entry / exit multiples, existing net debt, fees and exit costs. Leave override cells blank.",
                 "5. Check this page: Checks = All OK. If the data tables look stale, press F9."]:
        ws[f"B{r}"] = step
        r += 1
    r += 1
    section(ws, r, "Colour code", last_col=LASTC)
    r = colour_key(ws, r + 1) + 1
    section(ws, r, "Key outputs", last_col=LASTC)
    r += 1
    flags = ",".join(rng_(k) for k in ("cov_lev_b", "cov_icr_b"))
    for lab, u, f, fmt in [("Equity IRR", "%", f"=Model!$D${R['irr']}", PCT),
                           ("Equity multiple (MoM)", "x", f"=Model!$D${R['mom']}", MULT),
                           ("Sponsor equity at close", U, f"=Deal!{D['equity']}", NUM1),
                           ("Net debt / EBITDA at close", "x", f"=Model!${yc(0)}${R['cs_net']}", MULT),
                           (f"Equity value at exit (Year {n})", U, f"=Model!{M_}{R['x_eq']}", NUM1),
                           ("Covenant breaches (count of tests failed)", "#", f"=SUM({flags})", '0;-0;"None"'),
                           ("Checks (0 = all OK)", "", f"={tot_ref}", '[>=0.0005]0.000;[<=-0.0005](0.000);"All OK"')]:
        label(ws, r, lab, u)
        output_box(ws, f"D{r}", f, fmt)
        r += 1
    r += 1
    section(ws, r, "Assumptions log (anything the brief or the data does not settle)", last_col=LASTC)
    r += 1
    for i in range(1, 6):
        ws[f"B{r}"] = f"{i}."
        r += 1
    end_marker(ws, r + 1, LASTC)

    # resolve @@row@@ tokens that point at Model rows
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "@@" in cell.value:
                    cell.value = re.sub(r"@@(\w+)@@", lambda mm: str(R[mm.group(1)]), cell.value)
    wb.active = 0
    wb.save(path)
    return I, D, R, T


# ----------------------------------------------------------------------------------------------------------
def irr(flows):
    lo, hi = -0.99, 10.0
    f = lambda x: sum(cf / (1 + x) ** i for i, cf in enumerate(flows))  # noqa: E731
    for _ in range(300):
        mid = (lo + hi) / 2
        if f(lo) * f(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def mirror(c, n=5, entry=None, exit_=None, g_adj=0.0, m_adj=0.0):
    """Independent plain-Python calculation of the same LBO (no spreadsheet), for cross-checking."""
    h, dl = c["historical"], c["deal"]
    dr = c["drivers"][CASES[c.get("case", 1) - 1].lower()]
    ebitda0 = h["revenue"][-1] - h["cogs"][-1] - h["opex"][-1]
    entry = dl["entry_multiple"] if entry is None else entry
    exit_x = dl["exit_multiple"] if exit_ is None else exit_
    ev = ebitda0 * entry
    rcf = c["rcf"]
    tr = {k: dict(c[k]) for k, _ in TERM}
    for k in tr:
        tr[k]["amt"] = tr[k]["x_ebitda"] * ebitda0
    fees = {"rcf": rcf["commitment"] * rcf["fee_pct"], **{k: tr[k]["amt"] * tr[k]["fee_pct"] for k in tr}}
    years = {"rcf": rcf["fee_years"], **{k: tr[k]["fee_years"] for k in tr}}
    debt0 = rcf["drawn"] + sum(tr[k]["amt"] for k in tr)
    txn = ev * dl["transaction_fee_pct"]
    uses = ev + txn + sum(fees.values()) + c["min_cash"]
    equity = uses - debt0
    rate = lambda p, t: max(c["base_rate"][t - 1], p["floor"]) + p["margin"] if p["float"] == 1 else p["margin"]  # noqa
    bal = {"rcf": rcf["drawn"], **{k: tr[k]["amt"] for k in tr}}
    rev, nwc_prev, cash = h["revenue"][-1], h["nwc"][-1], c["min_cash"]
    out = {k: [] for k in ("rcf", "tla", "tlb", "notes", "cash", "ni", "fcf", "ebitda", "net_lev")}
    for t in range(1, n + 1):
        i = t - 1
        rev *= 1 + dr["growth"][i] + g_adj
        ebitda = rev * dr["gm"][i] - rev * (dr["opex"][i] - m_adj)
        da = rev * dr["da"][i]
        ebit = ebitda - da
        cash_int = 0.0
        interest = rate(rcf, t) * bal["rcf"] + rcf["undrawn_fee"] * max(0, rcf["commitment"] - bal["rcf"])
        interest += sum(rate(tr[k], t) * bal[k] for k in tr)
        fee_am = sum(fees[k] / years[k] if years[k] and t <= years[k] else 0 for k in fees)
        cash_int = cash * c["cash_rate"]
        pbt = ebit - interest - fee_am + cash_int
        ni = pbt - max(0, pbt) * c["tax_rate"]
        capex = rev * dr["capex"][i]
        nwc = rev * dr["nwc"][i]
        fcf = ni + da + fee_am - capex - (nwc - nwc_prev)
        mand = {k: min(bal[k], tr[k]["amt"] * tr[k]["amort"][i]) for k in ("tla", "tlb")}
        mand["notes"] = bal["notes"] if t == tr["notes"].get("maturity") else 0.0
        avail = fcf - sum(mand.values()) + cash - c["min_cash"]
        move = min(max(0, -avail), rcf["commitment"] - bal["rcf"]) - min(max(0, avail), bal["rcf"])
        sweep = max(0, avail + move)
        sw_a = min(bal["tla"] - mand["tla"], sweep * tr["tla"]["sweep"])
        sw_b = min(bal["tlb"] - mand["tlb"], (sweep - sw_a) * tr["tlb"]["sweep"])
        cash = cash + fcf - sum(mand.values()) + move - sw_a - sw_b
        bal["rcf"] += move
        bal["tla"] -= mand["tla"] + sw_a
        bal["tlb"] -= mand["tlb"] + sw_b
        bal["notes"] -= mand["notes"]
        nwc_prev = nwc
        for k in bal:
            out[k].append(bal[k])
        out["cash"].append(cash)
        out["ni"].append(ni)
        out["fcf"].append(fcf)
        out["ebitda"].append(ebitda)
        out["net_lev"].append((sum(bal.values()) - cash) / ebitda)
    ev_exit = ebitda * exit_x
    exit_eq = ev_exit * (1 - dl["exit_cost_pct"]) - sum(bal.values()) + cash
    flows = [-equity] + [0.0] * (n - 1) + [exit_eq]
    nd0 = debt0 - c["min_cash"]
    nd_n = sum(bal.values()) - cash
    attr = {"growth": (ebitda - ebitda0) * entry, "multiple": (exit_x - entry) * ebitda,
            "delever": nd0 - nd_n, "fees": -(txn + sum(fees.values()) + ev_exit * dl["exit_cost_pct"])}
    return {**out, "equity": equity, "exit_eq": exit_eq, "irr": irr(flows), "mom": exit_eq / equity, "attr": attr}


def run(py_args, **kw):
    return subprocess.run([sys.executable, *py_args], capture_output=True, text=True, **kw)


def xl_values(path, refs):
    res = run([os.path.join(HERE, "office.py"), "values", path, *refs], check=True)
    return json.loads(res.stdout)


def compare(c, path, R, n, extra=None):
    """Excel vs Python on every year's balances and key lines (plus any extra (name, ref, value) triples)."""
    last = yc(n)
    keys = {"rcf": "rcf_close", "tla": "tla_close", "tlb": "tlb_close", "notes": "notes_close", "cash": "cash",
            "ni": "ni", "fcf": "fcf", "ebitda": "ebitda", "net_lev": "cs_net"}
    refs = {k: f"Model!{yc(1)}{R[v]}:{last}{R[v]}" for k, v in keys.items()}
    extra = extra or []
    xl = xl_values(path, list(refs.values()) + [e[1] for e in extra])
    mr = mirror(c, n)
    ok, worst = True, 0.0
    for k, ref in refs.items():
        got = xl[ref][0]
        for a, b in zip(got, mr[k]):
            err = abs(a - b) / max(1, abs(b))
            worst = max(worst, err)
            ok &= err <= 1e-6
    for name, ref, want in extra:
        a = xl[ref]
        err = abs(a - want) / max(1, abs(want))
        worst = max(worst, err)
        ok &= err <= 1e-6
        print(f"    {name:<28}{a:>14,.6f}{want:>14,.6f}   {'OK' if err <= 1e-6 else 'MISMATCH'}")
    return ok, worst, mr


def variants(c):
    """The test case plus stress variants: each case of the toggle, and a revolver + bullet stress."""
    import copy
    out = [("case 1 (as filed)", c)]
    for k in (2, 3):
        v = copy.deepcopy(c)
        v["case"] = k
        out.append((f"case {k}", v))
    v = copy.deepcopy(c)
    v.update(case=3)
    v["notes"].update(x_ebitda=0.3, maturity=3)
    v["rcf"].update(drawn=20)
    out.append(("stress: RCF draw + bullet in Year 3", v))
    return out


def add_tables(path, T):
    """Create the three real Excel data tables (openpyxl cannot)."""
    if sys.platform != "win32":
        print("MAC BUILD: data tables need Excel, skipped. Finish on the PC: run this same command there.")
        return
    for rng, row_in, col_in in T["tables"]:
        run([os.path.join(HERE, "office.py"), "datatable", path, rng, row_in, col_in], check=True)


def verify(c, out, n, U):
    """Build every variant, recalc in Excel and compare with mirror(). The as-filed case (kept at `out`) also gets
    its data tables, a check that every grid moves on both axes, corner cells re-computed in Python, and Checks = 0."""
    ok_all = True
    cell = lambda ref: ref.replace("$", "")  # noqa: E731
    for i, (name, v) in enumerate(variants(c)):
        path = out if i == 0 else out.replace(".xlsx", f"_v{i}.xlsx")
        I, D, R, T = build(v, path, n, U)
        if i == 0:
            add_tables(path, T)
        rec = json.loads(run([os.path.join(HERE, "office.py"), "recalc", path]).stdout)
        mr = mirror(v, n)
        extra = [("Equity IRR", f"Model!D{R['irr']}", mr["irr"]), ("MoM", f"Model!D{R['mom']}", mr["mom"]),
                 ("Sponsor equity", "Deal!" + cell(D["equity"]), mr["equity"]),
                 ("Exit equity", f"Model!{yc(n)}{R['x_eq']}", mr["exit_eq"])]
        extra += [(f"Bridge – {k}", "Deal!" + cell(D[f"a_{k2}"]), mr["attr"][k])
                  for k, k2 in (("growth", "growth"), ("multiple", "mult"), ("delever", "delev"), ("fees", "fees"))]
        print(f"  {name}")
        ok, worst, _ = compare(v, path, R, n, extra)
        line_ok = ok and rec["status"] == "success"
        print(f"    recalc {rec['status']}, errors {rec['total_errors']}, circular {len(rec['circular_references'])} · "
              f"all years match Python: {'OK' if ok else 'MISMATCH'} (worst rel. err {worst:.1e}) · "
              f"peak RCF {max(mr['rcf']):,.1f}")
        if i == 0:
            st = v.get("steps", {})
            e0, x0 = v["deal"]["entry_multiple"], v["deal"]["exit_multiple"]
            corners = [(T["t1"], mirror(v, n, entry=e0 - 2 * st.get("entry", 0.5), exit_=x0 - 2 * st.get("exit", 0.5))["irr"]),
                       (T["t3"], mirror(v, n, g_adj=-2 * st.get("growth", 0.01), m_adj=-2 * st.get("margin", 0.01))["irr"])]
            grids = {k: f"Deal!{T[k].split(':')[0][0]}{int(T[k].split(':')[0][1:]) + 1}:{T[k].split(':')[1]}"
                     for k in ("t1", "t2", "t3")}
            xl = xl_values(path, list(grids.values()) + [T["check_total"]])
            for k, ref in grids.items():
                body = [row[1:] for row in xl[ref]]
                moves = all(len(set(round(x, 9) for x in row)) == 5 for row in body) and \
                    all(len(set(round(row[j], 9) for row in body)) == 5 for j in range(5))
                line_ok &= moves
                print(f"    data table {k}: {'moves on both axes' if moves else 'FLAT AXIS'}")
            for (rng, want), k in zip(corners, ("t1", "t3")):
                got = xl[grids[k]][0][1]
                good = abs(got - want) < 1e-6
                line_ok &= good
                print(f"    data table {k} corner (low/low) Excel {got:.6f} vs Python {want:.6f}: {'OK' if good else 'MISMATCH'}")
            chk = xl[T["check_total"]]
            line_ok &= abs(chk) < 1e-3
            print(f"    Checks total {chk:.6f}: {'all OK' if abs(chk) < 1e-3 else 'FAILED'}")
        ok_all &= line_ok
        if i:
            os.remove(path)
    return ok_all


if __name__ == "__main__":
    args = sys.argv[1:]
    out = os.path.abspath(args[0])
    opt = lambda k, d: args[args.index(k) + 1] if k in args else d  # noqa: E731
    cfg = json.load(open(opt("--fill", None), encoding="utf-8")) if "--fill" in args else None
    n_years, units = int(opt("--years", 5)), opt("--units", "$m")
    verify_on = "--no-verify" not in args and sys.platform == "win32"
    if cfg and verify_on:
        good = verify(cfg, out, n_years, units)
    else:
        *_, tables = build(cfg, out, n=n_years, U=units)
        add_tables(out, tables)
        good = True
    print(f"built {out}")
    if verify_on:
        res = run([os.path.join(HERE, "office.py"), "recalc", out])
        print(res.stdout)
        lint = run([os.path.join(HERE, "model_lint.py"), out])
        print(lint.stdout)
        good &= "success" in res.stdout and lint.returncode == 0
        print("VERIFIED" if good else "FAILED")
        sys.exit(0 if good else 1)
