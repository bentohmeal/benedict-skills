r"""Layout-choice eval: golden dataset, review sheet, fixed 80/20 split, scoring (references/deck-pipeline.md, "Eval").

  python layout_eval.py seed-pages <out.jsonl> [n]      # stratified real bank pages + their text, for blind chunking
  python layout_eval.py add-chunks <chunks.json>        # chunked real pages -> pending golden rows (bank's layout = label)
  python layout_eval.py add-plan <topics.json> --source NAME [--approved]
                                                        # picked slide plan -> rows (--approved: he OK'd it = accepted)
  python layout_eval.py review                          # pending rows -> golden_review.xlsx in the toolkit folder
  python layout_eval.py ingest                          # his accept / swap / reject decisions -> golden.jsonl
  python layout_eval.py examples [--all]                # train rows -> assets/eval/examples.md (few-shot for select)
  python layout_eval.py matcher [--all]                 # score layout_match.py on the hold-out -> scores.md
  python layout_eval.py blind [--all]                   # blind task for a subagent -> assets/eval/runs/<stamp>/task.md
  python layout_eval.py score <answers.json> [--all]    # score the subagent's answers -> scores.md
  python layout_eval.py stats

Only rows with status "accepted" count; --all also counts "pending" ones (a provisional score, labelled as such).
The split is a hash of the row id (1 in 5 -> hold-out), so it never moves when rows are added. Hold-out rows are
never written into examples.md or a task's examples.

Row: {"id", "source", "ref", "thumb", "section", "message", "block": {what, data_type, about, content, unit?},
      "label", "acceptable": [ids that also count as right], "why", "status", "reviewer", "added", "flag"?}
"""
import datetime
import hashlib
import itertools
import json
import os
import re
import sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import layout_match as lm  # noqa: E402

ASSETS = os.path.join(HERE, "..", "assets")
EVAL = os.path.join(ASSETS, "eval")
GOLDEN = os.path.join(EVAL, "golden.jsonl")
SCORES = os.path.join(EVAL, "scores.md")
EXAMPLES = os.path.join(EVAL, "examples.md")
SWAPS = os.path.join(EVAL, "swaps.md")
RUNS = os.path.join(EVAL, "runs")
PIPELINE = os.path.join(HERE, "..", "references", "deck-pipeline.md")
TOOLKIT = os.environ.get("DEAL_DESK_TOOLKIT",
                         r"C:\Users\user\Documents\Google Drive\Claude Things\output\_Deal Desk Toolkit")
REVIEW_XLSX = os.path.join(TOOLKIT, "golden_review.xlsx")

# layouts that do the same job (a bank's bar is as right as our column): counted as right by default
EQUIV = {"bar": ["column"], "column": ["bar"], "fin_table": ["table"], "table": ["fin_table"],
         "stat_strip": ["stat_icons"], "stat_icons": ["stat_strip"], "chevrons": ["numbered_steps"],
         "numbered_steps": ["chevrons"], "pill_rows": ["side_tabs"], "side_tabs": ["pill_rows"],
         "bullets": ["takeaway"]}
# a bank page's whole-page label -> the block layout that carries it
PAGE_TO_BLOCK = {"page_text": "bullets", "page_risks": "risk_matrix", "page_thesis": "side_tabs",
                 "page_deals": "table", "page_landscape": "logo_row", "page_geography": "map_pins",
                 "page_overview": "stat_icons", "page_market": "bullets"}
SHAPE_WORDS = re.compile(r"\b(\d+|items?|numbers?|figures?|series|unit|dated|dates|periods?|years?|words?|text|"
                         r"rows?|columns?|steps?|stages?|names?|kpi|trend|ranking|rank|compar\w*|composition|mix|"
                         r"range|matrix|grid|process|sequence|cycle|loop|hierarchy|geograph\w*|map|sites?)\b", re.I)


def today():
    return datetime.date.today().isoformat()


def holdout(rid):
    return int(hashlib.sha1(rid.encode()).hexdigest(), 16) % 5 == 0


def load_rows():
    if not os.path.exists(GOLDEN):
        return []
    with open(GOLDEN, encoding="utf-8") as f:
        return [json.loads(x) for x in f if x.strip()]


def save_rows(rows):
    os.makedirs(EVAL, exist_ok=True)
    with open(GOLDEN, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")


def add_rows(new):
    rows = load_rows()
    have = {r["id"] for r in rows}
    added = [r for r in new if r["id"] not in have]
    save_rows(rows + added)
    print(f"{len(added)} rows added ({len(new) - len(added)} already there); golden.jsonl now {len(rows) + len(added)}")
    return added


def usable(rows, include_pending):
    ok = {"accepted", "pending"} if include_pending else {"accepted"}
    return [r for r in rows if r["status"] in ok]


def _affinity(dt, lid, lib):
    r = lib.get(lid, {})
    return 1.0 if r.get("data_type") == dt else lm.NEAR.get(dt, {}).get(r.get("data_type"), 0)


# ---------- seeding ----------

def seed_pages(out, n=0):
    """Stratified pick of real bank pages by lead layout (rare layouts all in, common ones capped), with page text."""
    import pymupdf
    pg = lm._load(lm.PAGES_JSON)
    root, pages = pg["root"], pg["pages"]
    by_lead = defaultdict(list)
    for p in pages:
        by_lead[p["layouts"][0]].append(p)
    cap = {lead: 4 for lead in by_lead}
    for lead, ps in by_lead.items():
        if len(ps) > 30:
            cap[lead] = 6
        elif len(ps) > 9:
            cap[lead] = 5
    picked = []
    for lead, ps in sorted(by_lead.items(), key=lambda kv: len(kv[1])):
        ps = sorted(ps, key=lambda p: (-p["fill"], p["id"]))
        seen_decks = Counter()
        for p in ps:                                  # spread across decks before taking a second page from one
            if len([q for q in picked if q["layouts"][0] == lead]) >= cap[lead]:
                break
            if seen_decks[p["deck"]] >= 1 and len(ps) > cap[lead] * 2:
                continue
            seen_decks[p["deck"]] += 1
            picked.append(p)
    picked = picked[:n] if n else picked
    # most SEC decks are scans with no text layer: render those pages and read them with Windows OCR (ocr.ps1)
    import subprocess
    import tempfile
    texts, scans = {}, tempfile.mkdtemp(prefix="layout_eval_ocr_")
    for p in picked:
        doc = pymupdf.open(os.path.join(root, p["file"]))
        page = doc[p["page"] - 1]
        texts[p["id"]] = page.get_text()
        if len(texts[p["id"]].strip()) < 150:
            page.get_pixmap(matrix=pymupdf.Matrix(2400 / page.rect.width, 2400 / page.rect.width)).save(
                os.path.join(scans, p["id"] + ".png"))
        doc.close()
    if os.listdir(scans):
        ocr_out = os.path.join(scans, "ocr.tsv")
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                        os.path.join(HERE, "deck_library", "ocr.ps1"), scans, ocr_out], check=True)
        with open(ocr_out, encoding="utf-8-sig") as f:
            for ln in f:
                k, _, v = ln.rstrip("\n").partition("\t")
                if v and v != "ERR":
                    texts[k[:-4]] = v.replace(" | ", "\n")
    with open(out, "w", encoding="utf-8") as f:
        for p in picked:
            text = re.sub(r"[ \t]+", " ", re.sub(r"\n{2,}", "\n", texts[p["id"]])).strip()[:2400]
            f.write(json.dumps({"page_id": p["id"], "title": p["title"], "section": p["section"],
                                "exhibits": len(p["layouts"]), "text": text}, ensure_ascii=False) + "\n")
    print(f"{len(picked)} pages -> {out}")
    print(Counter(p["layouts"][0] for p in picked).most_common())


def add_chunks(path):
    """Chunker output [{"page_id", "section", "message", "blocks": [...]}] -> pending rows; block i <-> bank layout i."""
    lib = lm.Matcher().lib
    pg = lm._load(lm.PAGES_JSON)
    pages = {p["id"]: p for p in pg["pages"]}
    new = []
    for t in lm._load(path):
        p = pages[t["page_id"]]
        labs = [PAGE_TO_BLOCK.get(x, x) for x in p["layouts"]]
        # the bank tags say which layouts are on the page, not which exhibit is which: pair them by best data-type fit
        # (ties keep reading order)
        best = max(itertools.permutations(labs, min(len(labs), len(t["blocks"]))),
                   key=lambda perm: (sum(_affinity(b.get("data_type", ""), lab, lib) for b, lab in zip(t["blocks"], perm)),
                                     -sum(labs.index(lab) != i for i, lab in enumerate(perm))))
        for i, (b, lab) in enumerate(zip(t["blocks"], best)):
            row = {"id": f"rp_{p['id']}_b{i + 1}", "source": "real_page",
                   "ref": f"{p['bank']} · {p['deck']} p{p['page']}", "thumb": os.path.join(pg["root"], p["thumb"]),
                   "section": t.get("section") or p["section"], "message": t.get("message", p["title"]),
                   "block": b, "label": lab, "acceptable": EQUIV.get(lab, []),
                   "why": f"bank's choice: {p['note']}", "status": "pending", "reviewer": "", "added": today()}
            if _affinity(b.get("data_type", ""), lab, lib) < 0.3:
                row["flag"] = f"check: chunk says {b.get('data_type')}, bank layout is {lab}"
            new.append(row)
    add_rows(new)


def add_plan(path, source, approved=False):
    """Slide plan with picks -> rows. With --approved (he OK'd the plan table) they go straight in as accepted."""
    m = lm.Matcher()
    doc = lm._load(path)
    slug = re.sub(r"[^a-z0-9]+", "_", source.lower()).strip("_")
    new, swaps = [], []
    for t in doc["topics"]:
        for i, b in enumerate(t["blocks"]):
            pk = b.get("pick")
            if not pk:
                continue
            blk = {k: b[k] for k in ("what", "data_type", "about", "content", "unit") if k in b}
            top = m.candidates(b, t["section"], 1)
            top1 = top[0]["id"] if top else ""
            row = {"id": f"{slug}_s{int(t['slide']):02d}_b{i + 1}", "source": f"{'plan' if approved else 'case'}:{source}",
                   "ref": f"{doc.get('deck', source)} slide {t['slide']}", "thumb": "",
                   "section": t["section"], "message": t["message"], "block": blk, "label": pk,
                   "acceptable": sorted(set(b.get("acceptable", EQUIV.get(pk, [])))),
                   "why": b.get("why", ""), "status": "accepted" if approved else "pending",
                   "reviewer": "Benedict (plan approval)" if approved else "", "added": today(), "matcher_top1": top1}
            if b.get("swapped_from"):
                row["swapped_from"] = b["swapped_from"]
            new.append(row)
            if approved and (b.get("swapped_from") or (top1 and pk != top1 and top1 not in row["acceptable"])):
                swaps.append((row, b.get("swapped_from") or top1))
    add_rows(new)
    _log_swaps(swaps)


def _log_swaps(swaps):
    """Lessons loop: each swap is a candidate 'dont_use_when' line for the layout that lost."""
    if not swaps:
        return
    new = not os.path.exists(SWAPS)
    with open(SWAPS, "a", encoding="utf-8") as f:
        if new:
            f.write("# Layout swaps (lessons loop)\nEach line: the layout the matcher or Claude offered, what he chose "
                    "instead, and the block. Two or more swaps of the same pair -> add a `dont_use_when` line to the "
                    "losing layout in layout_tag.py, rebuild, and re-run the eval.\n\n")
        for row, lost in swaps:
            b = row["block"]
            f.write(f"- {row['added']} · {lost} -> **{row['label']}** · {b.get('what', '')} ({b.get('data_type', '')}, "
                    f"{len(b.get('content', []))} items) · {row['ref']}\n")
    pairs = Counter((lost, row["label"]) for row, lost in swaps)
    print("swaps logged to eval/swaps.md:", ", ".join(f"{a}->{b} x{n}" for (a, b), n in pairs.items()))


# ---------- his review ----------

def review():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    lib = lm.Matcher().lib
    rows = [r for r in load_rows() if r["status"] == "pending"]
    rows.sort(key=lambda r: (not holdout(r["id"]), not r.get("flag"), r["id"]))   # what the score rests on first
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws["A1"] = "Layout eval: review proposed labels"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A2"] = ("Hold-out rows (they set the score) come first, then flagged ones. Decision: accept (the label is right), swap (put the right layout id in 'Swap to'), reject (bad row). "
                "Blank = still pending. Then tell Claude 'ingest the layout review'.")
    ws["A2"].font = Font(italic=True, size=9, color="FFFFFF")
    for c in ("A1", "A2"):
        ws[c].alignment = Alignment(vertical="center")
    band = PatternFill("solid", fgColor="1F3864")
    for row in (1, 2):
        for col in range(1, 15):
            ws.cell(row, col).fill = band
    head = ["Row id", "Source", "Page / slide", "Section", "Slide message", "Block", "Type", "Items",
            "Content (first lines)", "Proposed layout", "Also right", "Check", "Decision", "Swap to"]
    for j, h in enumerate(head, 1):
        c = ws.cell(4, j, h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    dv = DataValidation(type="list", formula1='"accept,swap,reject"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, r in enumerate(rows, 5):
        b = r["block"]
        content = b.get("content", [])
        vals = [r["id"], r["source"], r["ref"], r["section"], r["message"], b.get("what", ""), b.get("data_type", ""),
                len(content) if isinstance(content, list) else "", " · ".join(map(str, content[:4]))[:220]
                if isinstance(content, list) else str(content)[:220], r["label"], ", ".join(r["acceptable"]),
                ("HOLD-OUT. " if holdout(r["id"]) else "") + r.get("flag", ""), "", ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if r.get("thumb") and os.path.exists(r["thumb"]):
            ws.cell(i, 3).hyperlink = r["thumb"]
            ws.cell(i, 3).font = Font(color="0563C1", underline="single")
        thumb = lib.get(r["label"], {}).get("thumb")
        if thumb:
            ws.cell(i, 10).hyperlink = os.path.abspath(os.path.join(ASSETS, thumb))
            ws.cell(i, 10).font = Font(color="0563C1", underline="single")
        if r.get("flag") or holdout(r["id"]):
            ws.cell(i, 12).font = Font(bold=True, color="C00000")
        dv.add(ws.cell(i, 13))
    for col, w in zip("ABCDEFGHIJKLMN", (20, 12, 30, 11, 34, 26, 12, 6, 50, 15, 15, 24, 10, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    os.makedirs(TOOLKIT, exist_ok=True)
    wb.save(REVIEW_XLSX)
    print(f"{len(rows)} pending rows -> {REVIEW_XLSX}")


def ingest():
    from openpyxl import load_workbook
    lib = lm.Matcher().lib
    ws = load_workbook(REVIEW_XLSX).active
    rows = load_rows()
    by_id = {r["id"]: r for r in rows}
    n, swaps, bad = Counter(), [], []
    for i in range(5, ws.max_row + 1):
        rid, dec, to = ws.cell(i, 1).value, (ws.cell(i, 13).value or "").strip().lower(), ws.cell(i, 14).value
        r = by_id.get(rid)
        if not r or not dec:
            continue
        if dec == "swap":
            to = (to or "").strip()
            if to not in lib:
                bad.append(f"{rid}: swap to {to!r} is not a layout id")
                continue
            swaps.append((r, r["label"]))
            r["why"] = f"{r['why']} | swapped from {r['label']} in review"
            r["label"], r["acceptable"] = to, EQUIV.get(to, [])
            dec = "accept"
        r["status"] = {"accept": "accepted", "reject": "rejected"}.get(dec, r["status"])
        r["reviewer"] = "Benedict"
        r["reviewed"] = today()
        r.pop("flag", None)
        n[dec] += 1
    save_rows(rows)
    _log_swaps([(r, lost) for r, lost in swaps])
    print(dict(n), "| still pending:", sum(r["status"] == "pending" for r in rows))
    for b in bad:
        print("SKIPPED", b)


# ---------- examples, scoring ----------

def _line(r):
    b = r["block"]
    sh = lm.shape(b)
    unit = f", unit {b['unit']}" if b.get("unit") else ""
    alts = f" (also right: {', '.join(r['acceptable'])})" if r["acceptable"] else ""
    return (f"- [{r['section']}] {b.get('what', '')}: {sh['data_type']}, {sh['items']} items, {sh['words']} words, "
            f"{'numbers' if sh['numbers'] else 'no numbers'}{unit} -> **{r['label']}**{alts}. {r['why'][:140]}")


def examples(include_pending=False):
    rows = [r for r in usable(load_rows(), include_pending) if not holdout(r["id"])]
    lines = ["# Layout choice examples (train split)",
             f"Built {today()} by `layout_eval.py examples` from {len(rows)} reviewed rows"
             + (" (includes unreviewed rows: provisional)" if include_pending else "") +
             ". Hold-out rows are never listed here. Read before the select step of deck-pipeline.md.", ""]
    by_dt = defaultdict(list)
    for r in rows:
        by_dt[r["block"].get("data_type", "")].append(r)
    for dt in sorted(by_dt):
        lines.append(f"## {dt}")
        lines += [_line(r) for r in sorted(by_dt[dt], key=lambda r: r["id"])]
        lines.append("")
    with open(EXAMPLES, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"{len(rows)} examples -> {EXAMPLES}")


def _test(include_pending):
    rows = usable(load_rows(), include_pending)
    test = [r for r in rows if holdout(r["id"])]
    if not test:
        sys.exit("no hold-out rows yet: review some rows (or pass --all for a provisional score)")
    return test


def _rubric(r, ans, lib):
    """Reasoning check: names the content shape, and the pick does not break its own 'don't use when'."""
    why = ans.get("why", "")
    cites = bool(SHAPE_WORDS.search(why))
    sh = lm.shape(r["block"])
    pk = ans.get("pick", "")
    dont = lm._sim(sh["query"], lm.toks(re.sub(r"\(use [^)]*\)", "", lib.get(pk, {}).get("dont_use_when", ""))))
    clean = dont < 0.25 and not (pk in lm.NUMERIC and not sh["numbers"])
    return (cites + clean) / 2


def _log(selector, rows, include_pending, top1, top3, rubric, note=""):
    new = not os.path.exists(SCORES)
    accepted = sum(r["status"] == "accepted" for r in rows)
    with open(SCORES, "a", encoding="utf-8") as f:
        if new:
            f.write("# Layout eval scores\nHold-out = the fixed 1-in-5 split of golden.jsonl. Keep a prompt or "
                    "library change only if it does not lower top-1 or top-3 on the same set.\n\n"
                    "| Date | Selector | Set | n (accepted) | Top-1 | Top-3 | Reasoning | Note |\n"
                    "|---|---|---|---|---|---|---|---|\n")
        rub = f"{rubric:.0%}" if rubric is not None else "n/a"
        f.write(f"| {today()} | {selector} | {'provisional (incl. pending)' if include_pending else 'accepted'} | "
                f"{len(rows)} ({accepted}) | {top1:.0%} | {top3:.0%} | {rub} | {note} |\n")


def _hit(r, ids):
    ok = {r["label"], *r["acceptable"]}
    return any(i in ok for i in ids)


def matcher(include_pending=False, note=""):
    m = lm.Matcher()
    test = _test(include_pending)
    t1 = t3 = 0
    misses = []
    for r in test:
        ids = [c["id"] for c in m.candidates(r["block"], r["section"], 3)]
        t1 += _hit(r, ids[:1])
        t3 += _hit(r, ids)
        if not _hit(r, ids[:1]):
            misses.append(f"  {r['id']}: label {r['label']}, matcher {', '.join(ids)}")
    n = len(test)
    print(f"matcher on {n} hold-out rows: top-1 {t1 / n:.0%}, top-3 {t3 / n:.0%}")
    print("\n".join(misses))
    _log("layout_match.py", test, include_pending, t1 / n, t3 / n, None, note)


def blind(include_pending=False):
    """Task file for a blind selector (a subagent): pipeline select rules + catalogue + train examples + hold-out blocks."""
    lib = lm.Matcher().lib
    rows = usable(load_rows(), include_pending)
    train = [r for r in rows if not holdout(r["id"])]
    test = [r for r in rows if holdout(r["id"])]
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    out = os.path.join(RUNS, stamp)
    os.makedirs(out, exist_ok=True)
    with open(PIPELINE, encoding="utf-8") as f:
        pipe = f.read()
    rules = pipe.split("## 4.")[0]
    cat = []
    for r in lib.values():
        if r.get("kind") not in ("layout", "map", "icon_set"):
            continue
        lo, hi = (r.get("items") or [1, 99])[:2]
        cat.append(f"- `{r['id']}` ({r.get('data_type', '')}, {lo}-{hi} items{', numbers' if r.get('numbers') else ''})"
                   f": {r.get('use_when', '')} | Don't: {r.get('dont_use_when', '')}")
    qs = []
    for i, r in enumerate(test, 1):
        sh = lm.shape(r["block"])
        qs.append({"q": f"q{i:02d}", "section": r["section"], "message": r["message"], "block": r["block"],
                   "shape": {k: sh[k] for k in ("data_type", "items", "words", "numbers", "series", "dated")}})
    with open(os.path.join(out, "task.md"), "w", encoding="utf-8") as f:
        f.write("# Blind layout selection\nYou are the select step of a slide pipeline. For each question, pick the one "
                "layout id that best carries the block, two alternatives, and a one-line reason that names the content "
                "shape (data type, item count, numbers) and why the pick beats the alternatives.\n"
                "Use only this file. Do not open golden.jsonl, scores.md or any other eval file.\n\n"
                f"Write `{os.path.join(out, 'answers.json')}` as "
                '{"q01": {"pick": "...", "alts": ["...", "..."], "why": "..."}, ...}.\n\n'
                "## Pipeline rules\n" + rules + "\n## Layout catalogue (id, data type, items: use when | don't)\n"
                + "\n".join(cat) + "\n\n## Worked examples (reviewed choices)\n" + "\n".join(_line(r) for r in train)
                + "\n\n## Questions\n```json\n" + json.dumps(qs, indent=1, ensure_ascii=False) + "\n```\n")
    with open(os.path.join(out, "key.json"), "w", encoding="utf-8") as f:
        json.dump({q["q"]: r["id"] for q, r in zip(qs, test)}, f, indent=1)
    print(f"{len(test)} questions, {len(train)} examples -> {os.path.join(out, 'task.md')}")


def score(answers_path, include_pending=False, note=""):
    lib = lm.Matcher().lib
    answers = lm._load(answers_path)
    key = lm._load(os.path.join(os.path.dirname(answers_path), "key.json"))
    by_id = {r["id"]: r for r in load_rows()}
    test = [by_id[key[q]] for q in key]
    t1 = t3 = rub = 0
    for q, rid in key.items():
        r, a = by_id[rid], answers.get(q, {})
        ids = [a.get("pick", "")] + list(a.get("alts", []))[:2]
        t1 += _hit(r, ids[:1])
        t3 += _hit(r, ids)
        rub += _rubric(r, a, lib)
        if not _hit(r, ids[:1]):
            print(f"  {q} {rid}: label {r['label']}, picked {', '.join(ids)} | {a.get('why', '')[:90]}")
    n = len(key)
    print(f"blind selector on {n}: top-1 {t1 / n:.0%}, top-3 {t3 / n:.0%}, reasoning {rub / n:.0%}")
    _log("blind subagent", test, include_pending, t1 / n, t3 / n, rub / n, note)


def stats():
    rows = load_rows()
    print("status:", dict(Counter(r["status"] for r in rows)))
    print("source:", dict(Counter(r["source"].split(":")[0] for r in rows)))
    print("hold-out:", sum(holdout(r["id"]) for r in rows), "of", len(rows))
    print("labels:", Counter(r["label"] for r in rows).most_common())
    print("flagged:", sum(bool(r.get("flag")) for r in rows))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    a = sys.argv[1:]
    allp = "--all" in a
    note = a[a.index("--note") + 1] if "--note" in a else ""
    cmd = a[0] if a else ""
    if cmd == "seed-pages" and len(a) > 1:
        seed_pages(a[1], int(a[2]) if len(a) > 2 else 0)
    elif cmd == "add-chunks" and len(a) > 1:
        add_chunks(a[1])
    elif cmd == "add-plan" and len(a) > 1 and "--source" in a:
        add_plan(a[1], a[a.index("--source") + 1], "--approved" in a)
    elif cmd == "review":
        review()
    elif cmd == "ingest":
        ingest()
    elif cmd == "examples":
        examples(allp)
    elif cmd == "matcher":
        matcher(allp, note)
    elif cmd == "blind":
        blind(allp)
    elif cmd == "score" and len(a) > 1:
        score(a[1], allp, note)
    elif cmd == "stats":
        stats()
    else:
        sys.exit(__doc__)
