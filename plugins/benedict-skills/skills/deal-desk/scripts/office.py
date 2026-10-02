"""Drive the real Excel / PowerPoint on the Windows PC; on the Mac fall back to LibreOffice (see "Mac" below).

Windows: full Excel / PowerPoint through COM (exact output, real What-If Data Tables).
Mac / Linux: no COM, so `recalc`, `values`, `render`, `pdf`, `preview`, `density` run through LibreOffice headless
  (`brew install --cask libreoffice`). Slides and sheets look close to Office, not identical, and recalc is a CHECK
  only: it never overwrites the model, so cached values are saved by the PC's recalc. `datatable` needs real Excel,
  so it is PC-only: build and style on either machine, then run `datatable` + `recalc` on the PC (the file is in Drive).

Usage (use the full Python path, see SKILL.md):
  python office.py recalc  model.xlsx            -> full recalc in Excel, save, JSON report of errors / circular refs
  python office.py values  model.xlsx "Sheet!A1" "Sheet!B2:D2" ...  -> recalc (no save) and print values
  python office.py render  deck.pptx  out_dir    -> one PNG per slide (1600x900) for visual QA
  python office.py pdf     file.pptx|file.xlsx out.pdf
  python office.py density deck.pptx [out_dir]  -> fill + empty band per slide vs the bank standard (exit 1 if sparse)
  python office.py preview model.xlsx out_dir [Sheet ...]  -> PNG per printed page of each sheet (look at the layout)
                                                  a sheet can be "Debt!A1:T70" to preview only that range (wide sheets)
  python office.py datatable model.xlsx "Sheet!B5:G10" "Sheet!rowInput" "Sheet!colInput"
                                                -> create a real Excel What-If Data Table on that range, save
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

WIN = sys.platform == "win32"
if WIN:
    import pythoncom
    import win32com.client as win32

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


_PIDS = {}   # id(Excel COM object) -> process id, filled by _excel()


def _abs(p):
    return os.path.abspath(p)


def _excel():
    pythoncom.CoInitialize()
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    xl.AskToUpdateLinks = False
    try:   # remember this instance's PID so _quit() can make sure it is gone (Quit alone leaves hidden EXCEL.EXE behind)
        import win32process
        _PIDS[id(xl)] = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
    except Exception:
        pass
    return xl


def _ppt_running():
    import subprocess
    return "POWERPNT.EXE" in subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/NH"],
                                            capture_output=True, text=True).stdout


def _powerpoint():
    """PowerPoint is single-instance: DispatchEx attaches to Benedict's open PowerPoint if there is one. Only an
    instance this script started is remembered for _quit_ppt() to close (and kill if it lingers)."""
    pythoncom.CoInitialize()
    was_running = _ppt_running()
    pp = win32.DispatchEx("PowerPoint.Application")
    if not was_running:
        try:
            import win32process
            _PIDS[id(pp)] = win32process.GetWindowThreadProcessId(pp.HWND)[1]
        except Exception:
            _PIDS[id(pp)] = None
    return pp


def _quit_ppt(pp):
    """Close PowerPoint only if _powerpoint() started it; never touch Benedict's own open session."""
    if id(pp) in _PIDS:
        _quit(pp)


def _quit(app):
    """Quit Excel / PowerPoint, then end the process if COM left it running (hidden instances blocked Benedict's
    Excel from opening files, 2026-09-27)."""
    pid = _PIDS.pop(id(app), None)
    try:
        app.Quit()
    except Exception:
        pass
    if pid:
        import subprocess
        import time
        for _ in range(20):   # up to ~5 s for a clean exit
            alive = str(pid) in subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/NH"], capture_output=True,
                                               text=True).stdout
            if not alive:
                return
            time.sleep(0.25)
        subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True)


# ------------------------------------------------------------------ Mac / Linux: LibreOffice instead of COM
def _soffice():
    for c in (shutil.which("soffice"), shutil.which("libreoffice"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and os.path.exists(c):
            return c
    sys.exit("LibreOffice not found. Install it once: brew install --cask libreoffice (see MAC-SETUP.md). "
             "Or run this step on the Windows PC.")


def _lo_convert(src, fmt, out_dir):
    """Convert with LibreOffice headless into out_dir; returns the new file's path."""
    os.makedirs(out_dir, exist_ok=True)
    prof = tempfile.mkdtemp(prefix="lo_profile_")        # private profile: avoids lock-ups if LibreOffice is open
    r = subprocess.run([_soffice(), f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", fmt,
                        "--outdir", out_dir, _abs(src)], capture_output=True, text=True, timeout=300)
    shutil.rmtree(prof, ignore_errors=True)
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + "." + fmt.split(":")[0])
    if not os.path.exists(out):
        sys.exit(f"LibreOffice conversion failed: {r.stderr or r.stdout}")
    return out


def _lo_pngs(pdf_path, out_dir, stem, width=None, dpi=None):
    import pymupdf
    doc, files = pymupdf.open(pdf_path), []
    for i, page in enumerate(doc, 1):
        f = os.path.join(out_dir, f"{stem}-{i:02d}.png")
        m = pymupdf.Matrix(width / page.rect.width, width / page.rect.width) if width else pymupdf.Matrix(dpi / 72, dpi / 72)
        page.get_pixmap(matrix=m).save(f)
        files.append(f)
    doc.close()
    return files


def _lo_recalculated(path):
    """A LibreOffice-recalculated COPY of a workbook (the original is never touched) -> path of the copy."""
    tmp = tempfile.mkdtemp(prefix="recalc_")
    return _lo_convert(path, "xlsx:Calc MS Excel 2007 XML", tmp)


def _split(ref):
    sheet, addr = ref.rsplit("!", 1)
    return sheet.strip("'"), addr


def recalc(path):
    """Full recalculation in Excel, save, then scan cached values with openpyxl."""
    from openpyxl import load_workbook

    path = _abs(path)
    circular = []
    if not WIN:
        print("(Mac: recalculated in LibreOffice as a check; the model file is left unchanged. "
              "Run recalc on the PC to save values.)", file=sys.stderr)
        vals_path = _lo_recalculated(path)
    else:
        vals_path = path
        xl = _excel()
    try:
        if not WIN:
            raise StopIteration
        wb = xl.Workbooks.Open(path, UpdateLinks=0)
        xl.CalculateFull()
        for ws in wb.Worksheets:
            try:
                c = ws.CircularReference
                if c is not None:
                    circular.append(f"{ws.Name}!{c.Address}")
            except Exception:
                pass
        wb.Save()
        wb.Close(False)
    except StopIteration:
        pass
    finally:
        if WIN:
            _quit(xl)

    vals = load_workbook(vals_path, data_only=True)
    forms = load_workbook(path, data_only=False)
    errors, n_formulas = {}, 0
    for ws in forms.worksheets:
        wv = vals[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formulas += 1
                v = wv[c.coordinate].value
                if isinstance(v, str) and v.strip() in ERRORS:
                    errors.setdefault(v.strip(), []).append(f"{ws.title}!{c.coordinate}")
    status = "success" if not errors and not circular else "errors_found"
    report = {
        "status": status,
        "total_formulas": n_formulas,
        "total_errors": sum(len(v) for v in errors.values()),
        "errors": {k: {"count": len(v), "locations": v[:50]} for k, v in errors.items()},
        "circular_references": circular,
    }
    print(json.dumps(report, indent=2))
    return report


def values(path, refs):
    if not WIN:
        from openpyxl import load_workbook
        wbv = load_workbook(_lo_recalculated(path), data_only=True)
        out = {}
        for ref in refs:
            sheet, addr = _split(ref)
            cells = wbv[sheet][addr]
            out[ref] = (cells.value if not isinstance(cells, tuple)
                        else [[c.value for c in row] if isinstance(row, tuple) else row.value for row in cells])
        print(json.dumps(out, indent=2, default=str))
        return out
    xl = _excel()
    out = {}
    try:
        wb = xl.Workbooks.Open(_abs(path), UpdateLinks=0, ReadOnly=True)
        xl.CalculateFull()
        for ref in refs:
            sheet, addr = _split(ref)
            v = wb.Worksheets(sheet).Range(addr).Value
            out[ref] = v
        wb.Close(False)
    finally:
        _quit(xl)
    print(json.dumps(out, indent=2, default=str))
    return out


def datatable(path, table_ref, row_input_ref, col_input_ref):
    """Range.Table(RowInput, ColumnInput): top row of the range = row-input values,
    left column = column-input values, top-left cell = the output formula.
    Input cells must be on the same sheet as the table (Excel rule)."""
    if not WIN:
        sys.exit("PC-ONLY STEP: What-If Data Tables need real Excel (COM). Everything else is built; open this "
                 "workbook on the Windows PC (it is in Drive) and run: office.py datatable ... then office.py recalc.")
    xl = _excel()
    try:
        wb = xl.Workbooks.Open(_abs(path), UpdateLinks=0)
        sheet, addr = _split(table_ref)
        ws = wb.Worksheets(sheet)
        row_in = ws.Range(_split(row_input_ref)[1]) if row_input_ref else None
        col_in = ws.Range(_split(col_input_ref)[1]) if col_input_ref else None
        ws.Range(addr).Table(row_in, col_in)
        xl.CalculateFull()
        wb.Save()
        wb.Close(False)
    finally:
        _quit(xl)
    print(f"data table created on {table_ref}")


def render(pptx, out_dir, width=1600, height=900):
    out_dir = _abs(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    if not WIN:
        tmp = tempfile.mkdtemp(prefix="render_")
        files = _lo_pngs(_lo_convert(pptx, "pdf", tmp), out_dir, "slide", width=width)
        print("\n".join(files))
        return files
    pp = _powerpoint()
    files = []
    try:
        pres = pp.Presentations.Open(_abs(pptx), True, False, False)  # ReadOnly, Untitled, WithWindow
        for i in range(1, pres.Slides.Count + 1):
            f = os.path.join(out_dir, f"slide-{i:02d}.png")
            pres.Slides(i).Export(f, "PNG", width, height)
            files.append(f)
        pres.Close()
    finally:
        _quit_ppt(pp)
    print("\n".join(files))
    return files


def pdf(src, out):
    src, out = _abs(src), _abs(out)
    if not WIN:
        tmp = tempfile.mkdtemp(prefix="pdf_")
        shutil.move(_lo_convert(src, "pdf", tmp), out)
        print(out)
        return
    if src.lower().endswith((".pptx", ".ppt")):
        pp = _powerpoint()
        try:
            pres = pp.Presentations.Open(src, True, False, False)
            pres.SaveAs(out, 32)  # ppSaveAsPDF
            pres.Close()
        finally:
            _quit_ppt(pp)
    else:
        xl = _excel()
        try:
            wb = xl.Workbooks.Open(src, UpdateLinks=0, ReadOnly=True)
            wb.ExportAsFixedFormat(0, out)  # xlTypePDF
            wb.Close(False)
        finally:
            _quit(xl)
    print(out)


# Density standard (lessons #32). Measured on 400 pages of 14 real bank decks (GS, JPM, MS, Citi, BofA, Barclays,
# Evercore, Centerview, Moelis, Qatalyst, PJT, PWP, HL; SEC filings via 10xebitda): body fill median 0.75, 25th pct
# 0.62, and no fully empty band. Qatalyst / GS market pages: 0.83-0.93.
FILL_MIN, FILL_TARGET, GAP_MAX = 0.70, 0.80, 2   # gap = fully empty grid rows (of 14) in the body
GRID_X, GRID_Y = 24, 14


def page_density(png):
    """(fill, empty rows) for one rendered slide/page: body = 15-88% of the height, 3-97% of the width, split into a
    24 x 14 grid; a cell counts as filled when >0.4% of its pixels are non-white. Same method as the research."""
    import numpy as np
    import pymupdf
    pix = pymupdf.Pixmap(png)
    a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)[:, :, :3]
    ink = a.min(axis=2) < 235
    h, w = ink.shape
    body = ink[int(.15 * h):int(.88 * h), int(.03 * w):int(.97 * w)]
    bh, bw = body.shape
    cells = [[body[int(i * bh / GRID_Y):int((i + 1) * bh / GRID_Y), int(j * bw / GRID_X):int((j + 1) * bw / GRID_X)].mean()
              > 0.004 for j in range(GRID_X)] for i in range(GRID_Y)]
    run = best = 0
    for row in cells:
        run = 0 if any(row) else run + 1
        best = max(best, run)
    return float(np.mean(cells)), best


def density(pptx, out_dir=None, skip_first=True):
    """Render the deck and check every content slide against the bank standard. Returns the failing slides."""
    import statistics
    import tempfile
    out_dir = out_dir or tempfile.mkdtemp(prefix="density_")
    files = render(pptx, out_dir) if pptx.lower().endswith((".pptx", ".ppt")) else sorted(
        os.path.join(pptx, f) for f in os.listdir(pptx) if f.endswith(".png"))
    rows, bad = [], []
    for i, f in enumerate(files, 1):
        if skip_first and i == 1:
            continue
        fill, gap = page_density(f)
        ok = fill >= FILL_MIN and gap < GAP_MAX
        rows.append(fill)
        if not ok:
            bad.append(i)
        print(f"slide {i:2d}: fill {fill:.2f}  empty rows {gap:2d}  {'OK' if ok else 'SPARSE'}")
    med = statistics.median(rows) if rows else 0
    print(f"median fill {med:.2f} (target {FILL_TARGET:.2f}; banks 0.75) | failing slides: {bad or 'none'}")
    return bad


def preview(path, out_dir, sheets=None, dpi=110):
    """Excel sheets -> PDF (landscape, 1 page wide) -> PNG per page, so a model's layout can be looked at."""
    import pymupdf

    out_dir = _abs(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    if not WIN:
        if sheets:
            print("(Mac: per-sheet / range previews are PC-only; showing every page of the workbook)", file=sys.stderr)
        tmp = tempfile.mkdtemp(prefix="preview_")
        files = _lo_pngs(_lo_convert(path, "pdf", tmp), out_dir, os.path.splitext(os.path.basename(path))[0], dpi=dpi)
        print("\n".join(files))
        return files
    xl = _excel()
    pdfs = []
    try:
        wb = xl.Workbooks.Open(_abs(path), UpdateLinks=0, ReadOnly=True)
        names = sheets or [ws.Name for ws in wb.Worksheets]
        for spec in names:
            name, area = (spec.split("!", 1) + [None])[:2]
            ws = wb.Worksheets(name)
            ps = ws.PageSetup
            if area:
                ps.PrintArea = area
            ps.Orientation, ps.Zoom, ps.FitToPagesWide, ps.FitToPagesTall = 2, False, 1, False
            f = os.path.join(out_dir, f"{name}.pdf")
            ws.ExportAsFixedFormat(0, f)
            pdfs.append((name, f))
        wb.Close(False)
    finally:
        _quit(xl)
    files = []
    for name, f in pdfs:
        doc = pymupdf.open(f)
        for i, page in enumerate(doc, 1):
            png = os.path.join(out_dir, f"{name}-{i:02d}.png")
            page.get_pixmap(dpi=dpi).save(png)
            files.append(png)
        doc.close()
        os.remove(f)
    print("\n".join(files))
    return files


if __name__ == "__main__":
    cmd, args = sys.argv[1], sys.argv[2:]
    if cmd == "recalc":
        r = recalc(args[0])
        sys.exit(0 if r["status"] == "success" else 1)
    elif cmd == "values":
        values(args[0], args[1:])
    elif cmd == "render":
        render(args[0], args[1])
    elif cmd == "pdf":
        pdf(args[0], args[1])
    elif cmd == "density":
        sys.exit(1 if density(args[0], args[1] if len(args) > 1 else None) else 0)
    elif cmd == "preview":
        preview(args[0], args[1], args[2:] or None)
    elif cmd == "datatable":
        datatable(args[0], args[1], args[2] or None, args[3] if len(args) > 3 else None)
    else:
        sys.exit(__doc__)
