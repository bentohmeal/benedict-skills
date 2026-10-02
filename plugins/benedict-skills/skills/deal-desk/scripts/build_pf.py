"""Build a project finance model (SMBC-type greenfield solar / power PPA) from a JSON case file (house style v2).

    python build_pf.py case.json out.xlsx              # build, solve debt/IDC/DSRA, recalc, cross-check in Python, lint
    python build_pf.py --resolve model.xlsx            # after changing inputs in Excel: re-solve the pasted values

Sheet order (FAST / Mazars / Gridlines practice): Cover · Inputs · Timing · Construction (monthly) · Operations ·
Debt · Waterfall (quarterly, same columns) · Annual (IS / CF / BS) · Returns (dated XIRR) · Scenarios · Dashboard · Checks.
Debt sized at MIN(gearing × total project cost, PV of CFADS / min DSCR); sculpted (flat DSCR) or straight-line;
DSRA (next N quarters of debt service) funded at COD as part of project cost, topped up / released each quarter;
DSCR, LLCR, PLCR; project and equity XIRR; scenario table re-solved by the engine; tariff × yield data table.
No circular references: debt size and total project cost are solved in Python and pasted (yellow input cells);
the Checks sheet proves pasted = live. See ../references/case-project-finance.md.
"""
import copy
import datetime as dt
import json
import os
import re
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlstyle import (CHECK, DATE, FIRST_PERIOD, FLAG, MULT, NAVY, NUM, NUM1, PCT, STEEL, STYLE_VERSION, TOTAL, check_cell, col,  # noqa: E402
                     colour_key, comment, end_marker, header_row, input_cell, label, note, output_box, put, section,
                     sensitivity, setup, sheet_header, workbook_defaults)

HERE = os.path.dirname(os.path.abspath(__file__))
M1 = NUM1
U = "USD m"


# ------------------------------------------------------------------ date helpers (mirror of Excel EOMONTH)
def eom(d, months):
    y, m = divmod(d.month - 1 + months, 12)
    y, m = d.year + y, m + 1
    nxt = dt.date(y + (m == 12), m % 12 + 1, 1)
    return nxt - dt.timedelta(days=1)


def xnpv(rate, cfs, dates):
    d0 = dates[0]
    return sum(cf / (1 + rate) ** ((d - d0).days / 365) for cf, d in zip(cfs, dates))


def xirr(cfs, dates):
    lo, hi = -0.99, 5.0
    for _ in range(300):
        mid = (lo + hi) / 2
        if xnpv(lo, cfs, dates) * xnpv(mid, cfs, dates) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def defaults(c):
    c.setdefault("tax_rate", 0.0)
    c.setdefault("dsra", {"on": 0, "quarters": 2})
    c["dsra"].setdefault("quarters", 2)
    c.setdefault("hours_per_year", 8760)
    return c


# ------------------------------------------------------------------ Python engine (solver + independent check)
class Engine:
    def __init__(self, c):
        self.c = c = defaults(c)
        self.start = dt.date(*c["start"])
        self.cm = c["construction_months"]
        self.cod = eom(self.start, self.cm - 1)
        self.ppa_end = eom(self.cod, 12 * c["ppa_years"])
        self.first_rep = eom(self.cod, c["first_repayment_months_after_cod"])
        self.maturity = eom(self.start, 12 * c["debt"]["tenor_years"] - 1)
        nq = (self.ppa_end.year - self.start.year + 1) * 4
        self.q = [eom(self.start, 3 * k + 2) for k in range(nq)]
        self.months = [self.start] + [eom(self.start, m - 1) for m in range(1, self.cm + 1)]
        dd = c["debt"]
        self.r_pre, self.r_post = dd["swap"] + dd["margin_pre"], dd["swap"] + dd["margin_post"]
        self.k_cod = next(k for k, d in enumerate(self.q) if d == self.cod) if self.cod in self.q else None

    def ops(self, tariff=None, dep_base=0.0):
        c = self.c
        tariff = c["tariff"] if tariff is None else tariff
        cuf = c["cuf"][c["energy_case"]]
        dep_q = dep_base / (c["depreciation_years"] * 4)
        rows = []
        for d in self.q:
            op = self.cod < d <= self.ppa_end
            oy = d.year - self.cod.year if op else 0
            qn = (d.month - 1) // 3
            mwh = c["capacity_mw"] * c["hours_per_year"] * cuf * c["seasonality"][qn] * (1 - c["degradation"]) ** max(oy - 1, 0) if op else 0
            idx = (1 + c["om"]["escalation"]) ** (d.year - c["om"]["base_year"])
            rev = mwh * tariff / 1e6
            fom = c["om"]["fixed_per_mw"] * c["capacity_mw"] * idx / 4 / 1e6 if op else 0
            vom = c["om"]["variable_per_mwh"] * mwh * idx / 1e6
            ebitda = rev - fom - vom
            dep = dep_q if op else 0.0
            tax = max(0.0, ebitda - dep) * c["tax_rate"]
            rows.append(dict(d=d, op=op, rev=rev, opex=fom + vom, ebitda=ebitda, dep=dep, tax=tax, cfads=ebitda - tax,
                             rep=self.first_rep <= d <= self.maturity, ll=op and d <= self.maturity))
        return rows

    def construction(self, D, TPC, dsra0):
        c, dd = self.c, self.c["debt"]
        E = TPC - D
        open_debt, cum_eq, out = 0.0, 0.0, []
        for m, d in enumerate(self.months):
            land = c["capex"]["land"] if m == 0 else 0
            epc = c["capex"]["epc"] / self.cm if m >= 1 else 0
            dev = c["capex"]["development"] / self.cm if m >= 1 else 0
            fee = dd["upfront_fee"] * D if m == 0 else 0
            idc = open_debt * self.r_pre / 12 if m >= 1 else 0
            dsra = dsra0 if m == self.cm else 0
            uses = land + epc + dev + fee + idc + dsra
            eq = min(uses, max(E - cum_eq, 0))
            draw = uses - eq
            cum_eq += eq
            out.append(dict(d=d, land=land, epc=epc, dev=dev, fee=fee, idc=idc, dsra=dsra, uses=uses, eq=eq, draw=draw))
            open_debt += draw
        return out

    def size(self, ops):
        r = self.r_post / 4
        cf = [o["cfads"] for o in ops if o["rep"]]
        pv = sum(v / (1 + r) ** (k + 1) for k, v in enumerate(cf))
        return pv / self.c["debt"]["min_dscr"], pv

    def schedule(self, D, ops, pv):
        r = self.r_post / 4
        flat = pv / D
        nrep = sum(o["rep"] for o in ops)
        bal, out = 0.0, []
        for k, o in enumerate(ops):
            first_op = o["op"] and (k == 0 or not ops[k - 1]["op"])
            if first_op:
                bal += D
            opening = bal
            interest = opening * r if o["op"] else 0.0
            if o["rep"]:
                prin = min(opening, o["cfads"] / flat - interest) if self.c["debt"]["repayment"] == "sculpted" else min(opening, D / nrep)
            else:
                prin = 0.0
            bal = opening - prin
            out.append(dict(open=opening, int=interest, prin=prin, ds=interest + prin, close=bal))
        return out, flat

    def dsra_targets(self, sched):
        n = self.c["dsra"]["quarters"]
        on = self.c["dsra"]["on"]
        return [on * sum(s["ds"] for s in sched[k + 1:k + 1 + n]) if self.cod <= d < self.maturity else 0.0
                for k, (d, _s) in enumerate(zip(self.q, sched))]

    def solve(self, tariff=None):
        base = sum(self.c["capex"][k] for k in ("land", "epc", "development"))
        D, TPC, dep_base = 0.8 * base, base, base - self.c["capex"]["land"]
        for _ in range(2000):
            ops = self.ops(tariff, dep_base)
            d_dscr, pv = self.size(ops)
            sched, _flat = self.schedule(D, ops, pv)
            tgt = self.dsra_targets(sched)
            dsra0 = tgt[self.k_cod] if self.k_cod is not None else 0.0
            con = self.construction(D, TPC, dsra0)
            tpc = sum(x["uses"] for x in con)
            dep_new = sum(x["epc"] + x["dev"] + x["idc"] + x["fee"] for x in con)
            d_new = min(self.c["debt"]["gearing"] * tpc, d_dscr)
            if abs(d_new - D) < 1e-11 and abs(tpc - TPC) < 1e-11 and abs(dep_new - dep_base) < 1e-11:
                break
            D, TPC, dep_base = d_new, tpc, dep_new
        return D, TPC, con, ops, d_dscr, pv, sched, tgt

    def run(self, tariff=None):
        c = self.c
        D, TPC, con, ops, d_dscr, pv, sched, tgt = self.solve(tariff)
        r = self.r_post / 4
        _s, flat = self.schedule(D, ops, pv)
        eq_cf = [(-x["eq"], x["d"]) for x in con]
        prj_cf = [(-(x["land"] + x["epc"] + x["dev"]), x["d"]) for x in con]
        dscrs, llcrs, plcrs, short = [], [], [], 0.0
        prev_close_dsra = 0.0
        for k, (o, s) in enumerate(zip(ops, sched)):
            move = 0.0
            if o["d"] == self.cod:
                close_dsra = tgt[k]
            elif o["op"]:
                need, post = tgt[k] - prev_close_dsra, o["cfads"] - s["ds"]
                move = min(need, max(0.0, post)) if need > 0 else need   # top up only from available cash
                close_dsra = prev_close_dsra + move
            else:
                close_dsra = prev_close_dsra
            if o["op"]:
                cafd = o["cfads"] - s["ds"] - move
                short += min(0.0, cafd)
                eq_cf.append((max(0.0, cafd), o["d"]))
                prj_cf.append((o["cfads"], o["d"]))
            if o["rep"]:
                dscrs.append(o["cfads"] / s["ds"])
                pvl = sum(ops[j]["cfads"] * ops[j]["ll"] / (1 + r) ** (j - k + 1) for j in range(k, len(ops)))
                pvp = sum(ops[j]["cfads"] * ops[j]["op"] / (1 + r) ** (j - k + 1) for j in range(k, len(ops)))
                llcrs.append((pvl + prev_close_dsra) / s["open"])
                plcrs.append(pvp / s["open"])
            prev_close_dsra = close_dsra
        e_irr = xirr([a for a, _ in eq_cf], [b for _, b in eq_cf])
        p_irr = xirr([a for a, _ in prj_cf], [b for _, b in prj_cf])
        eq_npv = xnpv(c["target_equity_irr"], [a for a, _ in eq_cf], [b for _, b in eq_cf])
        return dict(D=D, TPC=TPC, idc=sum(x["idc"] for x in con), fee=sum(x["fee"] for x in con), dsra0=sum(x["dsra"] for x in con),
                    d_dscr=d_dscr, d_gear=c["debt"]["gearing"] * TPC, pv=pv, flat=flat,
                    bind="DSCR" if d_dscr < c["debt"]["gearing"] * TPC else "Gearing",
                    e_irr=e_irr, p_irr=p_irr, eq_npv=eq_npv, min_dscr=min(dscrs), avg_dscr=sum(dscrs) / len(dscrs),
                    min_llcr=min(llcrs), min_plcr=min(plcrs), end_bal=sched[-1]["close"], short=short,
                    revenue=sum(o["rev"] for o in ops), opex=sum(o["opex"] for o in ops))

    def breakeven_tariff(self, target):
        lo, hi = 20.0, 200.0
        for _ in range(60):
            mid = (lo + hi) / 2
            if self.run(mid)["e_irr"] < target:
                lo = mid
            else:
                hi = mid
        return (lo + hi) / 2


SCENARIOS = [
    ("Base case", lambda c: None),
    ("P90 yield", lambda c: c.update(energy_case="P90")),
    ("Capex +10%", lambda c: c["capex"].update(epc=c["capex"]["epc"] * 1.1, development=c["capex"]["development"] * 1.1)),
    ("O&M +10%", lambda c: c["om"].update(fixed_per_mw=c["om"]["fixed_per_mw"] * 1.1, variable_per_mwh=c["om"]["variable_per_mwh"] * 1.1)),
    ("Rate +100bp", lambda c: c["debt"].update(swap=c["debt"]["swap"] + 0.01)),
    ("Tariff –5%", lambda c: c.update(tariff=c["tariff"] * 0.95)),
]


def run_scenarios(c):
    out = []
    for name, f in SCENARIOS:
        cc = copy.deepcopy(c)
        f(cc)
        out.append((name, Engine(cc).run()))
    return out


# ------------------------------------------------------------------ workbook
def build(c, path):
    c = defaults(c)
    eng = Engine(c)
    if c["dsra"]["on"] and eng.k_cod is None:
        raise SystemExit("DSRA funding at COD needs COD on a quarter end (construction months so that COD = Mar/Jun/Sep/Dec)")
    res = eng.run()
    D, TPC = res["D"], res["TPC"]
    be = eng.breakeven_tariff(c["target_equity_irr"])
    scen = run_scenarios(c)

    wb = Workbook()
    workbook_defaults(wb)
    cov = wb.active
    cov.title = "Cover"
    names = ("Inputs", "Timing", "Construction", "Operations", "Debt", "Waterfall", "Annual", "Returns", "Scenarios",
             "Dashboard", "Checks")
    inp, tim, con, opn, dbt, wf, ann, ret, scn, dash, chk = (wb.create_sheet(s) for s in names)
    ROWS = {s: {} for s in ("CON", "OPN", "DBT", "WF", "TIM", "ANN", "RET", "DASH", "SCN")}
    case = c.get("case_short", c["project"])

    nq, nm = len(eng.q), c["construction_months"] + 1
    years = list(range(eng.start.year, eng.ppa_end.year + 1))
    ny = len(years)
    Q0 = FIRST_PERIOD
    qc = lambda k: col(Q0 + k)  # noqa: E731
    lastq = qc(nq - 1)
    QL = Q0 + nq - 1
    for ws in (cov, inp, dash, chk):
        setup(ws, label_w=52, widths={"C": 11})
    setup(scn, label_w=52, period_w=12, widths={"D": 12, "E": 12})
    for ws in (tim, opn, dbt, wf):
        setup(ws, label_w=46, period_w=9, n_periods=nq, widths={"E": 2})
    setup(con, label_w=46, period_w=10, n_periods=nm, widths={"E": 2})
    setup(ann, label_w=46, period_w=10, n_periods=ny, widths={"E": 2})
    setup(ret, label_w=14, widths={"C": 2, "D": 13, "E": 13})

    # ---------------- Inputs
    r = sheet_header(inp, case, "Inputs · yellow = assumption (typed) · solver values are pasted here too", 8)
    I = {}

    def sec_i(text):
        nonlocal r
        section(inp, r, text, last_col=8)
        r += 1

    def inp_row(key, lab, unit, val, fmt=NUM1, formula=False, note_=None):
        nonlocal r
        label(inp, r, lab, unit)
        if formula:
            put(inp, f"D{r}", val, fmt)
        else:
            input_cell(inp, f"D{r}", val, fmt)
        if note_:
            comment(inp, f"E{r}", note_)
        I[key] = f"Inputs!$D${r}"
        r += 1

    sec_i("Timing")
    inp_row("start", "Financial close / construction start", "date", dt.datetime(*c["start"]), DATE)
    inp_row("cm", "Construction period", "months", c["construction_months"], "0")
    inp_row("ppa", "PPA term from completion", "years", c["ppa_years"], "0")
    inp_row("frep", "First repayment after completion", "months", c["first_repayment_months_after_cod"], "0")
    inp_row("tenor", "Loan tenor from construction start", "years", c["debt"]["tenor_years"], "0")
    r += 1
    sec_i("Generation and revenue")
    inp_row("cap", "Capacity (AC)", "MW", c["capacity_mw"], "#,##0")
    inp_row("hours", "Hours per year (leap years ignored)", "hours", c["hours_per_year"], "#,##0")
    inp_row("p50", "Capacity utilisation factor – P50", "%", c["cuf"]["P50"], PCT)
    inp_row("p90", "Capacity utilisation factor – P90", "%", c["cuf"]["P90"], PCT)
    inp_row("case", "Energy case (1 = P50, 2 = P90)", "switch", 1 if c["energy_case"] == "P50" else 2, "0")
    inp_row("deg", "Annual degradation (from the 2nd operating year)", "%", c["degradation"], '0.00%')
    for i, q in enumerate(("Q1 (Jan–Mar)", "Q2 (Apr–Jun)", "Q3 (Jul–Sep)", "Q4 (Oct–Dec)")):
        inp_row(f"s{i + 1}", f"Seasonality – {q}", "% of year", c["seasonality"][i], PCT)
    inp_row("tariff", "PPA tariff (fixed, no escalation)", "USD/MWh", c["tariff"], '#,##0.00')
    r += 1
    sec_i("Operating costs (real, escalated from the base year)")
    inp_row("fom", "Fixed O&M", "USD/MW/yr", c["om"]["fixed_per_mw"], "#,##0")
    inp_row("vom", "Variable O&M", "USD/MWh", c["om"]["variable_per_mwh"], "#,##0.00")
    inp_row("esc", "O&M escalation", "% p.a.", c["om"]["escalation"], PCT)
    inp_row("base", "O&M price base year", "year", c["om"]["base_year"], "0")
    inp_row("usd", "USD per USD m (unit conversion)", "USD", 1000000, "#,##0")
    r += 1
    sec_i("Capital costs")
    inp_row("land", "Land (paid at financial close)", U, c["capex"]["land"])
    inp_row("epc", "EPC contract (lump-sum turnkey, spread evenly)", U, c["capex"]["epc"])
    inp_row("dev", "Licences and development costs (spread evenly)", U, c["capex"]["development"])
    inp_row("life", "Depreciation life (straight line, excludes land)", "years", c["depreciation_years"], "0")
    inp_row("tax", "Tax rate on EBITDA – depreciation", "%", c["tax_rate"], PCT, note_="the case gives no tax")
    r += 1
    sec_i("Senior debt")
    dd = c["debt"]
    inp_row("gear", "Maximum gearing (% of total project cost)", "%", dd["gearing"], PCT)
    inp_row("mindscr", "Minimum DSCR for sizing", "x", dd["min_dscr"], MULT)
    inp_row("fee", "Upfront fee (% of debt, paid at close)", "%", dd["upfront_fee"], '0.00%')
    inp_row("swap", "SOFR swap rate (100% hedged)", "%", dd["swap"], '0.00%')
    inp_row("m_pre", "Margin pre-completion", "%", dd["margin_pre"], '0.00%')
    inp_row("m_post", "Margin post-completion", "%", dd["margin_post"], '0.00%')
    inp_row("rep", "Repayment (1 = sculpted to a flat DSCR, 2 = straight line)", "switch", 1 if dd["repayment"] == "sculpted" else 2, "0")
    r += 1
    sec_i("Debt service reserve account (DSRA)")
    inp_row("dsra_on", "DSRA required (1 = yes, 0 = no)", "switch", c["dsra"]["on"], "0")
    inp_row("dsra_q", "DSRA cover (quarters of senior debt service)", "quarters", c["dsra"]["quarters"], "0",
            note_="funded at COD as part of project cost; no interest earned on the balance")
    r += 1
    sec_i("Equity")
    inp_row("target", "Client minimum equity IRR", "%", c["target_equity_irr"], PCT)
    r += 1
    sec_i("Solver values (pasted – do not type over; run build_pf.py --resolve after changing inputs)")
    inp_row("Dp", "Debt amount (pasted)", U, round(D, 10), '#,##0.000', note_="solved by build_pf.py; Checks proves it equals the live sizing")
    inp_row("TPCp", "Total project cost (pasted)", U, round(TPC, 10), '#,##0.000', note_="includes IDC, fee and DSRA funding")
    r += 1
    sec_i("Derived (formulas)")
    for k, lab, f, fmt in [
            ("cod", "Construction completion (COD)", f"=EOMONTH({I['start']},{I['cm']}-1)", DATE),
            ("ppa_end", "PPA end", "=EOMONTH(@COD@,12*" + I["ppa"] + ")", DATE),
            ("first_rep", "First repayment date", "=EOMONTH(@COD@," + I["frep"] + ")", DATE),
            ("maturity", "Final maturity", f"=EOMONTH({I['start']},12*{I['tenor']}-1)", DATE),
            ("r_pre", "All-in rate pre-completion", f"={I['swap']}+{I['m_pre']}", '0.00%'),
            ("r_post", "All-in rate post-completion", f"={I['swap']}+{I['m_post']}", '0.00%'),
            ("cuf", "Capacity factor used", "=IF(ISNUMBER(@SCN_CUF@),@SCN_CUF@,CHOOSE(" + I["case"] + f",{I['p50']},{I['p90']}))", PCT),
            ("tariff_u", "Tariff used", "=IF(ISNUMBER(@SCN_TAR@),@SCN_TAR@," + I["tariff"] + ")", '#,##0.00')]:
        inp_row(k, lab, "", f, fmt, formula=True)
    comment(inp, f"E{I['cuf'].split('$')[-1]}", "the Scenarios sheet can override this for its data table")
    end_marker(inp, r + 1, 8)
    for row in inp.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "@COD@" in cell.value:
                cell.value = cell.value.replace("@COD@", I["cod"])

    # ---------------- helpers for time-based sheets
    def band_q(ws, desc):
        rows = [("Quarter ending", [f"=Timing!{qc(k)}$@TIM:end@" for k in range(nq)], DATE, False),
                ("Calendar year", [f"=Timing!{qc(k)}$@TIM:year@" for k in range(nq)], "0", False),
                ("Phase", [f"=Timing!{qc(k)}$@TIM:phase@" for k in range(nq)], "@", True)]
        return sheet_header(ws, case, desc, QL, rows, const_labels=("Total", ""))

    def writer(ws, key_map, first, n, colf):
        state = {"r": first}

        def row(key, lab, unit, fn, fmt=M1, total=False, bold=False, italic=False, border=None):
            rr = state["r"]
            label(ws, rr, lab, unit, bold=bold, italic=italic, indent=1 if italic else 0)
            key_map[key] = rr
            for k in range(n):
                cell = put(ws, f"{colf(k)}{rr}", fn(k, colf(k)), fmt, bold=bold, italic=italic)
                if border:
                    cell.border = border
            if total:
                put(ws, f"D{rr}", f"=SUM({colf(0)}{rr}:{colf(n - 1)}{rr})", fmt, bold=True)
            state["r"] += 1

        def sec(text, last):
            section(ws, state["r"], text, last_col=last)
            state["r"] += 1

        def gap(k=1):
            state["r"] += k

        return row, sec, gap, state

    T = ROWS["TIM"]
    prev = lambda k, L, rr, first="0": first if k == 0 else f"{qc(k - 1)}{rr}"  # noqa: E731

    # ---------------- Timing (quarterly)
    r0 = band_q(tim, "Timing: quarterly timeline and flags (1 = true) used by every quarterly sheet")
    qrow, qsec, qgap, st = writer(tim, T, r0, nq, qc)
    qsec("Timeline", QL)
    qrow("n", "Period counter", "#", lambda k, L: 1 if k == 0 else f"={qc(k - 1)}{st['r']}+1", "0")
    qrow("start", "Quarter start", "date", lambda k, L: f"=EOMONTH({I['start']},3*({L}{T['n']}-1)-1)+1", DATE)
    qrow("end", "Quarter end", "date", lambda k, L: f"=EOMONTH({I['start']},3*{L}{T['n']}-1)", DATE)
    qrow("year", "Calendar year", "year", lambda k, L: f"=YEAR({L}{T['end']})", "0")
    qrow("qy", "Quarter of the year", "#", lambda k, L: f"=ROUNDUP(MONTH({L}{T['end']})/3,0)", "0")
    qgap()
    qsec("Flags", QL)
    qrow("f_con", "Construction", "(1,0)", lambda k, L: f"=IF({L}{T['end']}<={I['cod']},1,0)", FLAG)
    qrow("f_cod", "COD quarter", "(1,0)", lambda k, L: f"=IF({L}{T['end']}={I['cod']},1,0)", FLAG)
    qrow("f_op", "Operations (after COD, to PPA end)", "(1,0)",
         lambda k, L: f"=IF(AND({L}{T['end']}>{I['cod']},{L}{T['end']}<={I['ppa_end']}),1,0)", FLAG, total=True)
    qrow("f_op1", "First operating quarter", "(1,0)",
         lambda k, L: f"={L}{T['f_op']}" if k == 0 else f"=IF(AND({L}{T['f_op']}=1,{qc(k - 1)}{T['f_op']}=0),1,0)", FLAG)
    qrow("f_rep", "Repayment period", "(1,0)",
         lambda k, L: f"=IF(AND({L}{T['end']}>={I['first_rep']},{L}{T['end']}<={I['maturity']}),1,0)", FLAG, total=True)
    qrow("f_ll", "Loan life (operations to final maturity)", "(1,0)", lambda k, L: f"={L}{T['f_op']}*IF({L}{T['end']}<={I['maturity']},1,0)", FLAG)
    qrow("f_dsra", "DSRA held (COD to the quarter before maturity)", "(1,0)",
         lambda k, L: f"=IF(AND({L}{T['end']}>={I['cod']},{L}{T['end']}<{I['maturity']}),1,0)*{I['dsra_on']}", FLAG)
    qrow("oy", "Operating year", "#", lambda k, L: f"=IF({L}{T['f_op']}=1,{L}{T['year']}-YEAR({I['cod']}),0)", "0")
    qrow("rep_n", "Repayment period counter", "#", lambda k, L: f"=SUM(${qc(0)}{T['f_rep']}:{L}{T['f_rep']})*{L}{T['f_rep']}", "0")
    qrow("idx", "O&M escalation index", "x", lambda k, L: f"=(1+{I['esc']})^({L}{T['year']}-{I['base']})", "0.0000")
    qrow("phase", "Phase", "", lambda k, L: f'=IF({L}{T["f_con"]}=1,"Constr.",IF({L}{T["f_op"]}=1,"Ops","Post-PPA"))', "@")
    qgap()
    end_marker(tim, st["r"], QL)
    tq = lambda key, L: f"Timing!{L}${T[key]}"  # noqa: E731

    # ---------------- Operations (quarterly)
    OP = ROWS["OPN"]
    r0 = band_q(opn, f"Operations: generation, revenue, O&M, EBITDA, tax, CFADS · {U} unless stated")
    orow, osec, ogap, st = writer(opn, OP, r0, nq, qc)
    osec("Generation", QL)
    orow("mwh", "Energy generated", "MWh",
         lambda k, L: f"={I['cap']}*{I['hours']}*{I['cuf']}*CHOOSE({tq('qy', L)},{I['s1']},{I['s2']},{I['s3']},{I['s4']})"
                      f"*(1-{I['deg']})^MAX({tq('oy', L)}-1,0)*{tq('f_op', L)}", "#,##0", total=True)
    ogap()
    osec("Revenue and operating costs", QL)
    orow("rev", "Revenue", U, lambda k, L: f"={L}{OP['mwh']}*{I['tariff_u']}/{I['usd']}", total=True, bold=True)
    orow("fom", "Fixed O&M", U, lambda k, L: f"={I['fom']}*{I['cap']}*{tq('idx', L)}/4/{I['usd']}*{tq('f_op', L)}", total=True)
    orow("vom", "Variable O&M", U, lambda k, L: f"={I['vom']}*{L}{OP['mwh']}*{tq('idx', L)}/{I['usd']}", total=True)
    orow("opex", "Total operating costs", U, lambda k, L: f"={L}{OP['fom']}+{L}{OP['vom']}", total=True, border=TOTAL)
    orow("ebitda", "EBITDA", U, lambda k, L: f"={L}{OP['rev']}-{L}{OP['opex']}", total=True, bold=True)
    orow("margin", "EBITDA margin", "%", lambda k, L: f"=IF({L}{OP['rev']}>0,{L}{OP['ebitda']}/{L}{OP['rev']},0)", PCT, italic=True)
    orow("dep", "Depreciation", U, lambda k, L: f"=Construction!$D$@CON:depbase@/({I['life']}*4)*{tq('f_op', L)}", total=True)
    orow("tax", "Tax paid (on EBITDA – depreciation)", U, lambda k, L: f"=MAX(0,{L}{OP['ebitda']}-{L}{OP['dep']})*{I['tax']}", total=True)
    orow("cfads", "CFADS (EBITDA – tax)", U, lambda k, L: f"={L}{OP['ebitda']}-{L}{OP['tax']}", total=True, bold=True, border=TOTAL)
    comment(opn, f"{col(QL + 1)}{OP['cfads']}", "cash flow available for debt service")
    ogap()
    end_marker(opn, st["r"], QL)

    # ---------------- Debt (quarterly)
    DB = ROWS["DBT"]
    r0 = band_q(dbt, f"Senior debt: sizing, sculpted repayment, DSCR / LLCR / PLCR · {U} unless stated")
    drow, dsec, dgap, st = writer(dbt, DB, r0, nq, qc)
    dsec("Debt sizing (lower of gearing and DSCR capacity)", QL)
    sz = {}
    for k_, lab, u_, f, fmt in [
            ("cf_rep", "CFADS in the repayment period", U, None, M1),
            ("pv", "PV of CFADS over the repayment period", U, None, M1),
            ("d_dscr", "Debt capacity at the minimum DSCR", U, None, M1),
            ("d_gear", "Debt capacity at maximum gearing", U, None, M1),
            ("d", "Maximum debt = lower of the two", U, None, M1),
            ("bind", "Binding constraint", "", None, "@"),
            ("flat", "Flat DSCR for sculpting (PV of CFADS / debt)", "x", None, MULT),
            ("nrep", "Number of repayment quarters", "#", None, "0")]:
        if k_ == "cf_rep":
            drow("cf_rep", lab, u_, lambda k, L: f"=Operations!{L}{OP['cfads']}*{tq('f_rep', L)}", total=True)
            continue
        label(dbt, st["r"], lab, u_, bold=(k_ == "d"))
        sz[k_] = st["r"]
        DB[k_] = st["r"]
        st["r"] += 1
    rpq = f"{I['r_post']}/4"
    put(dbt, f"D{sz['pv']}", f"=SUMPRODUCT({qc(0)}{DB['cf_rep']}:{lastq}{DB['cf_rep']}/(1+{rpq})^(Timing!{qc(0)}{T['rep_n']}:{lastq}{T['rep_n']}))", M1)
    put(dbt, f"D{sz['d_dscr']}", f"=D{sz['pv']}/{I['mindscr']}", M1)
    put(dbt, f"D{sz['d_gear']}", f"={I['gear']}*Construction!$D$@CON:uses@", M1)
    output_box(dbt, f"D{sz['d']}", f"=MIN(D{sz['d_dscr']},D{sz['d_gear']})", M1)
    put(dbt, f"D{sz['bind']}", f'=IF(D{sz["d_dscr"]}<D{sz["d_gear"]},"DSCR","Gearing")', "@", align="right")
    put(dbt, f"D{sz['flat']}", f"=D{sz['pv']}/{I['Dp']}", MULT)
    put(dbt, f"D{sz['nrep']}", f"=Timing!D{T['f_rep']}", "0")
    comment(dbt, f"E{sz['d']}", "solved and pasted on Inputs (no circularity); Checks proves pasted = this")
    comment(dbt, f"E{sz['pv']}", "discounted at the post-completion rate to the start of the first repayment quarter")
    comment(dbt, f"E{sz['d_dscr']}", "PV of CFADS / minimum DSCR")
    comment(dbt, f"E{sz['d_gear']}", "maximum gearing × total project cost")
    dgap()
    dsec("Senior debt schedule", QL)
    drow("cfads", "CFADS", U, lambda k, L: f"=Operations!{L}{OP['cfads']}", total=True)
    drow("xfer", "Debt drawn in construction (transferred at COD)", U, lambda k, L: f"={tq('f_op1', L)}*{I['Dp']}", total=True)
    drow("open", "Opening balance", U, lambda k, L: f"={L}@DBT:xfer@" if k == 0 else f"={qc(k - 1)}@DBT:close@+{L}@DBT:xfer@")
    drow("int", "Interest (post-completion rate, on the opening balance)", U, lambda k, L: f"={L}{DB['open']}*{rpq}*{tq('f_op', L)}", total=True)
    drow("tds", "Target debt service (CFADS / flat DSCR)", U, lambda k, L: f"={tq('f_rep', L)}*{L}{DB['cfads']}/$D${sz['flat']}", total=True)
    drow("prin", "Principal repaid", U,
         lambda k, L: f"={tq('f_rep', L)}*MIN({L}{DB['open']},IF({I['rep']}=1,{L}{DB['tds']}-{L}{DB['int']},{I['Dp']}/$D${sz['nrep']}))", total=True)
    drow("close", "Closing balance", U, lambda k, L: f"={L}{DB['open']}-{L}{DB['prin']}", bold=True, border=TOTAL)
    drow("ds", "Senior debt service (interest + principal)", U, lambda k, L: f"={L}{DB['int']}+{L}{DB['prin']}", total=True)
    dgap()
    dsec("Cover ratios", QL)
    drow("dscr", "DSCR (CFADS / debt service)", "x", lambda k, L: f'=IF({tq("f_rep", L)}=1,{L}{DB["cfads"]}/{L}{DB["ds"]},"")', MULT, bold=True)
    drow("pvl", "PV of CFADS to final maturity (loan life)", U,
         lambda k, L: f"=({L}{DB['cfads']}*{tq('f_ll', L)}" + (f"+{qc(k + 1)}{st['r']}" if k < nq - 1 else "") + f")/(1+{rpq})")
    drow("pvp", "PV of CFADS to PPA end (project life)", U,
         lambda k, L: f"=({L}{DB['cfads']}*{tq('f_op', L)}" + (f"+{qc(k + 1)}{st['r']}" if k < nq - 1 else "") + f")/(1+{rpq})")
    drow("llcr", "LLCR ((PV to maturity + DSRA) / debt outstanding)", "x",
         lambda k, L: f'=IF({tq("f_rep", L)}=1,({L}{DB["pvl"]}+Waterfall!{L}@WF:d_open@)/{L}{DB["open"]},"")', MULT, bold=True)
    drow("plcr", "PLCR (PV to PPA end / debt outstanding)", "x",
         lambda k, L: f'=IF({tq("f_rep", L)}=1,{L}{DB["pvp"]}/{L}{DB["open"]},"")', MULT, bold=True)
    comment(dbt, f"{col(QL + 1)}{DB['pvl']}", "discounted back one quarter at a time (no OFFSET, no circularity)")
    dgap()
    end_marker(dbt, st["r"], QL)
    # the 'open' row references close/xfer rows by token
    for row in dbt.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("@DBT:close@", str(DB["close"])).replace("@DBT:xfer@", str(DB["xfer"]))

    # ---------------- Waterfall (quarterly)
    W = ROWS["WF"]
    r0 = band_q(wf, f"Cash flow waterfall and DSRA · {U}")
    wrow, wsec, wgap, st = writer(wf, W, r0, nq, qc)
    wsec("Cash flow waterfall (order of priority)", QL)
    wrow("cfads", "CFADS", U, lambda k, L: f"=Operations!{L}{OP['cfads']}", total=True, bold=True)
    wrow("int", "Less: senior interest", U, lambda k, L: f"=-Debt!{L}{DB['int']}", total=True)
    wrow("prin", "Less: senior principal", U, lambda k, L: f"=-Debt!{L}{DB['prin']}", total=True)
    wrow("post_ds", "Cash flow after senior debt service", U, lambda k, L: f"=SUM({L}{W['cfads']}:{L}{W['prin']})", total=True, bold=True, border=TOTAL)
    wrow("dsra_mv", "Less: DSRA top-up / add: DSRA release", U, lambda k, L: "=-{L}@WF:d_top@".replace("{L}", L), total=True)
    wrow("cafd", "Cash available for distribution", U, lambda k, L: f"={L}{W['post_ds']}+{L}{W['dsra_mv']}", total=True, bold=True, border=TOTAL)
    wrow("div", "Dividends to equity", U, lambda k, L: f"=MAX(0,{L}{W['cafd']})*{tq('f_op', L)}", total=True)
    wrow("short", "Cash shortfall (must be zero)", U, lambda k, L: f"=MIN(0,{L}{W['cafd']})*{tq('f_op', L)}", total=True, italic=True)
    wgap()
    wsec("Debt service reserve account (DSRA)", QL)
    wrow("d_tgt", "Target balance (next N quarters of senior debt service)", U,
         lambda k, L: f"={tq('f_dsra', L)}*SUMPRODUCT((Timing!${qc(0)}${T['n']}:${lastq}${T['n']}>{tq('n', L)})"
                      f"*(Timing!${qc(0)}${T['n']}:${lastq}${T['n']}<={tq('n', L)}+{I['dsra_q']})*Debt!${qc(0)}${DB['ds']}:${lastq}${DB['ds']})")
    wrow("d_open", "Opening balance", U, lambda k, L: "=0" if k == 0 else f"={qc(k - 1)}@WF:d_close@")
    wrow("d_fund", "Funded at COD (from construction funding)", U, lambda k, L: f"={tq('f_cod', L)}*{L}{W['d_tgt']}", total=True)
    wrow("d_top", "Top-up (only from cash after debt service) / (release)", U,
         lambda k, L: f"={tq('f_op', L)}*IF({L}{W['d_tgt']}>{L}{W['d_open']},MIN({L}{W['d_tgt']}-{L}{W['d_open']},MAX(0,{L}{W['post_ds']})),"
                      f"{L}{W['d_tgt']}-{L}{W['d_open']})", total=True)
    wrow("d_close", "Closing balance", U, lambda k, L: f"={L}{W['d_open']}+{L}{W['d_fund']}+{L}{W['d_top']}", bold=True, border=TOTAL)
    wrow("d_gap", "Below target (refilled from later cash)", U,
         lambda k, L: f"=MAX(0,{L}{W['d_tgt']}-{L}{W['d_close']})", italic=True)
    comment(wf, f"{col(QL + 1)}{W['d_tgt']}", "released in full at final maturity")
    wgap()
    end_marker(wf, st["r"], QL)
    for row in wf.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("@WF:d_top@", str(W["d_top"])).replace("@WF:d_close@", str(W["d_close"]))
    for row in dbt.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("@WF:d_open@", str(W["d_open"]))

    # ---------------- Construction (monthly)
    CR = ROWS["CON"]
    mc = lambda m: col(Q0 + m)  # noqa: E731
    lastm = mc(nm - 1)
    band_m = [("Month ending", [f"={mc(m)}$@CON:date@" for m in range(nm)], DATE, False),
              ("Month", [f"={mc(m)}$@CON:n@" for m in range(nm)], '"M"0', True)]
    r0 = sheet_header(con, case, f"Construction (monthly): uses, equity-first funding, IDC, debt drawdown · {U}", Q0 + nm - 1,
                      band_m, const_labels=("Total", ""))
    crow, csec, cgap, st = writer(con, CR, r0, nm, mc)
    CL = Q0 + nm - 1
    csec("Timeline", CL)
    crow("n", "Month counter (M0 = financial close)", "#", lambda m, L: 0 if m == 0 else f"={mc(m - 1)}{st['r']}+1", "0")
    crow("date", "Month end", "date", lambda m, L: f"=IF({L}{CR['n']}=0,{I['start']},EOMONTH({I['start']},{L}{CR['n']}-1))", DATE)
    crow("year", "Calendar year", "year", lambda m, L: f"=YEAR({L}{CR['date']})", "0")
    crow("moy", "Month of the year", "#", lambda m, L: f"=MONTH({L}{CR['date']})", "0")
    crow("f_cod", "COD month", "(1,0)", lambda m, L: f"=IF({L}{CR['date']}={I['cod']},1,0)", FLAG)
    cgap()
    csec("Uses of funds", CL)
    crow("land", "Land", U, lambda m, L: f"=IF({L}{CR['n']}=0,{I['land']},0)", total=True)
    crow("epc", "EPC", U, lambda m, L: f"=IF({L}{CR['n']}>=1,{I['epc']}/{I['cm']},0)", total=True)
    crow("dev", "Licences and development", U, lambda m, L: f"=IF({L}{CR['n']}>=1,{I['dev']}/{I['cm']},0)", total=True)
    crow("fee", "Upfront financing fee", U, lambda m, L: f"=IF({L}{CR['n']}=0,{I['fee']}*{I['Dp']},0)", total=True)
    crow("idc", "Interest during construction (IDC)", U,
         lambda m, L: "=0" if m == 0 else f"={mc(m - 1)}@CON:close@*{I['r_pre']}/12", total=True)
    crow("dsra", "DSRA initial funding (COD month only)", U,
         lambda m, L: "=0" if m < nm - 1 else f"={L}{CR['f_cod']}*SUMIFS(Waterfall!${qc(0)}${W['d_fund']}:${lastq}${W['d_fund']},Timing!${qc(0)}${T['end']}:${lastq}${T['end']},{L}{CR['date']})",
         total=True)
    crow("uses", "Total uses (= total project cost)", U, lambda m, L: f"=SUM({L}{CR['land']}:{L}{CR['dsra']})", total=True, bold=True, border=TOTAL)
    cgap()
    csec("Sources of funds (equity first, then debt)", CL)
    crow("eq", "Equity drawn", U,
         lambda m, L: f"=MIN({L}{CR['uses']},MAX(0,({I['TPCp']}-{I['Dp']})" + ("" if m == 0 else f"-SUM(${mc(0)}{st['r']}:{mc(m - 1)}{st['r']})") + "))",
         total=True)
    crow("draw", "Debt drawn", U, lambda m, L: f"={L}{CR['uses']}-{L}{CR['eq']}", total=True)
    crow("src", "Total sources", U, lambda m, L: f"={L}{CR['eq']}+{L}{CR['draw']}", total=True, bold=True, border=TOTAL)
    cgap()
    csec("Senior debt during construction", CL)
    crow("open", "Opening balance", U, lambda m, L: "=0" if m == 0 else f"={mc(m - 1)}@CON:close@")
    crow("dd", "Drawdown", U, lambda m, L: f"={L}{CR['draw']}", total=True)
    crow("close", "Closing balance", U, lambda m, L: f"={L}{CR['open']}+{L}{CR['dd']}", bold=True, border=TOTAL)
    cgap()
    label(con, st["r"], "Depreciable cost (EPC + development + IDC + fee)", U, bold=True)
    CR["depbase"] = st["r"]
    output_box(con, f"D{st['r']}", f"=D{CR['epc']}+D{CR['dev']}+D{CR['idc']}+D{CR['fee']}", M1)
    comment(con, f"E{st['r']}", "land is not depreciated; the DSRA is cash, not plant")
    st["r"] += 2
    end_marker(con, st["r"], CL)
    for row in con.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = (cell.value.replace("@CON:close@", str(CR["close"])).replace("@CON:date@", str(CR["date"]))
                              .replace("@CON:n@", str(CR["n"])))

    # ---------------- Annual 3 statements
    A = ROWS["ANN"]
    ac = lambda i: col(Q0 + i)  # noqa: E731
    lastA = ac(ny - 1)
    AL = Q0 + ny - 1
    r0 = sheet_header(ann, case, f"Annual financial statements (calendar years) and debt metrics · {U}", AL,
                      [("Financial year", [f"={ac(i)}$@ANN:yr@" for i in range(ny)], "0", False)], const_labels=("Total", ""))
    arow, asec, agap, st = writer(ann, A, r0, ny, ac)
    asec("Timeline", AL)
    arow("yr", "Calendar year", "year", lambda i, L: f"=YEAR({I['start']})" if i == 0 else f"={ac(i - 1)}{st['r']}+1", "0")
    agap()
    sq = lambda sheet, rr: (lambda i, L: f"=SUMIF(Timing!${qc(0)}${T['year']}:${lastq}${T['year']},{L}${A['yr']},{sheet}!${qc(0)}${rr}:${lastq}${rr})")  # noqa: E731
    sc = lambda rr: (lambda i, L: f"=SUMIF(Construction!${mc(0)}${CR['year']}:${lastm}${CR['year']},{L}${A['yr']},Construction!${mc(0)}${rr}:${lastm}${rr})")  # noqa: E731
    neg = lambda f: (lambda i, L: f(i, L).replace("=SUMIF", "=-SUMIF", 1))  # noqa: E731
    ye_q = lambda sheet, rr: (lambda i, L: f"=SUMIFS({sheet}!${qc(0)}${rr}:${lastq}${rr},Timing!${qc(0)}${T['year']}:${lastq}${T['year']},{L}${A['yr']},Timing!${qc(0)}${T['qy']}:${lastq}${T['qy']},4)")  # noqa: E731
    ye_c = lambda rr: (f"SUMIFS(Construction!${mc(0)}${rr}:${lastm}${rr},Construction!${mc(0)}${CR['year']}:${lastm}${CR['year']},{{L}}${A['yr']},"  # noqa: E731
                       f"Construction!${mc(0)}${CR['moy']}:${lastm}${CR['moy']},12)")
    asec("Income statement", AL)
    arow("rev", "Revenue", U, sq("Operations", OP["rev"]), total=True)
    arow("opex", "Operating costs", U, neg(sq("Operations", OP["opex"])), total=True)
    arow("ebitda", "EBITDA", U, lambda i, L: f"={L}{A['rev']}+{L}{A['opex']}", total=True, bold=True, border=TOTAL)
    arow("dep", "Depreciation", U, neg(sq("Operations", OP["dep"])), total=True)
    arow("int", "Interest expense", U, neg(sq("Debt", DB["int"])), total=True)
    arow("pbt", "Profit before tax", U, lambda i, L: f"={L}{A['ebitda']}+{L}{A['dep']}+{L}{A['int']}", total=True, bold=True, border=TOTAL)
    arow("tax", "Tax", U, neg(sq("Operations", OP["tax"])), total=True)
    arow("npat", "Net profit", U, lambda i, L: f"={L}{A['pbt']}+{L}{A['tax']}", total=True, bold=True, border=TOTAL)
    agap()
    asec("Cash flow statement", AL)
    arow("cfo_np", "Net profit", U, lambda i, L: f"={L}{A['npat']}", total=True)
    arow("cfo_dep", "Add back depreciation", U, lambda i, L: f"=-{L}{A['dep']}", total=True)
    arow("cfo", "Cash flow from operations", U, lambda i, L: f"={L}{A['cfo_np']}+{L}{A['cfo_dep']}", total=True, bold=True, border=TOTAL)
    arow("capex", "Capex (land, EPC, development)", U,
         lambda i, L: "=-(" + "+".join(sc(CR[x])(i, L)[1:] for x in ("land", "epc", "dev")) + ")", total=True)
    arow("capfin", "Capitalised IDC and fees", U, lambda i, L: "=-(" + "+".join(sc(CR[x])(i, L)[1:] for x in ("idc", "fee")) + ")", total=True)
    arow("dsra", "(Increase) / decrease in DSRA", U,
         lambda i, L: "=-(" + sq("Waterfall", W["d_fund"])(i, L)[1:] + "+" + sq("Waterfall", W["d_top"])(i, L)[1:] + ")", total=True)
    arow("cfi", "Cash flow from investing", U, lambda i, L: f"=SUM({L}{A['capex']}:{L}{A['dsra']})", total=True, bold=True, border=TOTAL)
    arow("draw", "Debt drawn", U, sc(CR["draw"]), total=True)
    arow("prin", "Debt repaid", U, neg(sq("Debt", DB["prin"])), total=True)
    arow("eq", "Equity injected", U, sc(CR["eq"]), total=True)
    arow("div", "Dividends paid", U, neg(sq("Waterfall", W["div"])), total=True)
    arow("cff", "Cash flow from financing", U, lambda i, L: f"=SUM({L}{A['draw']}:{L}{A['div']})", total=True, bold=True, border=TOTAL)
    arow("net", "Net change in cash", U, lambda i, L: f"={L}{A['cfo']}+{L}{A['cfi']}+{L}{A['cff']}", total=True, bold=True)
    arow("cash", "Closing cash", U, lambda i, L: f"={L}{A['net']}" if i == 0 else f"={ac(i - 1)}{st['r']}+{L}{A['net']}")
    agap()
    asec("Balance sheet (31 December)", AL)
    arow("land", "Land", U, lambda i, L: f"=SUM(${ac(0)}$@ANN:landin@:{L}$@ANN:landin@)")
    arow("ppe", "Plant (cost less accumulated depreciation)", U,
         lambda i, L: f"=-SUM(${ac(0)}{A['capex']}:{L}{A['capex']})-SUM(${ac(0)}{A['capfin']}:{L}{A['capfin']})-{L}{A['land']}+SUM(${ac(0)}{A['dep']}:{L}{A['dep']})")
    arow("b_dsra", "DSRA (restricted cash)", U, ye_q("Waterfall", W["d_close"]))
    arow("b_cash", "Cash", U, lambda i, L: f"={L}{A['cash']}")
    arow("assets", "Total assets", U, lambda i, L: f"=SUM({L}{A['land']}:{L}{A['b_cash']})", bold=True, border=TOTAL)
    arow("debt", "Senior debt", U, lambda i, L: ye_q("Debt", DB["close"])(i, L) + "+" + ye_c(CR["close"]).replace("{L}", L))
    arow("sc", "Share capital", U, lambda i, L: f"=SUM(${ac(0)}{A['eq']}:{L}{A['eq']})")
    arow("re", "Retained earnings (net profit less dividends)", U, lambda i, L: f"=SUM(${ac(0)}{A['npat']}:{L}{A['npat']})+SUM(${ac(0)}{A['div']}:{L}{A['div']})")
    arow("le", "Total liabilities and equity", U, lambda i, L: f"=SUM({L}{A['debt']}:{L}{A['re']})", bold=True, border=TOTAL)
    arow("bal", "Balance check (assets – liabilities – equity)", U, lambda i, L: f"={L}{A['assets']}-{L}{A['le']}", CHECK, italic=True)
    agap()
    asec("Memo: land paid and annual debt metrics", AL)
    arow("landin", "Land paid", U, sc(CR["land"]), total=True)
    arow("m_cfads", "CFADS", U, sq("Operations", OP["cfads"]), total=True)
    arow("m_ds", "Senior debt service", U, sq("Debt", DB["ds"]), total=True)
    arow("m_dscr", "Annual DSCR", "x", lambda i, L: f'=IF({L}{A["m_ds"]}>0,{L}{A["m_cfads"]}/{L}{A["m_ds"]},"")', MULT)
    agap()
    end_marker(ann, st["r"], AL)
    for row in ann.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("@ANN:yr@", str(A["yr"])).replace("@ANN:landin@", str(A["landin"]))

    # ---------------- Returns (dated flows for XIRR)
    RT = ROWS["RET"]
    r = sheet_header(ret, case, "Returns: dated cash flows (construction months, then operating quarters) for XIRR / XNPV", 6,
                     units_label=False)
    section(ret, r, "Results", last_col=6)
    r += 1
    res_row = r
    r += 5
    section(ret, r, "Dated cash flows", last_col=6)
    r += 1
    header_row(ret, r, ["Date"], 2)
    header_row(ret, r, [f"Project CF ({U})", f"Equity CF ({U})"], 4)
    ret.column_dimensions["B"].width = 14
    r += 1
    f_row = r
    for m in range(nm):
        L = mc(m)
        put(ret, f"B{r}", f"=Construction!{L}{CR['date']}", DATE, align="right")
        put(ret, f"D{r}", f"=-(Construction!{L}{CR['land']}+Construction!{L}{CR['epc']}+Construction!{L}{CR['dev']})", M1)
        put(ret, f"E{r}", f"=-Construction!{L}{CR['eq']}", M1)
        r += 1
    for k, d in enumerate(eng.q):
        if not (eng.cod < d <= eng.ppa_end):
            continue
        L = qc(k)
        put(ret, f"B{r}", f"=Timing!{L}{T['end']}", DATE, align="right")
        put(ret, f"D{r}", f"=Operations!{L}{OP['cfads']}", M1)
        put(ret, f"E{r}", f"=Waterfall!{L}{W['div']}", M1)
        r += 1
    l_row = r - 1
    end_marker(ret, r + 1, 6)
    for i, (k, lab, f, fmt) in enumerate([
            ("p_irr", "Project IRR (pre-financing, XIRR)", f"=XIRR(D{f_row}:D{l_row},B{f_row}:B{l_row})", PCT),
            ("e_irr", "Equity IRR (XIRR)", f"=XIRR(E{f_row}:E{l_row},B{f_row}:B{l_row})", PCT),
            ("e_npv", "Equity NPV at the client minimum IRR (XNPV)", f"=XNPV({I['target']},E{f_row}:E{l_row},B{f_row}:B{l_row})", M1),
            ("eq_in", "Total equity injected", f"=-SUM(E{f_row}:E{f_row + nm - 1})", M1)]):
        ret[f"B{res_row + i}"] = lab
        ret[f"B{res_row + i}"].font = Font(name="Arial", size=10, bold=(i < 2))
        (output_box if i < 2 else put)(ret, f"E{res_row + i}", f, fmt)
        RT[k] = res_row + i

    # ---------------- Dashboard (answers + chart)
    O = ROWS["DASH"]
    r = sheet_header(dash, case, "Dashboard: one line per question in the brief, then lender metrics", 12, const_labels=("Value", ""))
    lk = lambda s, cref: f"={s}!{cref}"  # noqa: E731

    def orow(key, lab, unit, f, fmt=M1, box=False):
        nonlocal r
        label(dash, r, lab, unit, bold=box)
        O[key] = r
        (output_box if box else put)(dash, f"D{r}", f, fmt)
        r += 1

    def osec_(text):
        nonlocal r
        section(dash, r, text, last_col=12)
        r += 1

    osec_("1. Financing cost during construction")
    orow("idc", "Interest during construction", U, lk("Construction", f"$D${CR['idc']}"))
    orow("fee", "Upfront fee", U, lk("Construction", f"$D${CR['fee']}"))
    orow("fin", "Total financing cost during construction", U, f"=D{O['idc']}+D{O['fee']}", box=True)
    orow("dsra0", "DSRA initial funding at COD", U, lk("Construction", f"$D${CR['dsra']}"))
    orow("tpc", "Total project cost (incl. IDC, fee and DSRA)", U, lk("Construction", f"$D${CR['uses']}"))
    r += 1
    osec_("2. Revenue and operating costs over the project life")
    orow("rev", "Total revenue", U, lk("Operations", f"$D${OP['rev']}"), box=True)
    orow("opex", "Total operating costs", U, lk("Operations", f"$D${OP['opex']}"), box=True)
    r += 1
    osec_("3. Maximum debt (sizing)")
    orow("pv", "PV of CFADS over the repayment period", U, lk("Debt", f"$D${sz['pv']}"))
    orow("d_dscr", "Debt capacity at the minimum DSCR", U, lk("Debt", f"$D${sz['d_dscr']}"))
    orow("d_gear", "Debt capacity at maximum gearing", U, lk("Debt", f"$D${sz['d_gear']}"))
    orow("d", "Maximum debt = lower of the two", U, lk("Debt", f"$D${sz['d']}"), box=True)
    orow("bind", "Binding constraint", "", lk("Debt", f"$D${sz['bind']}"), "@")
    dash[f"D{O['bind']}"].alignment = Alignment(horizontal="right")
    orow("gear_act", "Actual gearing (debt / total project cost)", "%", f"={I['Dp']}/D{O['tpc']}", PCT)
    r += 1
    osec_("4–5. Repayment profile and DSCR")
    orow("flat", "Flat DSCR used for sculpting", "x", lk("Debt", f"$D${sz['flat']}"), MULT)
    orow("nrep", "Number of repayment quarters", "#", lk("Debt", f"$D${sz['nrep']}"), "0")
    orow("min_dscr", "Minimum DSCR (repayment period)", "x", f"=MIN(Debt!{qc(0)}{DB['dscr']}:{lastq}{DB['dscr']})", MULT, box=True)
    orow("avg_dscr", "Average DSCR (repayment period)", "x",
         f"=AVERAGEIFS(Debt!{qc(0)}{DB['dscr']}:{lastq}{DB['dscr']},Timing!{qc(0)}{T['f_rep']}:{lastq}{T['f_rep']},1)", MULT, box=True)
    orow("end_bal", "Debt outstanding at final maturity (must be 0)", U,
         f"=SUMIFS(Debt!{qc(0)}{DB['close']}:{lastq}{DB['close']},Timing!{qc(0)}{T['end']}:{lastq}{T['end']},{I['maturity']})")
    r += 1
    osec_("6. Returns")
    orow("p_irr", "Project IRR (pre-financing, XIRR)", "%", lk("Returns", f"$E${RT['p_irr']}"), PCT, box=True)
    orow("e_irr", "Equity IRR (XIRR)", "%", lk("Returns", f"$E${RT['e_irr']}"), PCT, box=True)
    orow("vs_target", "Equity IRR vs client minimum", "%", f"=D{O['e_irr']}-{I['target']}", PCT)
    orow("e_npv", "Equity NPV at the client minimum IRR", U, lk("Returns", f"$E${RT['e_npv']}"))
    r += 1
    osec_("7. Lender metrics")
    orow("min_llcr", "Minimum LLCR (includes the DSRA balance)", "x", f"=MIN(Debt!{qc(0)}{DB['llcr']}:{lastq}{DB['llcr']})", MULT, box=True)
    orow("min_plcr", "Minimum PLCR", "x", f"=MIN(Debt!{qc(0)}{DB['plcr']}:{lastq}{DB['plcr']})", MULT, box=True)
    orow("short", "Cash shortfall in any quarter (must be 0)", U, lk("Waterfall", f"$D${W['short']}"))
    r += 1
    osec_("Bonus: lowest tariff the client can bid (equity IRR = client minimum)")
    label(dash, r, "Lowest PPA tariff for the client's minimum equity IRR", "USD/MWh", bold=True)
    O["be"] = r
    output_box(dash, f"D{r}", round(be, 4), '#,##0.00')
    comment(dash, f"E{r}", "solved by build_pf.py (debt, IDC, fee and DSRA re-sized at each tariff); re-run after changing inputs")
    r += 2
    end_marker(dash, r, 12)
    # chart: CFADS vs debt service by year (repayment years)
    ch = BarChart()
    ch.type = "col"
    ch.y_axis.title = None
    ch.y_axis.majorGridlines = None
    ch.height, ch.width = 8.5, 17
    rep_years = [i for i, y in enumerate(years) if eng.first_rep.year <= y <= eng.maturity.year]
    i0, i1 = rep_years[0], rep_years[-1]
    data = Reference(ann, min_col=Q0 + i0, max_col=Q0 + i1, min_row=A["m_cfads"], max_row=A["m_ds"])
    cats = Reference(ann, min_col=Q0 + i0, max_col=Q0 + i1, min_row=A["yr"], max_row=A["yr"])
    ch.add_data(data, from_rows=True, titles_from_data=False)
    ch.set_categories(cats)
    for s_, colour, nm_ in zip(ch.series, (NAVY, STEEL), ("CFADS", "Senior debt service")):
        s_.graphicalProperties.solidFill = colour
        s_.graphicalProperties.line.solidFill = colour
        s_.tx = SeriesLabel(v=nm_)
    ch.legend.position = "t"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.numFmt = "0"
    ch.x_axis.tickLblPos = "low"
    ch.gapWidth = 60
    dash.add_chart(ch, f"G{O['idc']}")
    hdr = dash[f"G{O['idc'] - 1}"]
    hdr.value = "CFADS vs senior debt service by year (USD m)"
    hdr.font = Font(name="Arial", size=10, bold=True)

    # ---------------- Scenarios
    S = ROWS["SCN"]
    r = sheet_header(scn, case, "Scenarios: debt re-sized in every case · equity IRR sensitivity at the sized debt", 11)
    section(scn, r, "Scenario analysis (each case fully re-solved by build_pf.py: debt, IDC, fee and DSRA re-sized)", last_col=11)
    r += 1
    note(scn, f"B{r}", "Values are pasted from the solver run (blue). Re-run build_pf.py after changing inputs.")
    r += 1
    header_row(scn, r, [n_ for n_, _ in scen], Q0)
    for j in range(len(scen)):
        scn.cell(r, Q0 + j).alignment = Alignment(wrap_text=True, horizontal="right")
    scn.row_dimensions[r].height = 27
    r += 1
    metrics = [("D", "Senior debt", U, M1), ("bind", "Binding constraint", "", "@"), ("TPC", "Total project cost", U, M1),
               ("dsra0", "DSRA initial funding", U, M1), ("min_dscr", "Minimum DSCR", "x", MULT),
               ("avg_dscr", "Average DSCR", "x", MULT), ("min_llcr", "Minimum LLCR", "x", MULT),
               ("p_irr", "Project IRR", "%", PCT), ("e_irr", "Equity IRR", "%", PCT),
               ("vs", "Equity IRR vs client minimum", "%", PCT)]
    for key, lab, u_, fmt in metrics:
        label(scn, r, lab, u_, bold=key in ("D", "e_irr"))
        S[key] = r
        for j, (_n, rr_) in enumerate(scen):
            v = rr_["e_irr"] - c["target_equity_irr"] if key == "vs" else rr_[key]
            put(scn, f"{col(Q0 + j)}{r}", v if isinstance(v, str) else round(v, 12), fmt, bold=key in ("D", "e_irr"), align="right")
        r += 1
    rng_dscr = f"{col(Q0)}{S['min_dscr']}:{col(Q0 + len(scen) - 1)}{S['min_dscr']}"
    scn.conditional_formatting.add(rng_dscr, CellIsRule(operator="lessThan", formula=[I["mindscr"]], font=Font(color="FF0000", bold=True)))
    rng_vs = f"{col(Q0)}{S['vs']}:{col(Q0 + len(scen) - 1)}{S['vs']}"
    scn.conditional_formatting.add(rng_vs, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="FF0000", bold=True)))
    comment(scn, f"{col(Q0 + len(scen))}{S['min_dscr']}", "red = below the sizing DSCR")
    comment(scn, f"{col(Q0 + len(scen))}{S['vs']}", "red = below the client's minimum return")
    r += 1
    section(scn, r, "Equity IRR sensitivity – PPA tariff vs capacity factor (debt held at the sized amount)", last_col=11)
    r += 1
    label(scn, r, "Tariff override (blank = Inputs; flexed by the data table)", "USD/MWh")
    input_cell(scn, f"D{r}", None, '#,##0.00')
    S["tar_ovr"] = r
    r += 1
    label(scn, r, "Capacity factor override (blank = Inputs; flexed by the data table)", "%")
    input_cell(scn, f"D{r}", None, PCT)
    S["cuf_ovr"] = r
    r += 2
    base_cuf = c["cuf"][c["energy_case"]]
    tar_vals = [c["tariff"] + k * 5 for k in (-2, -1, 0, 1, 2)]
    cuf_vals = [base_cuf + k * 0.01 for k in (-2, -1, 0, 1, 2)]
    rng_dt, _b = sensitivity(scn, r, 4, f"=Returns!$E${RT['e_irr']}", cuf_vals, tar_vals, "Capacity factor", "PPA tariff (USD/MWh)",
                             PCT, '#,##0', PCT, base_side=base_cuf, base_top=c["tariff"], better="high")
    r += len(cuf_vals) + 2
    note(scn, f"D{r}", "Excel data table; base case boxed. Debt is not re-sized here (see the scenario table for re-sized cases).")
    r += 2
    end_marker(scn, r, 11)
    for ws in (inp,):
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = cell.value.replace("@SCN_CUF@", f"Scenarios!$D${S['cuf_ovr']}").replace("@SCN_TAR@", f"Scenarios!$D${S['tar_ovr']}")
    tables = [(f"Scenarios!{rng_dt}", f"Scenarios!D{S['tar_ovr']}", f"Scenarios!D{S['cuf_ovr']}")]

    # ---------------- Checks
    r = sheet_header(chk, case, "Checks: every line must show OK (0); the Cover shows the total", 8, const_labels=("Value", ""))
    section(chk, r, "Integrity checks", last_col=8)
    r += 1
    checks = [
        ("Pasted debt = live sizing", f"={I['Dp']}-Dashboard!D{O['d']}"),
        ("Pasted project cost = live project cost", f"={I['TPCp']}-Dashboard!D{O['tpc']}"),
        ("Debt drawn in construction = debt amount", f"=Construction!{lastm}{CR['close']}-{I['Dp']}"),
        ("Sources = uses in construction", f"=Construction!D{CR['src']}-Construction!D{CR['uses']}"),
        ("Debt repaid in full at final maturity", f"=Dashboard!D{O['end_bal']}"),
        ("No cash shortfall in any quarter", f"=Waterfall!D{W['short']}"),
        ("DSRA fully released by the last quarter", f"=Waterfall!{lastq}{W['d_close']}"),
        ("Balance sheet balances (all years)", f"=SUMPRODUCT(ABS(Annual!{ac(0)}{A['bal']}:{lastA}{A['bal']}))"),
        ("Cash never negative", f"=-MIN(0,MIN(Annual!{ac(0)}{A['cash']}:{lastA}{A['cash']}))"),
        ("Min DSCR ≥ sizing DSCR (when sculpted)", f"=IF({I['rep']}=1,MAX(0,{I['mindscr']}-Dashboard!D{O['min_dscr']}),0)"),
        ("Scenario table base case = live model (debt)", f"=Scenarios!{col(Q0)}{S['D']}-Dashboard!D{O['d']}"),
        ("Scenario table base case = live model (equity IRR)", f"=Scenarios!{col(Q0)}{S['e_irr']}-Dashboard!D{O['e_irr']}"),
    ]
    first = r
    for lab, f in checks:
        label(chk, r, lab)
        check_cell(chk, f"D{r}", f)
        r += 1
    label(chk, r, "Total (sum of absolute values)", bold=True)
    tot_chk = f"D{r}"
    check_cell(chk, tot_chk, f"=SUMPRODUCT(ABS(D{first}:D{r - 1}))")
    chk[tot_chk].border = TOTAL
    end_marker(chk, r + 2, 8)

    # ---------------- Cover
    r = sheet_header(cov, c.get("case", c["project"]), f"{U} · deal-desk house style v{STYLE_VERSION} · no circular references", 8,
                     units_label=False)
    section(cov, r, "Contents", last_col=8)
    r += 1
    for a, b in [("Inputs", "All assumptions; solver values (debt, project cost) pasted with a check"),
                 ("Timing", "Quarterly timeline and flags (construction, operations, repayment, loan life, DSRA)"),
                 ("Construction", "Monthly: capex, fee, IDC, DSRA funding, equity-first funding, debt drawdown"),
                 ("Operations", "Quarterly: generation, revenue, O&M, EBITDA, tax, CFADS"),
                 ("Debt", "Sizing, sculpted repayment, DSCR, LLCR, PLCR"),
                 ("Waterfall", "CFADS → senior debt service → DSRA → dividends; DSRA account"),
                 ("Annual", "Income statement, cash flow statement, balance sheet, annual DSCR"),
                 ("Returns", "Dated cash flows: project and equity XIRR, equity XNPV"),
                 ("Scenarios", "Downside cases re-solved; tariff × capacity factor data table"),
                 ("Dashboard", "Answers to each question in the brief, lender metrics, chart"),
                 ("Checks", "Integrity checks (all must be 0)")]:
        label(cov, r, a, bold=True)
        cov[f"C{r}"] = b
        r += 1
    r += 1
    section(cov, r, "Colour code", last_col=8)
    r = colour_key(cov, r + 1) + 1
    section(cov, r, "Key outputs", last_col=8)
    r += 1
    for lab, u_, f, fmt in [("Senior debt", U, f"=Dashboard!$D${O['d']}", M1), ("Total project cost", U, f"=Dashboard!$D${O['tpc']}", M1),
                            ("Minimum DSCR", "x", f"=Dashboard!$D${O['min_dscr']}", MULT),
                            ("Minimum LLCR", "x", f"=Dashboard!$D${O['min_llcr']}", MULT),
                            ("Project IRR", "%", f"=Dashboard!$D${O['p_irr']}", PCT), ("Equity IRR", "%", f"=Dashboard!$D${O['e_irr']}", PCT),
                            ("Checks (0 = all OK)", "", f"=Checks!${tot_chk[0]}${tot_chk[1:]}", '[>=0.0005]0.000;[<=-0.0005](0.000);"All OK"')]:
        label(cov, r, lab, u_)
        output_box(cov, f"D{r}", f, fmt)
        r += 1
    if c.get("assumptions_log"):
        r += 1
        section(cov, r, "Assumptions made (the brief was silent or ambiguous)", last_col=8)
        r += 1
        for i, a in enumerate(c["assumptions_log"]):
            cov[f"B{r}"] = f"{i + 1}. {a}"
            r += 1
    end_marker(cov, r + 1, 8)

    # resolve remaining cross-sheet row tokens
    tok = re.compile(r"@(CON|OPN|DBT|WF|TIM|ANN):(\w+)@")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "@" in cell.value:
                    cell.value = tok.sub(lambda mm: str(ROWS[mm.group(1)][mm.group(2)]), cell.value)
    wb.active = 0
    wb.save(path)
    return res, be, O, tot_chk, tables, scen


def verify(path, res, be, O, tot_chk, tables):
    py = sys.executable
    office = os.path.join(HERE, "office.py")
    for rng, row_in, col_in in tables:
        subprocess.run([py, office, "datatable", path, rng, row_in, col_in], check=True)
    rc = subprocess.run([py, office, "recalc", path], capture_output=True, text=True)
    print(rc.stdout)
    keys = ["d", "tpc", "idc", "fee", "dsra0", "p_irr", "e_irr", "e_npv", "min_dscr", "avg_dscr", "min_llcr", "min_plcr", "rev", "opex"]
    refs = [f"Dashboard!D{O[k]}" for k in keys] + [f"Checks!{tot_chk}"]
    xl = json.loads(subprocess.run([py, office, "values", path, *refs], capture_output=True, text=True, check=True).stdout)
    py_vals = {"d": res["D"], "tpc": res["TPC"], "idc": res["idc"], "fee": res["fee"], "dsra0": res["dsra0"], "p_irr": res["p_irr"],
               "e_irr": res["e_irr"], "e_npv": res["eq_npv"], "min_dscr": res["min_dscr"], "avg_dscr": res["avg_dscr"],
               "min_llcr": res["min_llcr"], "min_plcr": res["min_plcr"], "rev": res["revenue"], "opex": res["opex"]}
    ok = True
    print(f"{'Output':<12}{'Excel':>16}{'Python':>16}   match")
    for k, ref in zip(keys, refs):
        a, b = xl[ref], py_vals[k]
        tol = 1e-6 if k not in ("p_irr", "e_irr") else 1e-5
        good = a is not None and abs(a - b) <= tol * max(1, abs(b))
        ok &= good
        print(f"{k:<12}{(a if a is not None else float('nan')):>16,.6f}{b:>16,.6f}   {'OK' if good else 'MISMATCH'}")
    print(f"Lowest tariff for target equity IRR: {be:,.2f} USD/MWh")
    chk = xl[refs[-1]]
    print(f"Checks total: {chk} -> {'all OK' if chk is not None and abs(chk) < 1e-3 else 'FAILED'}")
    lint = subprocess.run([py, os.path.join(HERE, "model_lint.py"), path], capture_output=True, text=True)
    print(lint.stdout)
    return ok and chk is not None and abs(chk) < 1e-3 and '"status": "success"' in rc.stdout and lint.returncode == 0


def resolve(path, max_iter=100):
    """Re-solve pasted debt / project cost in an existing model via Excel (after inputs were changed by hand)."""
    from office import _excel, _quit   # _quit() makes sure no hidden EXCEL.EXE is left behind
    xl = _excel()
    try:
        wb = xl.Workbooks.Open(os.path.abspath(path))
        inp, out = wb.Worksheets("Inputs"), wb.Worksheets("Dashboard")
        find = lambda ws, text: next(r for r in range(1, 300) if (ws.Cells(r, 2).Value or "").startswith(text))  # noqa: E731
        rD, rT = find(inp, "Debt amount (pasted)"), find(inp, "Total project cost (pasted)")
        oD, oT = find(out, "Maximum debt"), find(out, "Total project cost")
        for i in range(max_iter):
            xl.CalculateFull()
            d_live, t_live = out.Cells(oD, 4).Value, out.Cells(oT, 4).Value
            d_p, t_p = inp.Cells(rD, 4).Value, inp.Cells(rT, 4).Value
            if abs(d_live - d_p) < 1e-9 and abs(t_live - t_p) < 1e-9:
                break
            inp.Cells(rD, 4).Value, inp.Cells(rT, 4).Value = d_live, t_live
        wb.Save()
        wb.Close(False)
        print(f"resolved after {i} iterations: debt {d_live:,.4f}, project cost {t_live:,.4f}")
        print("Note: the scenario table and the bonus 'lowest tariff' are not re-solved here; re-run build_pf.py for those.")
    finally:
        _quit(xl)


if __name__ == "__main__":
    if sys.argv[1] == "--resolve":
        resolve(sys.argv[2])
        sys.exit(0)
    cfg = json.load(open(sys.argv[1], encoding="utf-8"))
    out_path = os.path.abspath(sys.argv[2])
    R, BE, Orows, Tot, Tables, _S = build(cfg, out_path)
    print(f"built {out_path}")
    if sys.platform != "win32" and "--no-verify" not in sys.argv:
        print("MAC BUILD: workbook written. The data tables, Excel recalc and independent check need Excel, so they are "
              "skipped here (sensitivity tables stay empty, cached values missing). Finish on the PC: run this same command "
              "there; it rebuilds identically from the case file, then verifies.")
        sys.argv.append("--no-verify")
    if "--no-verify" not in sys.argv:
        sys.exit(0 if verify(out_path, R, BE, Orows, Tot, Tables) else 1)
