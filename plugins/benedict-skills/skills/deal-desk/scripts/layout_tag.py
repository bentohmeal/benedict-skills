"""Layout library: one catalogue of every slide layout deal-desk can build (assets/layout_library.json).

    python layout_tag.py scan                 -> assets/layout_scan.json: variants + text slots per Infographic Library slide
    python layout_tag.py sheets  [out_dir]    -> contact sheets (12 thumbnails each, variants boxed and numbered) for tagging
    python layout_tag.py name                 -> rename text shapes in Infographic Library.pptx to "IG slot:v<k>.<role><n>"
                                                 (backs the file up first; re-run after infographic_library.py rebuilds it)
    python layout_tag.py thumbs               -> assets/layout_thumbs/: ig-NNN.png (library) + gal-NN.png (gallery), 480 px
    python layout_tag.py build                -> assets/layout_library.json from scan + hand tags + coded components
    python layout_tag.py validate             -> check layout_library.json against the schema (exit 1 on any error)
    python layout_tag.py find <section> [data_type] [items]
                                              -> shortlist for one slide, layouts seen on real bank pages first

Sources of records:
  - Infographic Library (211 PwC originals): geometry from `scan`, meaning from assets/layout_tags.json (hand-written by
    looking at the thumbnails: use when / don't use when, section, data type, item counts, slot order).
  - Coded ibdeck components (fin_chart, football_field, risk_matrix...): COMPONENTS below.
  - The gallery's page recipes (Deck Layout Gallery.pptx): layout_tags.json "gallery" (recipe: true), which also maps
    each component to its gallery thumbnail.

Slots: inside one variant (a sub-diagram; PwC pages often show 2-5 side by side), text shapes are named in reading
order (or clockwise from 12 o'clock when the tag says order = "clockwise") with a role:
  head = short bold / larger-font label, num = number or % placeholder, text = everything else.
Deck.layout(s, "ig045", variant=1, data={"head": [...], "text": [...]}) fills them (ibdeck.py).
"""
import json
import math
import os
import re
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ASSETS = os.path.join(HERE, "..", "assets")
SCAN = os.path.join(ASSETS, "layout_scan.json")
TAGS = os.path.join(ASSETS, "layout_tags.json")
LIB_JSON = os.path.join(ASSETS, "layout_library.json")
THUMBS = os.path.join(ASSETS, "layout_thumbs")
NOTES = os.path.join(ASSETS, "eval", "layout_notes.json")  # wording fixes from the tuning loop (layout_loop.py)
PAGES_JSON = os.path.join(ASSETS, "real_deck_pages.json")   # exemplar thumbs are relative to its "root"
EMU = 914400
GAP = 0.12   # inches: shapes closer than this belong to the same sub-diagram
NUM_RE = re.compile(r"^[\s$€£]*(\d+([.,]\d+)?\s*(%|x|m|bn|k)?|xx?%)\s*$", re.I)


def _lib():
    from ibdeck import LIBRARY
    return LIBRARY


def _gallery():
    from ibdeck import TOOLKIT
    return os.path.join(TOOLKIT, "Deck Layout Gallery.pptx")


# ---------------------------------------------------------------- geometry
def _walk(sh):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    yield sh
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        for c in sh.shapes:
            yield from _walk(c)


def _box(sh):
    return (sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU)


def _near(a, b, gap=GAP):
    return not (a[2] + gap < b[0] or b[2] + gap < a[0] or a[3] + gap < b[1] or b[3] + gap < a[1])


def _union(boxes):
    return (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes))


def _font(sh):
    """(max font pt, any bold) over the runs of a text frame."""
    size, bold = 0, False
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                size = max(size, r.font.size.pt)
            bold = bold or bool(r.font.bold)
    return size, bold


def _filled(sh):
    """True when a top-level shape draws something (fill, line, picture, group, table, chart), not a bare text box."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if sh.shape_type in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE) or getattr(sh, "has_chart", False) \
            or getattr(sh, "has_table", False):
        return True
    xml = sh._element.xml
    if "<a:noFill/>" in xml.split("<a:ln")[0] and not re.search(r"<a:ln[^>]*>\s*<a:solidFill", xml):
        return False
    return "<a:solidFill>" in xml or "<a:gradFill>" in xml or "<a:custGeom>" in xml or sh.shape_type != 17


def _components(boxes, gap=GAP):
    """Connected components (lists of indices) of boxes that touch or sit within `gap` of each other."""
    parent = list(range(len(boxes)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            if _near(boxes[i], boxes[j], gap):
                parent[find(i)] = find(j)
    comps = {}
    for i in range(len(boxes)):
        comps.setdefault(find(i), []).append(i)
    return list(comps.values())


def clusters(shapes, crops=None):
    """Group top-level shapes into sub-diagrams. Default: connected components of touching / near boxes (a component
    of bare text boxes only is a PwC annotation such as "7 petals", not a layout). Where that split is wrong
    (rotated freeforms and tightly packed diagrams chain together), the hand tag gives `crops`: one box per variant in
    library-slide inches; a shape belongs to the first crop holding its centre."""
    boxes = [_box(sh) for sh in shapes]
    if crops:
        groups = [[] for _ in crops]
        for i, b in enumerate(boxes):
            cx, cy = (b[0] + b[2]) / 2, (b[1] + b[3]) / 2
            for g, c in zip(groups, crops):
                if c[0] <= cx <= c[2] and c[1] <= cy <= c[3]:
                    g.append(i)
                    break
        comps = [g for g in groups if g]
    else:
        comps = _components(boxes)
    out = []
    for idx in comps:
        members = [shapes[i] for i in idx]
        out.append({"shapes": members, "box": _union([boxes[i] for i in idx]),
                    "annotation": not any(_filled(sh) for sh in members)})
    if not crops:   # reading order of the variants: rows (by top, 0.6 in tolerance), then left to right
        out.sort(key=lambda c: (round(c["box"][1] / 0.6), c["box"][0]))
    return out


def variants(slide, tag=None):
    """[(variant no, [shape], box, [(slot name, text shape, box)])] for one library slide (annotations dropped).
    tag = the slide's hand tag (order, crops)."""
    tag = tag or {}
    shapes = [sh for sh in slide.shapes if sh.name.startswith("IG ")]
    out, k = [], 0
    for c in clusters(shapes, tag.get("crops")):
        if c["annotation"]:
            continue
        k += 1
        frames = _order([f for sh in c["shapes"] for f in text_frames(sh)], tag.get("order", "rows"))
        n, slots = {}, []
        for (sh, b), r in zip(frames, roles(frames)):
            n[r] = n.get(r, 0) + 1
            slots.append((f"v{k}.{r}{n[r]}", sh, b))
        out.append((k, c["shapes"], c["box"], slots))
    return out


def _absolute_box(sh, chain):
    """Box in slide coordinates for a shape inside groups (python-pptx gives child-space coordinates)."""
    x0, y0, x1, y1 = _box(sh)
    for g in reversed(chain):
        xfrm = g._element.grpSpPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
        if xfrm is None:
            continue
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        off, ext, choff, chext = (xfrm.find(ns + t) for t in ("off", "ext", "chOff", "chExt"))
        if None in (off, ext, choff, chext) or int(chext.get("cx")) == 0 or int(chext.get("cy")) == 0:
            continue
        sx, sy = int(ext.get("cx")) / int(chext.get("cx")), int(ext.get("cy")) / int(chext.get("cy"))
        ox, oy, cx, cy = (int(off.get("x")) / EMU, int(off.get("y")) / EMU, int(choff.get("x")) / EMU,
                          int(choff.get("y")) / EMU)
        x0, x1 = ox + (x0 - cx) * sx, ox + (x1 - cx) * sx
        y0, y1 = oy + (y0 - cy) * sy, oy + (y1 - cy) * sy
    return x0, y0, x1, y1


def text_frames(shape):
    """[(shape, absolute box)] for every text-bearing shape inside a top-level shape (groups walked)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []

    def rec(sh, chain):
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for c in sh.shapes:
                rec(c, chain + [sh])
        elif sh.has_text_frame and sh.text_frame.text.strip():
            out.append((sh, _absolute_box(sh, chain)))
    rec(shape, [])
    return out


def _order(frames, how):
    if how == "clockwise" and len(frames) > 2:
        cx = sum((b[0] + b[2]) / 2 for _, b in frames) / len(frames)
        cy = sum((b[1] + b[3]) / 2 for _, b in frames) / len(frames)

        def ang(f):
            b = f[1]
            mx, my = (b[0] + b[2]) / 2 - cx, (b[1] + b[3]) / 2 - cy
            if math.hypot(mx, my) < 0.25:
                return -1                      # the centre label comes first
            return (math.degrees(math.atan2(mx, -my)) + 360) % 360
        return sorted(frames, key=ang)
    if how == "columns":
        return sorted(frames, key=lambda f: (round(f[1][0] / 0.5), f[1][1]))
    return sorted(frames, key=lambda f: (round(f[1][1] / 0.25), f[1][0]))   # rows: top to bottom, left to right


def roles(frames):
    """head / num / text per frame: numbers by pattern; heads = bold or larger than the variant's typical size."""
    sizes = sorted(_font(sh)[0] for sh, _ in frames) or [10]
    typical = sizes[len(sizes) // 2]
    out = []
    for sh, _ in frames:
        t = sh.text_frame.text.strip()
        size, bold = _font(sh)
        if NUM_RE.match(t.replace("\n", "")):
            out.append("num")
        elif (bold or size > typical) and len(t) <= 40:
            out.append("head")
        else:
            out.append("text")
    return out


# ---------------------------------------------------------------- scan / name
def scan(path=None, tags=None, rename=False):
    from pptx import Presentation
    path = path or _lib()
    tags = tags if tags is not None else _load(TAGS, {}).get("ig", {})
    prs = Presentation(path)
    out = {}
    for no, s in enumerate(prs.slides, 1):
        tag = tags.get(f"ig{no:03d}", {})
        rec = []
        for k, shp, box, slots in variants(s, tag):
            counts = {}
            for name, sh, b in slots:
                r = name.split(".")[1].rstrip("0123456789")
                counts[r] = counts.get(r, 0) + 1
                if rename:
                    sh._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr") \
                        .set("name", f"IG slot:{name}")
            x0, y0, x1, y1 = box
            rec.append({"variant": k, "box": [round(v, 2) for v in box], "w": round(x1 - x0, 2),
                        "h": round(y1 - y0, 2), "shapes": len(shp),
                        "has_table": any(getattr(sh, "has_table", False) and sh.has_table for sh in shp),
                        "has_chart": any(getattr(sh, "has_chart", False) and sh.has_chart for sh in shp),
                        "slots": [{"slot": name, "sample": sh.text_frame.text.strip()[:40],
                                   "chars": len(sh.text_frame.text.strip()), "box": [round(v, 2) for v in b]}
                                  for name, sh, b in slots],
                        "counts": counts})
        out[f"ig{no:03d}"] = rec
    if rename:
        bak = path.replace(".pptx", " (backup before slot names).pptx")
        if not os.path.exists(bak):
            shutil.copy2(path, bak)
        prs.save(path)
        print(f"renamed slots in {path} (backup: {os.path.basename(bak)})")
    return out


def _load(p, default):
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    return default


def _dump(p, obj):
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=1, ensure_ascii=False)
        f.write("\n")


# ---------------------------------------------------------------- thumbnails / contact sheets
def thumbs(tmp=None):
    import tempfile
    from office import render
    import pymupdf
    os.makedirs(THUMBS, exist_ok=True)
    tmp = tmp or tempfile.mkdtemp(prefix="thumbs_")
    for src, prefix in ((_lib(), "ig"), (_gallery(), "gal")):
        d = os.path.join(tmp, prefix)
        files = sorted((f for f in os.listdir(d) if f.endswith(".png")),       # slide-9 before slide-10
                       key=lambda f: int(re.sub(r"\D", "", f) or 0)) if os.path.isdir(d) else None
        files = [os.path.join(d, f) for f in files] if files else render(src, d)
        for i, f in enumerate(files, 1):
            pix = pymupdf.Pixmap(f)
            fac = 480 / pix.width
            small = pymupdf.Pixmap(pix, int(pix.width * fac), int(pix.height * fac), None) if fac < 1 else pix
            small.save(os.path.join(THUMBS, f"{prefix}-{i:03d}.png" if prefix == "ig" else f"{prefix}-{i:02d}.png"))
    print(f"thumbnails in {os.path.abspath(THUMBS)}")


def sheets(out_dir, render_dir, per=12, cols=3):
    """Contact sheets for tagging: 1600x900 renders, variant boxes drawn and numbered (needs `scan` output)."""
    import pymupdf
    os.makedirs(out_dir, exist_ok=True)
    sc = _load(SCAN, {})
    ids = sorted(sc)
    W_, H_ = 800, 450
    for start in range(0, len(ids), per):
        chunk = ids[start:start + per]
        rows = math.ceil(len(chunk) / cols)
        doc = pymupdf.open()
        page = doc.new_page(width=cols * W_, height=rows * (H_ + 24))
        for j, lid in enumerate(chunk):
            no = int(lid[2:])
            x, y = (j % cols) * W_, (j // cols) * (H_ + 24)
            rect = pymupdf.Rect(x, y + 24, x + W_, y + 24 + H_)
            page.insert_image(rect, filename=os.path.join(render_dir, f"slide-{no:02d}.png" if no < 100
                                                          else f"slide-{no}.png"))
            page.insert_text((x + 6, y + 18), lid, fontsize=18, color=(0.8, 0, 0))
            f = W_ / 13.333
            for inch in range(1, 14):        # 1-inch grid, for reading crop boxes off the sheet
                page.draw_line((x + inch * f, y + 24), (x + inch * f, y + 24 + H_), color=(0.2, 0.7, 0.2), width=0.4)
                page.insert_text((x + inch * f + 2, y + 24 + 10), str(inch), fontsize=9, color=(0.2, 0.6, 0.2))
            for inch in range(1, 8):
                page.draw_line((x, y + 24 + inch * f), (x + W_, y + 24 + inch * f), color=(0.2, 0.7, 0.2), width=0.4)
                page.insert_text((x + 2, y + 24 + inch * f - 2), str(inch), fontsize=9, color=(0.2, 0.6, 0.2))
            for v in sc[lid]:
                if v.get("annotation"):
                    continue
                b = v["box"]
                r = pymupdf.Rect(x + b[0] * f, y + 24 + b[1] * f, x + b[2] * f, y + 24 + b[3] * f)
                page.draw_rect(r, color=(0.85, 0, 0), width=1.5)
                page.insert_text((r.x0 + 3, r.y0 + 16), f"v{v['variant']}", fontsize=16, color=(0.85, 0, 0))
        pix = page.get_pixmap(dpi=72)
        f = os.path.join(out_dir, f"sheet-{chunk[0]}-{chunk[-1]}.png")
        pix.save(f)
        print(f)


# ---------------------------------------------------------------- coded components and gallery recipes
SECTIONS = ["cover", "exec_summary", "company", "market", "thesis", "financials", "valuation", "returns", "structure",
            "risks", "process", "recommendation", "appendix", "any"]
DATA_TYPES = ["process", "cycle", "comparison", "hierarchy", "geography", "time_series", "kpi_set", "matrix",
              "composition", "list", "relationship", "text", "bridge", "range", "icons"]

# id, call, data type, sections, items (min, max), numbers needed, use when, don't use when, slots (name: words max)
COMPONENTS = [
    ("stat_icons", "d.stat_icons(s, x, y, w, h, [(icon, value, label)], cols=2)", "kpi_set",
     ["company", "exec_summary", "returns"], (2, 8), True,
     "Company at a glance or key numbers, each with a meaningful icon; return drivers in numbers on a thesis page",
     "More than 8 numbers (use a table); numbers that need a trend (use fin_chart)",
     {"value": 3, "label": 10}),
    ("stat_strip", "d.stat_strip(s, x, y, w, h, [(value, label)])", "kpi_set",
     ["financials", "valuation", "risks", "market"], (3, 6), True,
     "A row of 3-6 headline numbers under a chart or risk table",
     "Numbers that need icons to be understood; more than 6 items", {"value": 3, "label": 8}),
    ("stat_rings", "d.stat_rings(s, x, y, w, h, [(pct, label)])", "composition",
     ["company", "market", "structure"], (2, 5), True,
     "Shares, margins, utilisation or a debt / equity split shown as % rings",
     "Values that are not percentages; shares that must add to 100% across many parts (use stacked or pie)",
     {"label": 8}),
    ("hero_number", "d.hero_number(s, x, y, w, h, value, label, points=[...])", "kpi_set",
     ["recommendation", "returns", "exec_summary", "valuation"], (1, 1), True,
     "One headline figure (IRR, offer price, EV) with 2-3 supporting points",
     "Several numbers of equal weight (use stat_icons)", {"value": 3, "label": 10, "point": 20}),
    ("pictogram", "d.pictogram(s, x, y, w, h, share, total=20, icon='person-filled')", "composition",
     ["market", "company"], (1, 1), True, "One share of a population (1 in 4 customers, 35% of capacity)",
     "Several shares to compare", {"label": 12}),
    ("tam_sam_som", "d.tam_sam_som(s, x, y, w, h, [(name, value, text)])", "hierarchy", ["market"], (3, 3), True,
     "Market sizing from total to serviceable to obtainable market",
     "A market trend over time (use fin_chart / column)", {"value": 3, "text": 12}),
    ("pill_rows", "d.pill_rows(s, x, y, w, h, [(label, text)])", "list",
     ["company", "thesis", "exec_summary"], (3, 6), False,
     "Business overview or summary: a short label pill beside 1-2 lines of text per row",
     "Rows needing long argued paragraphs (use side_tabs)", {"label": 3, "text": 30}),
    ("side_tabs", "d.side_tabs(s, x, y, w, h, [(tab, (claim, text))])", "list",
     ["thesis", "recommendation", "exec_summary"], (3, 5), False,
     "Investment thesis pillars: bold claim + evidence -> so what -> implication (35+ words per pillar)",
     "Short facts or one-liners (use pill_rows); a process (use chevrons)", {"tab": 3, "claim": 12, "text": 60}),
    ("chevrons", "d.chevrons(s, x, y, w, h, [(title, [bullets])], highlight=None)", "process",
     ["process", "thesis", "company"], (3, 6), False,
     "A sequence or plan (value creation plan, 100-day plan, deal steps), each step with 1-3 bullets",
     "Items with no order (use pill_rows / stat_icons)", {"title": 4, "bullet": 12}),
    ("timeline", "d.timeline(s, x, y, w, h, [(date, title, text)], highlight=None)", "time_series",
     ["company", "process", "market"], (4, 8), False,
     "History and milestones (company, policy or market events), or a deal timetable; highlight today",
     "Fewer than 4 dates; numeric trends (use a chart)", {"date": 2, "title": 4, "text": 8}),
    ("numbered_steps", "d.numbered_steps(s, x, y, w, h, [(title, text)], direction='down')", "process",
     ["process", "recommendation"], (3, 6), False, "Next steps, diligence plan or workstreams in order",
     "Steps that need sub-bullets each (use chevrons)", {"title": 4, "text": 15}),
    ("hub_spoke", "d.hub_spoke(s, x, y, w, h, centre, [(icon, title, text)])", "relationship",
     ["company", "thesis", "market"], (4, 6), False,
     "Growth drivers or a business model around one centre (the company, the platform)",
     "A sequence (use chevrons); more than 6 spokes", {"centre": 3, "title": 3, "text": 6}),
    ("pros_cons", "d.pros_cons(s, x, y, w, h, pros, cons)", "comparison", ["risks", "thesis", "recommendation"],
     (2, 6), False, "Strengths vs watch-points, or two sides of a decision", "A full risk register (use risk_matrix)",
     {"item": 15}),
    ("risk_matrix", "d.risk_matrix(s, x, y, w, h, [((title, text), (title, text), level)])", "matrix", ["risks"],
     (3, 6), False, "Key risks with a quantified impact, a mitigant + what remains, and a severity Harvey ball",
     "Fewer than 3 risks; risks with no mitigant", {"risk": 30, "mitigant": 30}),
    ("harvey", "d.harvey(s, x, y, w, rows, cols, scores)", "matrix", ["market", "company", "process"], (3, 10), False,
     "Scoring peers, buyers or options on several criteria (Harvey balls)", "Exact numbers (use a table)",
     {"row": 4, "col": 3}),
    ("logo_row", "d.logo_row(s, x, y, w, h, names, logos=..., highlight=None)", "comparison",
     ["market", "process", "company"], (3, 16), False,
     "Peers, buyers, sponsors, customers or tenants as logo tiles (target highlighted)",
     "Players that need numbers beside them (use a bar chart with the target navy)", {"name": 3}),
    ("table", "d.table(s, x, y, w, rows)", "comparison", ["any"], (2, 15), True,
     "Comparison tables, sources & uses, comps, precedent transactions, deal terms, requirements with detail",
     "A trend or one number per item (use a chart)", {}),
    ("fin_table", "d.fin_table(s, x, y, w, periods, lines)", "time_series", ["financials"], (4, 20), True,
     "Income statement / cash flow line items in bank format (A / E periods, % rows)",
     "A one-metric trend (use fin_chart)", {"line": 4}),
    ("fin_chart", "d.fin_chart(s, x, y, w, h, periods, values, strip=..., cagr=..., split=...)", "time_series",
     ["financials", "market"], (4, 12), True,
     "One metric's trend with CAGR, actual / estimate split and a % margin strip",
     "Several unrelated metrics (use fin_table)", {}),
    ("column", "d.column(s, x, y, w, h, categories, series)", "time_series", ["market", "financials"], (3, 12), True,
     "Market size or volumes over time, or a few series side by side", "Shares of a whole (use stacked / rings)", {}),
    ("bar", "d.bar(s, x, y, w, h, categories, series)", "comparison", ["market", "valuation"], (3, 15), True,
     "Ranking players by one metric (target navy, the rest grey)", "A time series (use column)", {}),
    ("stacked", "d.stacked(s, x, y, w, h, categories, series)", "composition", ["market", "financials"], (2, 10),
     True, "Mix over time (revenue by segment, capacity by operator)", "A single-period split (use stat_rings)", {}),
    ("line", "d.line(s, x, y, w, h, categories, series)", "time_series", ["market"], (6, 40), True,
     "Prices, rates or indices over many periods", "Fewer than 6 points (use column)", {}),
    ("football_field", "d.football_field(s, x, y, w, h, [(method, low, high)])", "range", ["valuation"], (3, 8),
     True, "Valuation ranges by method against the offer / current price", "A single point estimate", {"method": 4}),
    ("heatmap", "d.heatmap(s, x, y, w, corner, col_head, row_head, values)", "matrix", ["returns", "valuation"],
     (9, 49), True, "IRR / MOIC or price sensitivity grid on two drivers (red-white-green)",
     "One-way sensitivities (use a table or tornado)", {}),
    ("photo", "d.photo(s, path, x, y, w, h, credit=...)", "geography", ["company", "market", "cover"], (1, 1), False,
     "An asset, site or map photo beside facts (photo + fact strip + chart recipe)",
     "Decoration with no information", {}),
    ("map_pins", "d.photo(s, map_png, x, y, w, h); d.pin(s, px, py, [name, fact], label_box, hot=...)", "geography",
     ["market", "company"], (2, 6), True,
     "Sites, markets or a corridor on a map, each pinned with a name and one fact (target pin in gold)",
     "Places with no spatial story (use a table)", {"name": 4, "fact": 8}),
    ("bullets", "d.bullets(s, x, y, w, h, ['**Lead-in:** evidence and so what', ...])", "text", ["any"], (2, 6), False,
     "Numbered commentary beside an exhibit: bold lead-in + evidence + so what, 2-6 points",
     "The main exhibit of a page (pick a chart, table or diagram); more than 2 bullet blocks on one slide",
     {"lead": 4, "text": 25}),
    ("takeaway", "d.takeaway(s, x, y, w, h, '**So what:** ...')", "text", ["any"], (1, 1), False,
     "One framing sentence or 'so what' in a dashed outline (Why invest line, page conclusion)",
     "Lists or several points (use bullets / pill_rows)", {"text": 45}),
]

# Gallery page recipes live in layout_tags.json "gallery" (recipe: true); component thumbs map there too.


def build():
    sc = _load(SCAN, None)
    if sc is None:
        sys.exit("run `layout_tag.py scan` first")
    tags = _load(TAGS, {})
    igt, galt = tags.get("ig", {}), tags.get("gallery", {})
    recs = []
    for lid, variants in sorted(sc.items()):
        t = igt.get(lid, {})
        no = int(lid[2:])
        vs = [v for v in variants if not v.get("annotation")]
        recs.append({
            "id": lid, "source": "infographic", "library_no": no, "name": t.get("name", ""),
            "kind": t.get("kind", "layout"), "data_type": t.get("data_type", ""), "sections": t.get("sections", []),
            "use_when": t.get("use_when", ""), "dont_use_when": t.get("dont_use_when", ""),
            "items": t.get("items"), "numbers": t.get("numbers", False), "order": t.get("order", "rows"),
            "words": t.get("words", {}),
            "text_free": bool(vs) and not any(v["slots"] for v in vs),   # label beside it with d.text
            "variants": [{"variant": v["variant"], "items": t.get("variant_items", {}).get(str(v["variant"])),
                          "aspect": round(v["w"] / v["h"], 2) if v["h"] else None,
                          "table": v["has_table"], "chart": v["has_chart"], "counts": v["counts"],
                          # slot -> "sample text | chars the original held" (chars = rough capacity)
                          "slots": {x["slot"]: f"{x['sample'][:24]} | {x['chars']}" for x in v["slots"]}}
                         for v in vs],
            "call": f"d.layout(s, '{lid}', variant=1, box=(x, y, w, h), data={{'head': [...], 'text': [...]}})",
            "thumb": f"layout_thumbs/ig-{no:03d}.png", "exemplars": t.get("exemplars", []),
        })
    for c in COMPONENTS:
        cid, call, dt, secs, items, nums, use, dont, words = c
        recs.append({"id": cid, "source": "component", "name": cid.replace("_", " "), "kind": "layout",
                     "data_type": dt, "sections": secs, "use_when": use, "dont_use_when": dont,
                     "items": list(items), "numbers": nums, "words": words, "call": call,
                     "thumb": galt.get(cid, {}).get("thumb", ""), "exemplars": galt.get(cid, {}).get("exemplars", [])})
    for gid, g in sorted(galt.items()):
        if g.get("recipe"):
            recs.append({"id": gid, "source": "recipe", "name": g["name"], "kind": "page", "data_type": g["data_type"],
                         "sections": g["sections"], "use_when": g["use_when"], "dont_use_when": g["dont_use_when"],
                         "items": g.get("items"), "numbers": g.get("numbers", True), "words": g.get("words", {}),
                         "call": g["call"], "components": g.get("components", []), "thumb": g["thumb"],
                         "exemplars": g.get("exemplars", [])})
    # bank designs rebuilt from the real pages (bank_layouts.py; what each is for in bank_specs.py)
    from bank_specs import SPECS
    for bid, g in sorted(SPECS.items()):
        recs.append({"id": bid, "source": "bank", "name": g["name"], "group": g["group"], "kind": g["kind"],
                     "data_type": g["data_type"], "sections": g["sections"], "use_when": g["use_when"],
                     "dont_use_when": g["dont_use_when"], "look": g["look"], "items": g["items"],
                     "numbers": g["numbers"], "words": {},
                     "call": f"d.bank(s, '{bid}', box=(x, y, w, h), data={{...}})  # data shape: bank_samples.SAMPLE['{bid}']",
                     "thumb": f"layout_thumbs/bk-{bid[2:]}.png", "exemplars": [], "examples": g["examples"]})
    # real bank pages (deck_library/pages.py build): best 4 per layout, pages where it is the lead layout first;
    # bank designs use the pages they were drawn from
    pages = _load(PAGES_JSON, {}).get("pages", [])
    for r in recs:
        if r["source"] == "bank":
            hits = [p for p in pages if p["id"] in r["examples"]]
            hits.sort(key=lambda p: r["examples"].index(p["id"]))
        else:
            hits = [p for p in pages if r["id"] in p["layouts"]]
            hits.sort(key=lambda p: (p["layouts"].index(r["id"]) > 0, -p["fill"]))
        r["real_pages"] = len(hits)
        r["exemplars"] = r["exemplars"] + [
            {"deck": p["deck"], "page": p["page"], "bank": p["bank"], "file": p["file"], "title": p["title"],
             "note": p["note"], "thumb": p["thumb"]} for p in hits[:4]]
    # tuning-loop wording fixes override use_when / dont_use_when (kept only when the dev score did not drop)
    notes = _load(NOTES, {}).get("notes", {})
    for r in recs:
        for k in ("use_when", "dont_use_when"):
            if notes.get(r["id"], {}).get(k):
                r[k] = notes[r["id"]][k]
    lib = {"version": 1, "note": "Built by scripts/layout_tag.py build; edit layout_tags.json / COMPONENTS / "
           "eval/layout_notes.json, not this file",
           "sections": SECTIONS, "data_types": DATA_TYPES, "layouts": recs}
    _dump(LIB_JSON, lib)
    print(f"{len(recs)} records -> {os.path.abspath(LIB_JSON)}")


def validate(path=LIB_JSON):
    lib = _load(path, None)
    if lib is None:
        sys.exit("no layout_library.json; run build")
    errs, ids = [], set()
    for r in lib["layouts"]:
        rid = r.get("id", "?")
        if rid in ids:
            errs.append(f"{rid}: duplicate id")
        ids.add(rid)
        if r.get("kind") not in ("layout", "page", "icon_set", "map", "reference"):
            errs.append(f"{rid}: kind {r.get('kind')!r}")
        if r.get("kind") in ("layout", "page"):
            for k in ("name", "use_when", "dont_use_when", "data_type", "sections", "call"):
                if not r.get(k):
                    errs.append(f"{rid}: missing {k}")
            if r.get("data_type") and r["data_type"] not in lib["data_types"]:
                errs.append(f"{rid}: data_type {r['data_type']!r} not in the list")
            for sct in r.get("sections", []):
                if sct not in lib["sections"]:
                    errs.append(f"{rid}: section {sct!r} not in the list")
            it = r.get("items")
            if not (isinstance(it, list) and len(it) == 2 and all(isinstance(v, int) for v in it) and it[0] <= it[1]):
                errs.append(f"{rid}: items must be [min, max]")
            if r["source"] == "infographic" and not r.get("variants"):
                errs.append(f"{rid}: no variants")
        if r.get("thumb") and not os.path.exists(os.path.join(ASSETS, r["thumb"])):
            errs.append(f"{rid}: thumbnail {r['thumb']} missing")
        for e in r.get("exemplars", []):
            if not {"deck", "page"} <= set(e):
                errs.append(f"{rid}: exemplar needs deck + page")
    for e in errs:
        print("ERROR:", e)
    print(f"{len(lib['layouts'])} records, {len(errs)} errors")
    return not errs


def find(section=None, data_type=None, items=None):
    """Shortlist for a slide: layouts tagged with this section (and data type / item count), real-bank-proven first."""
    lib = _load(LIB_JSON, None) or sys.exit("no layout_library.json; run build")
    n = int(items) if items else None
    hits = [r for r in lib["layouts"] if r.get("kind") in ("layout", "page")
            and (not section or section in r["sections"] or "any" in r["sections"])
            and (not data_type or r["data_type"] == data_type)
            and (n is None or r["items"][0] <= n <= r["items"][1])]
    hits.sort(key=lambda r: (-r.get("real_pages", 0), r["source"] == "infographic", r["id"]))
    for r in hits:
        ex = r["exemplars"][0] if r["exemplars"] else None
        seen = f"  e.g. {ex['bank']} {ex['deck'][:40]} p{ex['page']}" if ex else ""
        print(f"{r['id']:<15} {r.get('real_pages', 0):>3} real | {r['name'][:30]:<30} | {r['use_when'][:80]}{seen}")
    print(f"{len(hits)} layouts")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "scan":
        _dump(SCAN, scan())
        print(f"scan -> {os.path.abspath(SCAN)}")
    elif cmd == "name":
        _dump(SCAN, scan(rename=True))
    elif cmd == "sheets":
        sheets(sys.argv[2] if len(sys.argv) > 2 else "layout_sheets", sys.argv[3] if len(sys.argv) > 3 else "render_ig")
    elif cmd == "thumbs":
        thumbs(sys.argv[2] if len(sys.argv) > 2 else None)
    elif cmd == "build":
        build()
    elif cmd == "validate":
        sys.exit(0 if validate() else 1)
    elif cmd == "find":
        find(*sys.argv[2:5])
    else:
        sys.exit(__doc__)
