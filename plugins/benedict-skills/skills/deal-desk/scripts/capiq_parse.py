"""Capital IQ Excel exports -> deal_data.json (feeds the model and deck scripts).

Recipe for pulling the exports: references/capiq-pull.md.

  python capiq_parse.py build <company_dir> [--accessed YYYY-MM-DD]
      reads <company_dir>/raw/*.xlsx (+ curation.json if present), writes <company_dir>/deal_data.json
  python capiq_parse.py check <company_dir>
      prints the headline figures to compare against the CapIQ screen, and the tie-out checks

Raw files (names fixed by the recipe): income_statement, balance_sheet, cash_flow, multiples,
peer_comps, precedents (.xlsx). Any missing file is skipped with a note.

Every money figure is in US$ m (CapIQ's $000 is divided by 1,000). Per-share figures stay in $,
share counts stay actual, margins in %, multiples in x. Each table carries "src"
("S&P Capital IQ, accessed <date> (<file>)"); every figure inside it takes that source.

curation.json also adds precedents by hand from filings (manual_precedents, each with its own src),
splits peers into segments (peer_segments) and blanks multiples that are not meaningful (peer_nm). The build
then writes the median/mean stats per segment and for the core precedents.
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

FILES = ["income_statement", "balance_sheet", "cash_flow", "multiples", "peer_comps", "precedents"]


def _num(v):
    if v is None or isinstance(v, str) and v.strip().upper() in ("NA", "NM", "", "-"):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return v


def _scale(v, unit):
    """Convert to US$ m where the unit is $000 or $M; leave everything else."""
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return v
    if unit == "000":
        return round(v / 1000, 4)
    return round(v, 6) if isinstance(v, float) else v


def _unit_of(label):
    lab = label.replace(" ", "")
    if "(%)" in lab:
        return "%"
    if "(x)" in lab:
        return "x"
    if "(actual)" in lab:
        return "actual"
    if "($000)" in lab:
        return "000"
    if "($M)" in lab or "($m)" in lab:
        return "m"
    if "($)" in lab:
        return "$"
    return None


def _clean(label):
    return re.sub(r"\s+", " ", re.sub(r"\n", " ", str(label))).strip()


def _src(accessed, f):
    return f"S&P Capital IQ, accessed {accessed} ({f.name})"


# ---------- company financial statements (Income Statement / Balance Sheet / Cash Flow) ----------
def parse_statement(f, accessed):
    ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    meta, hdr = {}, None
    for i, r in enumerate(rows):
        a = r[0]
        if isinstance(a, str) and ":" in a and hdr is None and i < 14 and r[1] is None:
            k, _, v = a.partition(":")
            meta[k.strip()] = v.strip()
        if isinstance(r[1], str) and re.match(r"\d{4} FY|LTM|\d{4} Q", r[1].strip()):
            hdr = i
            break
    if hdr is None:
        raise ValueError(f"{f.name}: no period header row found")
    cols = [j for j in range(1, len(rows[hdr])) if rows[hdr][j]]
    periods = [{"label": rows[hdr][j].strip()} for j in cols]
    lines, unit = {}, "000"  # CapIQ statements default to $000 (Magnitude row)
    if "Thousands" not in meta.get("Magnitude", "Thousands"):
        unit = "m" if "Million" in meta.get("Magnitude", "") else None
    base_unit = unit
    for r in rows[hdr + 1:]:
        a = r[0]
        if a is None or not str(a).strip():
            continue
        label = _clean(a)
        vals = [r[j] if j < len(r) else None for j in cols]
        if label in ("Period Ended", "Financial Filing Date"):
            key = "period_end" if label == "Period Ended" else "filing_date"
            for p, v in zip(periods, vals):
                p[key] = _num(v)
            continue
        if all(v is None for v in vals):  # section header, e.g. "Per Share Items ($)"
            unit = _unit_of(label) or base_unit
            continue
        u = _unit_of(label) or unit
        if u == "$" and not re.search(r"per share|EPS|price", label, re.I):
            u = base_unit  # CapIQ heads the BS supplemental block "($)" but its money lines are $000
        if label in lines:  # CapIQ repeats a few labels; keep the first, suffix the rest
            n = 2
            while f"{label} #{n}" in lines:
                n += 1
            label = f"{label} #{n}"
        lines[label] = [_scale(_num(v), u) for v in vals]
    title = _clean(rows[1][0]) if len(rows) > 1 and rows[1][0] else ""
    ident = _clean(rows[2][0]) if len(rows) > 2 and rows[2][0] else ""
    return {"src": _src(accessed, f), "unit": "US$ m unless the line says ($), (actual), (%) or (x)",
            "basis": meta.get("Reporting Basis", ""), "periods": periods, "lines": lines,
            "_title": title, "_ident": ident}


# ---------- multiples history ----------
def parse_multiples(f, accessed):
    ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = next(i for i, r in enumerate(rows) if r[0] and "Standard" in str(r[0]) and r[2])
    dates = [str(d).strip() for d in rows[hdr][2:] if d]
    out, metric = {}, None
    for r in rows[hdr + 1:]:
        if r[0] and str(r[0]).strip():
            metric = _clean(r[0])
        if metric and r[1] and str(r[1]).strip():
            out.setdefault(metric, {})[str(r[1]).strip().lower()] = dict(zip(dates, [_num(v) for v in r[2:2 + len(dates)]]))
    return {"src": _src(accessed, f), "unit": "x", "dates": dates, "series": out}


# ---------- screener exports (Companies peer comps, Transactions precedents) ----------
def _screener_sheet(ws):
    """Return (labels, codes, data_rows). Label row sits just above the field-code row."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    ci = next(i for i, r in enumerate(rows) if r[0] and re.fullmatch(r"[A-Z0-9_]+", str(r[0])))
    labels = [_clean(x) if x else "" for x in rows[ci - 1]]
    codes = [str(x) if x is not None else "" for x in rows[ci]]
    data, start = [], ci + 1
    if start < len(rows) and rows[start][0] is None and any(rows[start]):  # parameter row (e.g. "Announcement")
        start += 1
    for r in rows[start:]:
        if r[0] is None or not str(r[0]).strip():
            if data:
                break
            continue
        data.append(r)
    return labels, codes, data


def _field(label, code):
    unit = _unit_of(label)
    name = re.sub(r"\s*\((%|x|\$000|\$M|\$|actual|MM/dd/yyyy)\)\s*$", "", label.replace(" MM/dd/yyyy", "")).strip()
    return name, unit


def _ticker(name):
    m = re.search(r"\(([A-Z]+):([A-Z0-9.]+)\)\s*$", name)
    return (m.group(2), m.group(1)) if m else (None, None)


def _criteria(wb):
    if "Screening Criteria" not in wb.sheetnames:
        return []
    return [_clean(r[0]) for r in wb["Screening Criteria"].iter_rows(values_only=True) if r[0]][2:]


def parse_comps(f, accessed):
    wb = openpyxl.load_workbook(f, data_only=True)
    ents = {}
    for ws in wb.worksheets:
        if ws.title == "Screening Criteria":
            continue
        labels, codes, data = _screener_sheet(ws)
        for r in data:
            name = _clean(r[0])
            if name in ("Minimum", "Median", "Average", "Maximum"):
                continue
            e = ents.setdefault(name, {"name": name})
            for lab, code, v in zip(labels[1:], codes[1:], r[1:]):
                if not code or code == "SP_ENTITY_NAME":
                    continue
                fname, u = _field(lab, code)
                e[fname] = _scale(_num(v), u)
    out = []
    for e in ents.values():
        e["ticker"], e["exchange"] = _ticker(e["name"])
        e["name"] = re.sub(r"\s*\([A-Z]+:[A-Z0-9.]+\)\s*$", "", e["name"])
        out.append(e)
    return {"src": _src(accessed, f), "unit": "US$ m, $ per share, x, %; LTM unless named Forward (NTM)",
            "screen": _criteria(wb), "rows": out}


def parse_precedents(f, accessed):
    wb = openpyxl.load_workbook(f, data_only=True)
    labels, codes, data = _screener_sheet(wb.worksheets[0])
    rows, seen = [], set()
    for r in data:
        d = {}
        for lab, code, v in zip(labels, codes, r):
            fname, u = _field(lab, code)
            if fname in d and d[fname] == _scale(_num(v), u):  # CapIQ repeats Total Transaction Value
                continue
            if fname in d:
                fname += " (incl. estimated)"
            d[fname] = _scale(_num(v), u)
        if d.get("MI Transaction ID") in seen:
            continue
        seen.add(d.get("MI Transaction ID"))
        rows.append(d)
    return {"src": _src(accessed, f), "unit": "US$ m, multiples x (at announcement)",
            "screen": _criteria(wb), "rows": rows}


# ---------- curation (peer set, precedent set; written by hand per company) ----------
def apply_curation(data, cur):
    if "comps" in data:
        core = set(cur.get("peers_core", []))
        notes = cur.get("peer_notes", {})
        for e in data["comps"]["rows"]:
            t = e.get("ticker")
            e["role"] = "target" if t == cur.get("ticker") else ("core" if t in core else "excluded")
            if t in notes:
                e["note"] = notes[t]
    if "precedents" in data:
        core = set(cur.get("precedents_core", []))
        notes = cur.get("precedent_notes", {})
        for d in data["precedents"]["rows"]:
            i = d.get("MI Transaction ID")
            d["role"] = "core" if i in core else ("reference" if i in cur.get("precedents_reference", []) else "excluded")
            if i in notes:
                d["note"] = notes[i]
        for m in cur.get("manual_precedents", []):
            data["precedents"]["rows"].append(manual_row(m))
        data["precedents"]["not_in_screen"] = cur.get("precedents_missing", [])
        data["precedents"]["stats"] = {"core": stats([r for r in data["precedents"]["rows"] if r.get("role") == "core"], DEAL_MULTS)}
    if "comps" in data:
        for e in data["comps"]["rows"]:
            for k in cur.get("peer_nm", {}).get(e.get("ticker"), {}).get("fields", []):
                e[k] = None  # not meaningful (e.g. near-zero EBITDA); reason kept in nm_reason
                e["nm_reason"] = cur["peer_nm"][e["ticker"]]["why"]
        core = [e for e in data["comps"]["rows"] if e.get("role") == "core"]
        segs = {"All core": core}
        for name, seg in cur.get("peer_segments", {}).items():
            tick = seg["tickers"]
            for e in core:
                if e.get("ticker") in tick:
                    e["segment"] = name
            segs[name] = [e for e in core if e.get("ticker") in tick]
        data["comps"]["segments"] = {n: {"tickers": [e["ticker"] for e in rows],
                                         "why": cur.get("peer_segments", {}).get(n, {}).get("why", ""),
                                         **stats(rows, COMP_MULTS)} for n, rows in segs.items()}


COMP_MULTS = ["TEV/ EBITDA", "TEV/ Forward EBITDA", "TEV/ Total Revenue", "TEV/ Forward Total Revenue"]
DEAL_MULTS = ["Implied Enterprise Value/ EBITDA", "Implied Enterprise Value/ NTM EBITDA",
              "Implied Enterprise Value/ Total Revenue", "Implied Enterprise Value/ NTM Revenue"]


def stats(rows, fields):
    """n, median, mean, low, high per multiple; blanks and NM are left out."""
    out = {}
    for f in fields:
        v = sorted(r[f] for r in rows if isinstance(r.get(f), (int, float)))
        if v:
            n = len(v)
            med = v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2
            out[f] = {"n": n, "median": round(med, 2), "mean": round(sum(v) / n, 2),
                      "low": round(v[0], 2), "high": round(v[-1], 2)}
    return out


def manual_row(m):
    """A deal added by hand from filings (not in the Capital IQ screen). Multiples land in the LTM or NTM
    column per their stated basis; a 'less than' multiple is used at its cap and flagged."""
    ltm = m["ebitda_basis_type"] == "LTM"
    row = {"Target/Issuer Name": m["target"], "MI Transaction ID": m["id"], "Announced Date": m["announced"],
           "Transaction Type": "M&A", "Transaction Status": "Completed", "Total Transaction Value": m["ev"],
           "Buyers/Investors Name": m["buyer"], "Transaction Industry (MI)": m.get("industry"),
           "Implied Enterprise Value/ EBITDA": m["ev_ebitda"] if ltm else None,
           "Implied Enterprise Value/ NTM EBITDA": None if ltm else m["ev_ebitda"],
           "Implied Enterprise Value/ Total Revenue": m.get("ev_revenue") if m.get("revenue_basis_type") == "LTM" else None,
           "Implied Enterprise Value/ NTM Revenue": m.get("ev_revenue") if m.get("revenue_basis_type") == "NTM" else None,
           "role": "core", "note": m.get("note", ""), "source": "manual (filings)",
           "multiple_note": f"{m.get('ev_ebitda_qualifier', '')}{m['ev_ebitda']}x {m['ebitda_basis']}".strip(),
           "implied_ebitda": round(m["ev"] / m["ev_ebitda"], 1), "src": "; ".join(m["sources"])}
    return row


def build(cdir, accessed):
    raw = cdir / "raw"
    data = {"company": {}, "accessed": accessed,
            "licence": "S&P Capital IQ via SMU; personal and academic use only, not for redistribution",
            "missing": []}
    for name in FILES:
        f = raw / f"{name}.xlsx"
        if not f.exists():
            data["missing"].append(name)
            continue
        if name in ("income_statement", "balance_sheet", "cash_flow"):
            s = parse_statement(f, accessed)
            title, ident = s.pop("_title"), s.pop("_ident")
            if title and not data["company"]:
                m = re.search(r"MI KEY: (\d+); SPCIQ KEY: (\d+)", ident)
                data["company"] = {"name": title.split("|")[0].strip(), "listing": ident.split("(")[0].strip(),
                                   "mi_key": m.group(1) if m else None, "ciq_key": m.group(2) if m else None}
            data[name] = s
        elif name == "multiples":
            data[name] = parse_multiples(f, accessed)
        elif name == "peer_comps":
            data["comps"] = parse_comps(f, accessed)
        else:
            data["precedents"] = parse_precedents(f, accessed)
    cur_f = cdir / "curation.json"
    if cur_f.exists():
        cur = json.loads(cur_f.read_text(encoding="utf-8"))
        apply_curation(data, cur)
        data["curation"] = {"file": cur_f.name, "by": cur.get("by", ""), "note": cur.get("note", "")}
    out = cdir / "deal_data.json"
    out.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    return out, data


# ---------- check ----------
def _line(st, *names):
    for n in names:
        if n in st["lines"]:
            return st["lines"][n]
    raise KeyError(names[0])


def _print_stats(title, st):
    print(f"  {title}:")
    for f, v in st.items():
        if isinstance(v, dict) and "median" in v:
            print(f"    {f:40} n={v['n']}  median {v['median']:>6}x  mean {v['mean']:>6}x  range {v['low']}-{v['high']}x")


def check(cdir):
    d = json.loads((cdir / "deal_data.json").read_text(encoding="utf-8"))
    fails = []

    def tie(desc, a, b, tol=0.5):
        ok = a is not None and b is not None and abs(a - b) <= tol
        print(f"  {'OK  ' if ok else 'FAIL'} {desc}: {a} vs {b}")
        if not ok:
            fails.append(desc)

    print(f"{d['company'].get('name')} ({d['company'].get('listing')}) - accessed {d['accessed']}")
    if d["missing"]:
        print("  missing raw files:", ", ".join(d["missing"]))
    if "income_statement" in d:
        s = d["income_statement"]
        labs = [p["label"] for p in s["periods"]]
        rev, cogs, gp = _line(s, "Total Revenue"), _line(s, "Cost Of Goods Sold"), _line(s, "Gross Profit")
        ebitda = _line(s, "EBITDA")
        print("Income statement (US$ m):", " | ".join(labs))
        print("  Revenue ", rev)
        print("  EBITDA  ", ebitda)
        for i, l in enumerate(labs):
            tie(f"{l} revenue - COGS = gross profit", round(rev[i] - cogs[i], 3), gp[i], 0.01)
    if "balance_sheet" in d:
        b = d["balance_sheet"]
        ta, tle = _line(b, "Total Assets"), _line(b, "Total Liabilities And Equity")
        tl, te = _line(b, "Total Liabilities"), _line(b, "Total Equity")
        for i, p in enumerate(b["periods"]):
            tie(f"{p['label']} assets = liabilities + equity", ta[i], round(tl[i] + te[i], 3), 0.01)
            tie(f"{p['label']} balance sheet total", ta[i], tle[i], 0.01)
    if "cash_flow" in d:
        c = d["cash_flow"]
        cfo, cfi, cff = _line(c, "Cash from Ops."), _line(c, "Cash from Investing"), _line(c, "Cash from Financing")
        fx, net = _line(c, "Foreign Exchange Rate Adj."), _line(c, "Net Change in Cash")
        for i, p in enumerate(c["periods"]):
            parts = [x or 0 for x in (cfo[i], cfi[i], cff[i], fx[i])]
            tie(f"{p['label']} CFO + CFI + CFF + FX = change in cash", round(sum(parts), 3), net[i], 0.01)
    if "comps" in d:
        rows = d["comps"]["rows"]
        print("Trading comps (role · ticker · TEV US$ m · TEV/EBITDA LTM · NTM):")
        for e in sorted(rows, key=lambda e: (e.get("role", ""), e.get("ticker") or "")):
            print(f"  {e.get('role','-'):9} {e.get('ticker') or '-':6} {e.get('Total Enterprise Value (CIQ)')!s:>10}"
                  f" {e.get('TEV/ EBITDA')!s:>10} {e.get('TEV/ Forward EBITDA')!s:>10}")
    if "precedents" in d:
        rows = d["precedents"]["rows"]
        core = [r for r in rows if r.get("role") == "core"]
        print(f"Precedents: {len(rows)} rows, {len(core)} core ({sum(1 for r in core if r.get('source'))} added by hand)")
        for r in core:
            print(f"  {r['Announced Date'][:7]} {r['Target/Issuer Name'][:40]:40} {r['Total Transaction Value']:>9.0f}"
                  f"  EV/EBITDA LTM {r.get('Implied Enterprise Value/ EBITDA')} NTM {r.get('Implied Enterprise Value/ NTM EBITDA')}"
                  + (f"  [{r['multiple_note']}]" if r.get("multiple_note") else ""))
        _print_stats("Core precedents", d["precedents"].get("stats", {}).get("core", {}))
    for name, seg in d.get("comps", {}).get("segments", {}).items():
        _print_stats(f"Comps - {name} ({', '.join(seg['tickers'])})", seg)
    print("All ties OK" if not fails else f"{len(fails)} FAILED: " + "; ".join(fails))
    return not fails


if __name__ == "__main__":
    if len(sys.argv) < 3 or sys.argv[1] not in ("build", "check"):
        print(__doc__)
        sys.exit(1)
    cdir = Path(sys.argv[2])
    if sys.argv[1] == "build":
        acc = sys.argv[sys.argv.index("--accessed") + 1] if "--accessed" in sys.argv else dt.date.today().isoformat()
        out, _ = build(cdir, acc)
        print("wrote", out)
        sys.exit(0 if check(cdir) else 1)
    sys.exit(0 if check(cdir) else 1)
