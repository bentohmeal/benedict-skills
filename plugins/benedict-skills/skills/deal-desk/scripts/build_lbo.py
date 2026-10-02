"""Build a formula-driven PE buyout (LBO) model from a JSON case file, then verify it (house style v2).

    python build_lbo.py case.json out.xlsx            # build + data tables + recalc + independent check + lint
    python build_lbo.py case.json out.xlsx --no-verify

Covers the classic timed PE test (Keppel "Tough Love" type): 5-yr P&L, D&A schedule, NWC, FCFF / FCFE,
one amortising term loan, covenants (Net debt/EBITDA, DSCR), Sources & Uses, IRR / MoM, two real Excel
data tables with heat maps, WACC + DCF (perpetuity and exit-multiple terminal values) and a Checks sheet.
Sheets: Cover · Inputs · Deal · Model · Checks. Same column grid on every sheet: A markers · B labels · C units ·
D/E constants · F = Year 0 … Year n. See ../references/case-lbo.md for the JSON fields.
"""
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlstyle import (FIRST_PERIOD, MULT, NUM, NUM1, PCT, STYLE_VERSION, TOTAL, YEAR_FMT, check_cell, col, colour_key,  # noqa: E402
                     comment, end_marker, input_cell, label, na_cell, note, output_box, put, row_of, section,
                     sensitivity, setup, sheet_header, sources_uses, workbook_defaults)

HERE = os.path.dirname(os.path.abspath(__file__))
Y0 = FIRST_PERIOD  # column F = Year 0


def yc(t):
    return col(Y0 + t)


# ----------------------------------------------------------------------------------------------------------
def defaults(c):
    n = c["years"]
    c.setdefault("units", "$'000")
    c.setdefault("tags", {})
    d = c["deal"]
    d.setdefault("entry_multiple_override", None)
    d.setdefault("entry_cost_pct", 0.0)
    d.setdefault("exit_cost_pct", 0.0)
    d.setdefault("min_cash", 0.0)
    d.setdefault("distribute_fcfe", 1)
    ln = c["loan"]
    ln.setdefault("upfront_fee_pct", 0.0)
    ln.setdefault("interest_on_average", 0)
    if "amort_pct" not in ln:
        tenor = ln.get("tenor", n)
        ln["amort_pct"] = [1.0 / tenor if t <= tenor else 0.0 for t in range(1, n + 1)]
    c.setdefault("dep", {}).setdefault("same_year", 1)
    c.setdefault("dcf", {})
    c["dcf"].setdefault("terminal_growth", 0.02)
    c["dcf"].setdefault("mid_year", 0)
    # Sensitivity axes: base case exactly in the centre, symmetric steps (lessons.md #6)
    s = c.setdefault("sensitivity", {})
    h = c["historical"]
    ebitda0 = h["revenue"] - h["cogs"] - h["opex"]
    price = d["entry_multiple_override"] * ebitda0 if d.get("entry_multiple_override") else d["purchase_price"]
    base_entry = price / ebitda0
    steps = s.get("steps", {"entry": 0.25, "exit": 0.25, "gearing": 0.05, "rate": 0.01})
    s["entry"] = [base_entry + k * steps["entry"] for k in (-2, -1, 0, 1, 2)]
    s["exit"] = [d["exit_multiple"] + k * steps["exit"] for k in (-2, -1, 0, 1, 2)]
    s["gearing"] = [ln["amount"] / price + k * steps["gearing"] for k in (-2, -1, 0, 1, 2)]
    s["rate"] = [ln["rate"] + k * steps["rate"] for k in (-2, -1, 0, 1, 2)]
    s["base"] = {"entry": base_entry, "exit": d["exit_multiple"], "gearing": ln["amount"] / price, "rate": ln["rate"]}
    return c


# ----------------------------------------------------------------------------------------------------------
def build(c, path):
    n = c["years"]
    last = yc(n)
    LASTC = Y0 + n  # last period column index (bars run to here)
    U = c["units"]
    wb = Workbook()
    workbook_defaults(wb)
    cover = wb.active
    cover.title = "Cover"
    inp, deal, m, chk = (wb.create_sheet(s) for s in ("Inputs", "Deal", "Model", "Checks"))
    for ws in (cover, inp, deal, m, chk):
        setup(ws)
    tag = lambda k, default: (c["tags"].get(k, "") + " – " if c["tags"].get(k) else "") + default  # noqa: E731
    case = c.get("case_short", f"{c['company']} – LBO case study")
    fy0 = c.get("first_year")   # fiscal year of Year 0 (e.g. 2025): labels read FY2025, FY2026 ... (lessons #30)
    YL = lambda t: f"FY{fy0 + t}" if fy0 else f"Year {t}"  # noqa: E731
    timeline = [("Financial year", [fy0 or 0] + [f"={yc(t - 1)}5+1" for t in range(1, n + 1)], YEAR_FMT, False),
                ("Period", ["Actual"] + ["Forecast"] * n, "@", True)]

    # ---------------- Inputs ----------------
    r = sheet_header(inp, case, f"Inputs · {U} unless stated", LASTC, timeline)
    section(inp, r, f"{YL(0)} (LTM) P&L – costs entered as positive numbers", last_col=LASTC)
    h = c["historical"]
    I = {}
    r += 1
    for k, lab in [("revenue", "Revenue"), ("cogs", "Cost of sales"), ("opex", "Operating expenses"),
                   ("dep", "Depreciation"), ("interest", "Interest expense"), ("tax", "Income tax")]:
        label(inp, r, lab, U)
        input_cell(inp, f"D{r}", h[k if k != "dep" else "depreciation"])
        I[k] = f"Inputs!$D${r}"
        r += 1
    r += 1
    section(inp, r, "Operating drivers (by forecast year)", last_col=LASTC)
    r += 1
    drv = c["drivers"]
    for k, lab, u, f in [("revenue_growth", "Revenue growth", "%", PCT), ("gross_margin", "Gross margin (% of revenue)", "%", PCT),
                         ("opex_pct", "Operating expenses (% of revenue)", "%", PCT), ("capex", "Capex (new fixed assets)", U, NUM),
                         ("nwc_pct", "Net working capital (% of revenue)", "%", PCT)]:
        label(inp, r, lab, u)
        na_cell(inp, f"{yc(0)}{r}")
        for t in range(1, n + 1):
            input_cell(inp, f"{yc(t)}{r}", drv[k][t - 1], f)
        I[k] = r
        r += 1
    label(inp, r, "Loan amortisation (% of original loan)", "%")
    na_cell(inp, f"{yc(0)}{r}")
    for t in range(1, n + 1):
        input_cell(inp, f"{yc(t)}{r}", c["loan"]["amort_pct"][t - 1], PCT)
    I["amort"] = r
    r += 2
    section(inp, r, "Other assumptions", last_col=LASTC)
    r += 1
    for k, lab, u, v, f in [
            ("opening_nwc", f"{YL(0)} net working capital", U, c["opening"]["nwc"], NUM),
            ("book_equity", f"{YL(0)} shareholders' equity (info only)", U, c["opening"].get("equity_book", 0), NUM),
            ("existing_dep", "Depreciation of existing fixed assets (per year)", U, c["dep"]["existing_per_year"], NUM),
            ("life", "Useful life of new fixed assets", "years", c["dep"]["new_asset_life"], "0"),
            ("same_year", "New capex depreciated from year of spend", "(1,0)", c["dep"]["same_year"], "0"),
            ("tax_rate", "Tax rate", "%", c["tax_rate"], PCT),
            ("avg_int", "Interest on average (1) or opening (0) balance", "(1,0)", c["loan"]["interest_on_average"], "0"),
            ("min_cash", "Minimum cash kept in the business", U, c["deal"]["min_cash"], NUM),
            ("distribute", "Distribute excess cash to sponsor each year", "(1,0)", c["deal"]["distribute_fcfe"], "0"),
            ("ke", "Cost of equity", "%", c["dcf"]["cost_of_equity"], PCT),
            ("g", "Terminal growth rate (perpetuity method)", "%", c["dcf"]["terminal_growth"], PCT),
            ("mid", "Mid-year discounting (1) or year-end (0)", "(1,0)", c["dcf"]["mid_year"], "0")]:
        label(inp, r, lab, u)
        input_cell(inp, f"D{r}", v, f)
        I[k] = f"Inputs!$D${r}"
        if u == "(1,0)":
            comment(inp, f"E{r}", {"same_year": "1 = full-year depreciation in the year of spend, 0 = from the next year",
                                   "avg_int": "0 avoids a circular reference", "distribute": "1 = sweep all FCFE to the sponsor, 0 = keep cash to exit",
                                   "mid": "1 = mid-year convention"}[k])
        r += 1
    end_marker(inp, r + 1, LASTC)

    # ---------------- Deal ----------------
    r = sheet_header(deal, case, "Deal: assumptions, sources & uses, returns, sensitivities", LASTC)
    note(deal, f"B{r}", "Deal assumptions sit on this sheet because Excel data tables need their input cells on the same sheet.")
    r += 1
    section(deal, r, "Deal assumptions", last_col=LASTC)
    r += 1
    D = {}
    dl = c["deal"]
    for k, lab, u, v, f in [
            ("price", "Purchase price (enterprise value, cash-free / debt-free)", U, dl["purchase_price"], NUM),
            ("entry_ovr", "Entry EV / EBITDA override", "x", dl["entry_multiple_override"], MULT),
            ("price_used", "Purchase price used", U, None, NUM),
            ("entry_x", f"Implied entry EV / {YL(0)} EBITDA", "x", None, MULT),
            ("entry_cost", "Entry transaction costs (% of purchase price)", "%", dl["entry_cost_pct"], PCT),
            ("exit_x", "Exit EV / EBITDA", "x", dl["exit_multiple"], MULT),
            ("exit_cost", "Exit transaction costs (% of exit EV)", "%", dl["exit_cost_pct"], PCT),
            ("loan_base", f"{c['loan'].get('name', 'Term loan')} amount", U, c["loan"]["amount"], NUM),
            ("gear_ovr", "Gearing override (loan as % of purchase price)", "%", None, PCT),
            ("loan", "Loan used", U, None, NUM),
            ("gearing", "Gearing (loan / purchase price)", "%", None, PCT),
            ("rate", "All-in interest rate", "%", c["loan"]["rate"], PCT),
            ("fee", "Upfront fee (% of loan)", "%", c["loan"]["upfront_fee_pct"], PCT)]:
        label(deal, r, lab, u)
        D[k] = f"Deal!$D${r}"
        if v is not None or k in ("entry_ovr", "gear_ovr"):
            input_cell(deal, f"D{r}", v, f)
        r += 1
    ref = lambda k: D[k].split("!")[1]  # noqa: E731
    put(deal, ref("price_used"), f"=IF(ISNUMBER({ref('entry_ovr')}),{ref('entry_ovr')}*Model!$F$EBITDA,{ref('price')})")
    put(deal, ref("entry_x"), f"={ref('price_used')}/Model!$F$EBITDA", MULT)
    put(deal, ref("loan"), f"=IF(ISNUMBER({ref('gear_ovr')}),{ref('gear_ovr')}*{ref('price_used')},{ref('loan_base')})")
    put(deal, ref("gearing"), f"={ref('loan')}/{ref('price_used')}", PCT)
    comment(deal, ref("entry_ovr").replace("$D$", "E"), "leave blank = use the purchase price; the sensitivity table flexes this cell")
    comment(deal, ref("gear_ovr").replace("$D$", "E"), "leave blank = use the loan amount; the sensitivity table flexes this cell")
    r += 1

    # Sources & uses (stacked, bank layout)
    section(deal, r, tag("su", f"Sources and uses at close (end of {YL(0)})"), last_col=LASTC)
    r += 2
    loan_name = c["loan"].get("name", "Term loan")
    su, tot_s, tot_u, r = sources_uses(
        deal, r, U,
        sources=[(loan_name, f"={ref('loan')}"), ("Sponsor equity", "PLUG")],
        uses=[("Purchase price (100% stake)", f"={ref('price_used')}"),
              ("Entry transaction costs", f"={ref('price_used')}*{ref('entry_cost')}"),
              ("Upfront financing fee", f"={ref('loan')}*{ref('fee')}"),
              ("Cash to balance sheet", f"={I['min_cash']}")],
        ebitda_ref="Model!$F$EBITDA")
    comment(deal, f"G{su['Sponsor equity'][1:]}", "sponsor equity = total uses less debt (the plug)")
    D["equity"] = f"Deal!${su['Sponsor equity'][0]}${su['Sponsor equity'][1:]}"
    D["uses"] = f"Deal!{tot_u}"
    D["sources"] = f"Deal!{tot_s}"
    cb = su["Cash to balance sheet"]
    D["cash_bs"] = f"Deal!${cb[0]}${cb[1:]}"
    deal_ret_row = r

    # ---------------- Model ----------------
    r = sheet_header(m, case, f"Model: P&L, cash flow, debt, returns, valuation · {U}", LASTC, timeline, const_labels=("Value", ""))
    R = {}

    def sec(key, text):
        nonlocal r
        section(m, r, tag(key, text), last_col=LASTC)
        r += 1

    def line(key, lab, unit, y0, fn, fmt=NUM, bold=False, total=False, italic=False, grand=False):
        """y0: Year-0 formula/value or None; fn(t, colletter) -> formula for years 1..n."""
        nonlocal r
        label(m, r, lab, unit, bold=bold, italic=italic, indent=1 if italic else 0)
        R[key] = r
        vals = [y0] + [fn(t, yc(t)) for t in range(1, n + 1)]
        row_of(m, r, Y0, vals, fmt, bold=bold, total=total, italic=italic, grand=grand)
        r += 1

    pc = lambda t: yc(t - 1)  # noqa: E731
    F0 = yc(0)
    iv = lambda key, t: f"Inputs!{yc(t)}${I[key]}"  # noqa: E731

    sec("pl", "Income statement")
    line("rev", "Revenue", U, f"={I['revenue']}", lambda t, L: f"={pc(t)}{R['rev']}*(1+{iv('revenue_growth', t)})", bold=True)
    line("rev_g", "Growth", "%", None, lambda t, L: f"={L}{R['rev']}/{pc(t)}{R['rev']}-1", PCT, italic=True)
    line("cogs", "Cost of sales", U, f"=-{I['cogs']}", lambda t, L: f"=-{L}{R['rev']}*(1-{iv('gross_margin', t)})")
    line("gp", "Gross profit", U, f"={F0}{R['rev']}+{F0}{R['cogs']}", lambda t, L: f"={L}{R['rev']}+{L}{R['cogs']}", bold=True, total=True)
    line("gm", "Gross margin", "%", f"={F0}{R['gp']}/{F0}{R['rev']}", lambda t, L: f"={L}{R['gp']}/{L}{R['rev']}", PCT, italic=True)
    line("opex", "Operating expenses", U, f"=-{I['opex']}", lambda t, L: f"=-{L}{R['rev']}*{iv('opex_pct', t)}")
    line("ebitda", "EBITDA", U, f"={F0}{R['gp']}+{F0}{R['opex']}", lambda t, L: f"={L}{R['gp']}+{L}{R['opex']}", bold=True, total=True)
    line("ebitda_m", "EBITDA margin", "%", f"={F0}{R['ebitda']}/{F0}{R['rev']}", lambda t, L: f"={L}{R['ebitda']}/{L}{R['rev']}", PCT, italic=True)
    dep_row = r
    line("dep", "Depreciation", U, f"=-{I['dep']}", lambda t, L: "=0")
    line("ebit", "EBIT", U, f"={F0}{R['ebitda']}+{F0}{R['dep']}", lambda t, L: f"={L}{R['ebitda']}+{L}{R['dep']}", bold=True, total=True)
    int_row = r
    line("int", "Interest expense", U, f"=-{I['interest']}", lambda t, L: "=0")
    line("pbt", "Profit before tax", U, f"={F0}{R['ebit']}+{F0}{R['int']}", lambda t, L: f"={L}{R['ebit']}+{L}{R['int']}", bold=True, total=True)
    line("tax", "Income tax", U, f"=-{I['tax']}", lambda t, L: f"=-MAX(0,{L}{R['pbt']})*{I['tax_rate']}")
    line("npat", "Profit after tax", U, f"={F0}{R['pbt']}+{F0}{R['tax']}", lambda t, L: f"={L}{R['pbt']}+{L}{R['tax']}", bold=True, grand=True)
    r += 1

    sec("da", "Depreciation schedule")
    line("capex_in", "Capex (new fixed assets)", U, None, lambda t, L: f"={iv('capex', t)}")
    line("dep_old", "Depreciation – existing assets", U, None, lambda t, L: f"={I['existing_dep']}")
    line("dep_new", "Depreciation – new capex", U, None,
         lambda t, L: f"=(SUM(${yc(1)}{R['capex_in']}:{L}{R['capex_in']})-(1-{I['same_year']})*{L}{R['capex_in']})/{I['life']}")
    line("dep_tot", "Total depreciation", U, None, lambda t, L: f"={L}{R['dep_old']}+{L}{R['dep_new']}", bold=True, total=True)
    for t in range(1, n + 1):
        put(m, f"{yc(t)}{dep_row}", f"=-{yc(t)}{R['dep_tot']}")
    r += 1

    sec("nwc", "Net working capital")
    line("nwc", "Net working capital balance", U, f"={I['opening_nwc']}", lambda t, L: f"={L}{R['rev']}*{iv('nwc_pct', t)}")
    line("dnwc", "(Increase) / decrease in NWC", U, None, lambda t, L: f"=-({L}{R['nwc']}-{pc(t)}{R['nwc']})")
    r += 1

    sec("debt", f"Debt schedule – {loan_name}")
    line("d_open", "Opening balance", U, None, lambda t, L: f"={pc(t)}{r + 3}")
    line("d_draw", "Drawdown", U, f"={D['loan']}", lambda t, L: "=0")
    line("d_rep", "Scheduled repayment", U, None, lambda t, L: f"=-MIN({L}{R['d_open']},{D['loan']}*Inputs!{L}${I['amort']})")
    line("d_close", "Closing balance", U, f"={F0}{R['d_draw']}",
         lambda t, L: f"={L}{R['d_open']}+{L}{R['d_draw']}+{L}{R['d_rep']}", bold=True, total=True)
    line("d_int", "Interest expense", U, None,
         lambda t, L: f"={D['rate']}*IF({I['avg_int']}=1,AVERAGE({L}{R['d_open']},{L}{R['d_close']}),{L}{R['d_open']})")
    comment(m, f"{col(LASTC + 1)}{R['d_int']}", "on the opening balance by default, so no circular reference")
    for t in range(1, n + 1):
        put(m, f"{yc(t)}{int_row}", f"=-{yc(t)}{R['d_int']}")
    r += 1

    sec("cf", "Free cash flow to firm (FCFF) and to equity (FCFE)")
    line("f_ebit", "EBIT", U, None, lambda t, L: f"={L}{R['ebit']}")
    line("f_tax", "Less: tax on EBIT (unlevered)", U, None, lambda t, L: f"=-MAX(0,{L}{R['ebit']})*{I['tax_rate']}")
    line("f_nopat", "NOPAT", U, None, lambda t, L: f"={L}{R['f_ebit']}+{L}{R['f_tax']}", bold=True, total=True)
    line("f_da", "Add: depreciation", U, None, lambda t, L: f"=-{L}{R['dep']}")
    line("f_capex", "Less: capex", U, None, lambda t, L: f"=-{L}{R['capex_in']}")
    line("f_nwc", "Less: increase in NWC", U, None, lambda t, L: f"={L}{R['dnwc']}")
    line("fcff", "Free cash flow to firm (FCFF)", U, None,
         lambda t, L: f"=SUM({L}{R['f_nopat']}:{L}{R['f_nwc']})", bold=True, grand=True)
    comment(m, f"{col(LASTC + 1)}{R['fcff']}", "unlevered: used in the DCF")
    r += 1
    line("e_npat", "Profit after tax", U, None, lambda t, L: f"={L}{R['npat']}")
    line("e_da", "Add: depreciation", U, None, lambda t, L: f"=-{L}{R['dep']}")
    line("e_capex", "Less: capex", U, None, lambda t, L: f"=-{L}{R['capex_in']}")
    line("e_nwc", "Less: increase in NWC", U, None, lambda t, L: f"={L}{R['dnwc']}")
    line("e_rep", "Less: debt repayment", U, None, lambda t, L: f"={L}{R['d_rep']}")
    line("fcfe", "Free cash flow to equity (FCFE)", U, None,
         lambda t, L: f"=SUM({L}{R['e_npat']}:{L}{R['e_rep']})", bold=True, grand=True)
    comment(m, f"{col(LASTC + 1)}{R['fcfe']}", "levered: cash available to the sponsor")
    line("fcfe_chk", "Check: FCFF – after-tax interest – repayment = FCFE", U, None,
         lambda t, L: f"={L}{R['fcff']}+{L}{R['int']}+({L}{R['tax']}-{L}{R['f_tax']})+{L}{R['d_rep']}-{L}{R['fcfe']}",
         NUM1, italic=True)
    r += 1

    sec("cash", "Cash and distributions to sponsor")
    line("c_open", "Opening cash", U, None, lambda t, L: f"={pc(t)}{r + 3}")
    line("c_fcfe", "Add: FCFE", U, None, lambda t, L: f"={L}{R['fcfe']}")
    line("c_dist", "Less: distribution to sponsor", U, None,
         lambda t, L: f"=-IF({I['distribute']}=1,MAX(0,{L}{R['c_open']}+{L}{R['c_fcfe']}-{I['min_cash']}),0)")
    line("c_close", "Closing cash", U, f"={D['cash_bs']}",
         lambda t, L: f"=SUM({L}{R['c_open']}:{L}{R['c_dist']})", bold=True, total=True)
    r += 1

    sec("cov", "Financial covenants (annual, last twelve months)")
    line("nd", "Net debt (closing debt – closing cash)", U, f"={F0}{R['d_close']}-{F0}{R['c_close']}",
         lambda t, L: f"={L}{R['d_close']}-{L}{R['c_close']}")
    line("lev", "Net debt / EBITDA", "x", f"={F0}{R['nd']}/{F0}{R['ebitda']}", lambda t, L: f"={L}{R['nd']}/{L}{R['ebitda']}", MULT, bold=True)
    line("cfads", "CFADS (EBITDA – tax paid – capex – increase in NWC)", U, None,
         lambda t, L: f"={L}{R['ebitda']}+{L}{R['tax']}-{L}{R['capex_in']}+{L}{R['dnwc']}")
    line("ds", "Debt service (interest + scheduled repayment)", U, None,
         lambda t, L: f"={L}{R['d_int']}-{L}{R['d_rep']}")
    line("dscr", "DSCR (CFADS / debt service)", "x", None,
         lambda t, L: f'=IF({L}{R["ds"]}>0,{L}{R["cfads"]}/{L}{R["ds"]},"n.a.")', MULT, bold=True)
    r += 1

    sec("ret", f"Returns – exit at end of {YL(n)}")
    L = last
    rr = {}
    for k, lab, f, fmt in [
            ("x_ebitda", f"{YL(n)} EBITDA", f"={L}{R['ebitda']}", NUM),
            ("x_mult", "Exit EV / EBITDA", f"={D['exit_x']}", MULT),
            ("x_ev", "Exit enterprise value", None, NUM),
            ("x_cost", "Less: exit transaction costs", None, NUM),
            ("x_debt", "Less: debt outstanding at exit", f"=-{L}{R['d_close']}", NUM),
            ("x_cash", "Add: cash at exit", f"={L}{R['c_close']}", NUM),
            ("x_eq", "Equity value at exit", None, NUM)]:
        label(m, r, lab, "x" if fmt == MULT else U, bold=(k == "x_eq"))
        rr[k] = r
        if f:
            put(m, f"{L}{r}", f, fmt)
        r += 1
    put(m, f"{L}{rr['x_ev']}", f"={L}{rr['x_ebitda']}*{L}{rr['x_mult']}")
    put(m, f"{L}{rr['x_cost']}", f"=-{L}{rr['x_ev']}*{D['exit_cost']}")
    put(m, f"{L}{rr['x_eq']}", f"=SUM({L}{rr['x_ev']}:{L}{rr['x_cash']})", bold=True).border = TOTAL
    R.update(rr)
    r += 1
    line("eq_cf", "Sponsor cash flows (equity in, distributions, exit)", U, f"=-{D['equity']}",
         lambda t, L2: f"=-{L2}{R['c_dist']}" + (f"+{L2}{rr['x_eq']}" if t == n else ""), bold=True)
    R["irr"], R["mom"] = r + 1, r + 2
    label(m, r + 1, "Equity IRR", "%", bold=True)
    output_box(m, f"D{r + 1}", f"=IRR({F0}{R['eq_cf']}:{L}{R['eq_cf']})", PCT)
    label(m, r + 2, "Equity multiple (MoM)", "x", bold=True)
    output_box(m, f"D{r + 2}", f"=SUMIF({yc(1)}{R['eq_cf']}:{L}{R['eq_cf']},\">0\")/-{F0}{R['eq_cf']}", MULT)
    r += 4

    sec("wacc", "WACC")
    wr = {}
    for k, lab, f, fmt in [
            ("ke", "Cost of equity", f"={I['ke']}", PCT),
            ("kd", "Pre-tax cost of debt (all-in loan rate)", f"={D['rate']}", PCT),
            ("t", "Tax rate", f"={I['tax_rate']}", PCT),
            ("kd_at", "After-tax cost of debt", None, PCT),
            ("e", "Equity (sponsor equity at close)", f"={D['equity']}", NUM),
            ("d", "Debt (loan at close)", f"={D['loan']}", NUM),
            ("we", "Equity weight  E / (D + E)", None, PCT),
            ("wd", "Debt weight  D / (D + E)", None, PCT),
            ("wacc", "WACC", None, PCT)]:
        label(m, r, lab, "%" if fmt == PCT else U, bold=(k == "wacc"))
        wr[k] = r
        if f:
            put(m, f"D{r}", f, fmt)
        r += 1
    put(m, f"D{wr['kd_at']}", f"=D{wr['kd']}*(1-D{wr['t']})", PCT)
    put(m, f"D{wr['we']}", f"=D{wr['e']}/(D{wr['e']}+D{wr['d']})", PCT)
    put(m, f"D{wr['wd']}", f"=D{wr['d']}/(D{wr['e']}+D{wr['d']})", PCT)
    output_box(m, f"D{wr['wacc']}", f"=D{wr['we']}*D{wr['ke']}+D{wr['wd']}*D{wr['kd_at']}", PCT)
    comment(m, f"E{wr['e']}", "weights from the deal's capital structure at close")
    R.update({"w_" + k: v for k, v in wr.items()})
    W = f"$D${wr['wacc']}"
    r += 1

    sec("dcf", f"DCF valuation (unlevered FCFF, valuation date = end of {YL(0)})")
    line("t", "Discount period", "years", None, lambda t, L2: f"={L2}$5-${yc(0)}$5-0.5*{I['mid']}", NUM1)
    line("df", "Discount factor", "x", None, lambda t, L2: f"=1/(1+{W})^{L2}{R['t']}", '0.0000')
    line("pv", "PV of FCFF", U, None, lambda t, L2: f"={L2}{R['fcff']}*{L2}{R['df']}")
    r += 1
    vr = {}
    for k, lab, f, fmt in [
            ("sum_pv", "Sum of PV of FCFF", f"=SUM({yc(1)}{R['pv']}:{L}{R['pv']})", NUM),
            ("tv_g", "Terminal value – perpetuity growth", f"={L}{R['fcff']}*(1+{I['g']})/({W}-{I['g']})", NUM),
            ("tv_m", "Terminal value – exit multiple", f"={L}{R['ebitda']}*{D['exit_x']}", NUM),
            ("pv_tv_g", "PV of terminal value (perpetuity)", None, NUM),
            ("pv_tv_m", "PV of terminal value (exit multiple)", None, NUM),
            ("ev_g", "Enterprise value – perpetuity method", None, NUM),
            ("ev_m", "Enterprise value – exit multiple method", None, NUM),
            ("imp_x", "Implied exit multiple from perpetuity method", None, MULT),
            ("imp_g", "Implied growth from exit multiple method", None, PCT),
            ("ev_vs_price", "Perpetuity EV vs purchase price", None, PCT)]:
        label(m, r, lab, "x" if fmt == MULT else ("%" if fmt == PCT else U), bold=k.startswith("ev_"))
        vr[k] = r
        if f:
            put(m, f"D{r}", f, fmt)
        r += 1
    tN = f"{L}{R['t']}"
    put(m, f"D{vr['pv_tv_g']}", f"=D{vr['tv_g']}/(1+{W})^({tN}+0.5*{I['mid']})")
    put(m, f"D{vr['pv_tv_m']}", f"=D{vr['tv_m']}/(1+{W})^({tN}+0.5*{I['mid']})")
    output_box(m, f"D{vr['ev_g']}", f"=D{vr['sum_pv']}+D{vr['pv_tv_g']}", NUM)
    output_box(m, f"D{vr['ev_m']}", f"=D{vr['sum_pv']}+D{vr['pv_tv_m']}", NUM)
    put(m, f"D{vr['imp_x']}", f"=D{vr['tv_g']}/{L}{R['ebitda']}", MULT)
    put(m, f"D{vr['imp_g']}", f"=(D{vr['tv_m']}*{W}-{L}{R['fcff']})/(D{vr['tv_m']}+{L}{R['fcff']})", PCT)
    put(m, f"D{vr['ev_vs_price']}", f"=D{vr['ev_g']}/{D['price_used']}-1", PCT)
    comment(m, f"E{vr['imp_x']}", "sanity check: compare with the entry and exit multiples")
    comment(m, f"E{vr['tv_g']}", f"{YL(n)} FCFF × (1 + g) / (WACC – g)")
    comment(m, f"E{vr['tv_m']}", f"{YL(n)} EBITDA × exit EV / EBITDA")
    R.update({"v_" + k: v for k, v in vr.items()})
    r += 1
    if c.get("alt_method"):
        sec("alt", "Alternative valuation method")
        for i, line_txt in enumerate(c["alt_method"]):
            label(m, r, f"({i + 1}) {line_txt}")
            r += 1
        r += 1
    end_marker(m, r, LASTC)

    # replace the EBITDA placeholder used on Deal
    for row in deal.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "$F$EBITDA" in cell.value:
                cell.value = cell.value.replace("$F$EBITDA", f"$F${R['ebitda']}")

    # ---------------- Deal: returns + data tables ----------------
    r = deal_ret_row
    section(deal, r, tag("ret", "Returns"), last_col=LASTC)
    r += 1
    for i, (lab, u, f, fmt) in enumerate([("Equity IRR", "%", f"=Model!$D${R['irr']}", PCT),
                                          ("Equity multiple (MoM)", "x", f"=Model!$D${R['mom']}", MULT),
                                          ("Exit equity value", U, f"=Model!${last}${R['x_eq']}", NUM),
                                          ("DCF enterprise value (perpetuity)", U, f"=Model!$D${R['v_ev_g']}", NUM)]):
        label(deal, r + i, lab, u, bold=(i < 2))
        (output_box if i < 2 else put)(deal, f"D{r + i}", f, fmt)
    D["irr"] = f"Deal!$D${r}"
    r += 5
    tables = []
    s = c["sensitivity"]
    b = s["base"]
    names = {}
    for key, text, top_vals, side_vals, top_in, side_in, top_fmt, side_fmt, top_t, side_t, bt, bs in [
            ("sens1", "Equity IRR sensitivity – purchase price multiple vs exit multiple", s["entry"], s["exit"],
             ref("entry_ovr"), ref("exit_x"), MULT, MULT, f"Purchase price multiple (EV / {YL(0)} EBITDA)", "Exit multiple",
             b["entry"], b["exit"]),
            ("sens2", "Equity IRR sensitivity – gearing vs interest rate", s["gearing"], s["rate"],
             ref("gear_ovr"), ref("rate"), PCT, PCT, "Gearing (loan / purchase price)", "Interest rate",
             b["gearing"], b["rate"])]:
        section(deal, r, tag(key, text), last_col=LASTC)
        top = r + 2
        rng, _base = sensitivity(deal, top, 4, f"={ref('irr')}", side_vals, top_vals, side_t, top_t,
                                 side_fmt, top_fmt, PCT, base_side=bs, base_top=bt, better="high")
        tables.append((f"Deal!{rng}", f"Deal!{top_in}", f"Deal!{side_in}"))
        names[key] = f"Deal!${rng.split(':')[0][0]}${rng.split(':')[0][1:]}:${rng.split(':')[1][0]}${rng.split(':')[1][1:]}"
        r = top + len(side_vals) + 2
        note(deal, f"D{r}", "Excel data table (What-If Analysis); base case boxed in the centre; green = higher IRR.")
        r += 2
    end_marker(deal, r, LASTC)

    # ---------------- Checks ----------------
    r = sheet_header(chk, case, "Checks: every line must show OK (0); the Cover shows the total", LASTC, const_labels=("Value", ""))
    section(chk, r, "Integrity checks", last_col=LASTC)
    r += 1
    checks = [
        ("Sources = uses", f"={D['sources']}-{D['uses']}"),
        ("Debt never negative", f"=-MIN(0,MIN(Model!{yc(0)}{R['d_close']}:{last}{R['d_close']}))"),
        ("Cash never below minimum", f"=-MIN(0,MIN(Model!{yc(1)}{R['c_close']}:{last}{R['c_close']})-{I['min_cash']})"),
        ("FCFE reconciles to FCFF", f"=SUMPRODUCT(ABS(Model!{yc(1)}{R['fcfe_chk']}:{last}{R['fcfe_chk']}))"),
        ("NPV of sponsor cash flows at IRR = 0",
         f"=NPV(Model!$D${R['irr']},Model!{yc(1)}{R['eq_cf']}:{last}{R['eq_cf']})+Model!{yc(0)}{R['eq_cf']}"),
        ("Asset life ≥ forecast years (depreciation formula valid)", f"=IF({I['life']}>=Model!{last}$5-Model!${yc(0)}$5,0,1)"),
        ("WACC weights sum to 100%", f"=Model!$D${R['w_we']}+Model!$D${R['w_wd']}-1"),
    ]
    first = r
    for lab, f in checks:
        label(chk, r, lab)
        check_cell(chk, f"D{r}", f)
        r += 1
    label(chk, r, "Total (sum of absolute values)", bold=True)
    tot = f"D{r}"
    check_cell(chk, tot, f"=SUMPRODUCT(ABS(D{first}:D{r - 1}))")
    chk[tot].border = TOTAL
    end_marker(chk, r + 2, LASTC)

    # ---------------- Cover ----------------
    r = sheet_header(cover, c.get("case", c["company"]), f"{c['company']} · {U} · deal-desk house style v{STYLE_VERSION}", LASTC,
                     units_label=False)
    section(cover, r, "Contents", last_col=LASTC)
    r += 1
    for a, bb in [("Inputs", f"{YL(0)} P&L, operating drivers, other assumptions"),
                  ("Deal", "Deal assumptions, sources & uses, returns, IRR sensitivities (data tables)"),
                  ("Model", "P&L, depreciation, NWC, debt, FCFF / FCFE, covenants, returns, WACC, DCF"),
                  ("Checks", "Integrity checks (all must be 0)")]:
        label(cover, r, a, bold=True)
        cover[f"D{r}"] = bb
        r += 1
    r += 1
    section(cover, r, "Colour code", last_col=LASTC)
    r = colour_key(cover, r + 1) + 1
    section(cover, r, "Key outputs", last_col=LASTC)
    r += 1
    for lab, u, f, fmt in [("Equity IRR", "%", f"=Model!$D${R['irr']}", PCT), ("Equity multiple (MoM)", "x", f"=Model!$D${R['mom']}", MULT),
                           ("DCF EV (perpetuity)", U, f"=Model!$D${R['v_ev_g']}", NUM),
                           ("DCF EV (exit multiple)", U, f"=Model!$D${R['v_ev_m']}", NUM),
                           ("Checks (0 = all OK)", "", f"=Checks!${tot[0]}${tot[1:]}", '[>=0.0005]0.000;[<=-0.0005](0.000);"All OK"')]:
        label(cover, r, lab, u)
        output_box(cover, f"D{r}", f, fmt)
        r += 1
    if c.get("assumptions_log"):
        r += 1
        section(cover, r, "Assumptions made (the brief was silent or ambiguous)", last_col=LASTC)
        r += 1
        for i, a in enumerate(c["assumptions_log"]):
            cover[f"B{r}"] = f"{i + 1}. {a}"
            r += 1
    end_marker(cover, r + 1, LASTC)

    # named ranges (slides and graders read these instead of fixed addresses)
    for nm, rg in [("Sens_Entry_Exit", names["sens1"]), ("Sens_Gearing_Rate", names["sens2"])]:
        wb.defined_names[nm] = DefinedName(nm, attr_text=rg)
    wb.active = 0
    wb.save(path)
    return R, tables, tot, D


# ----------------------------------------------------------------------------------------------------------
def irr(flows):
    lo, hi = -0.99, 10.0
    f = lambda r: sum(cf / (1 + r) ** i for i, cf in enumerate(flows))  # noqa: E731
    for _ in range(300):
        mid = (lo + hi) / 2
        if f(lo) * f(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def mirror(c):
    """Independent plain-Python calculation of the same case (no spreadsheet) for cross-checking."""
    n, h, d, ln, dr = c["years"], c["historical"], c["deal"], c["loan"], c["drivers"]
    tax = c["tax_rate"]
    ebitda0 = h["revenue"] - h["cogs"] - h["opex"]
    price = d["entry_multiple_override"] * ebitda0 if d.get("entry_multiple_override") else d["purchase_price"]
    loan = ln["amount"]
    uses = price * (1 + d["entry_cost_pct"]) + loan * ln["upfront_fee_pct"] + d["min_cash"]
    equity = uses - loan
    rev, nwc_prev, bal, cash = h["revenue"], c["opening"]["nwc"], loan, d["min_cash"]
    capex_cum, flows, fcffs = 0.0, [-equity], []
    for t in range(1, n + 1):
        rev *= 1 + dr["revenue_growth"][t - 1]
        ebitda = rev * dr["gross_margin"][t - 1] - rev * dr["opex_pct"][t - 1]
        capex = dr["capex"][t - 1]
        capex_cum += capex
        dep = c["dep"]["existing_per_year"] + (capex_cum - (1 - c["dep"]["same_year"]) * capex) / c["dep"]["new_asset_life"]
        ebit = ebitda - dep
        rep = min(bal, loan * ln["amort_pct"][t - 1])
        close = bal - rep
        interest = ln["rate"] * ((bal + close) / 2 if ln["interest_on_average"] else bal)
        pbt = ebit - interest
        npat = pbt - max(0, pbt) * tax
        nwc = rev * dr["nwc_pct"][t - 1]
        dnwc = nwc - nwc_prev
        fcff = ebit - max(0, ebit) * tax + dep - capex - dnwc
        fcffs.append(fcff)
        fcfe = npat + dep - capex - dnwc - rep
        dist = max(0, cash + fcfe - d["min_cash"]) if d["distribute_fcfe"] else 0
        cash = cash + fcfe - dist
        bal, nwc_prev = close, nwc
        flows.append(dist)
    ev_exit = ebitda * d["exit_multiple"]
    flows[-1] += ev_exit * (1 - d["exit_cost_pct"]) - bal + cash
    ke, kd = c["dcf"]["cost_of_equity"], ln["rate"] * (1 - tax)
    wacc = equity / (equity + loan) * ke + loan / (equity + loan) * kd
    mid = c["dcf"]["mid_year"]
    pv = sum(f / (1 + wacc) ** (t + 1 - 0.5 * mid) for t, f in enumerate(fcffs))
    g = c["dcf"]["terminal_growth"]
    tv = fcffs[-1] * (1 + g) / (wacc - g)
    ev = pv + tv / (1 + wacc) ** n
    return {"irr": irr(flows), "mom": sum(f for f in flows[1:] if f > 0) / equity, "equity": equity, "wacc": wacc, "ev_g": ev}


def verify(c, path, R, tables, tot, D):
    py = sys.executable
    office = os.path.join(HERE, "office.py")
    for rng, row_in, col_in in tables:
        subprocess.run([py, office, "datatable", path, rng, row_in, col_in], check=True)
    res = subprocess.run([py, office, "recalc", path], capture_output=True, text=True)
    print(res.stdout)
    eq = D["equity"].replace("$", "")
    refs = [f"Model!D{R['irr']}", f"Model!D{R['mom']}", eq, f"Model!D{R['w_wacc']}", f"Model!D{R['v_ev_g']}", f"Checks!{tot}"]
    out = subprocess.run([py, office, "values", path, *refs], capture_output=True, text=True, check=True)
    xl = json.loads(out.stdout)
    mr = mirror(c)
    pairs = [("IRR", xl[refs[0]], mr["irr"]), ("MoM", xl[refs[1]], mr["mom"]), ("Sponsor equity", xl[refs[2]], mr["equity"]),
             ("WACC", xl[refs[3]], mr["wacc"]), ("DCF EV (perpetuity)", xl[refs[4]], mr["ev_g"])]
    ok = True
    print(f"{'Output':<22}{'Excel':>16}{'Python':>16}   match")
    for name, a, b in pairs:
        good = abs(a - b) <= 1e-6 * max(1, abs(b))
        ok &= good
        print(f"{name:<22}{a:>16,.6f}{b:>16,.6f}   {'OK' if good else 'MISMATCH'}")
    chk = xl[refs[5]]
    print(f"Checks total: {chk}  -> {'all OK' if abs(chk) < 1e-3 else 'FAILED'}")
    lint = subprocess.run([py, os.path.join(HERE, "model_lint.py"), path], capture_output=True, text=True)
    print(lint.stdout)
    return ok and abs(chk) < 1e-3 and "success" in res.stdout and lint.returncode == 0


if __name__ == "__main__":
    cfg = defaults(json.load(open(sys.argv[1], encoding="utf-8")))
    out_path = os.path.abspath(sys.argv[2])
    Rows, Tables, Tot, Dd = build(cfg, out_path)
    print(f"built {out_path}")
    if sys.platform != "win32" and "--no-verify" not in sys.argv:
        print("MAC BUILD: workbook written. The data tables, Excel recalc and independent check need Excel, so they are "
              "skipped here (sensitivity tables stay empty, cached values missing). Finish on the PC: run this same command "
              "there; it rebuilds identically from the case file, then verifies.")
        sys.argv.append("--no-verify")
    if "--no-verify" not in sys.argv:
        sys.exit(0 if verify(cfg, out_path, Rows, Tables, Tot, Dd) else 1)
