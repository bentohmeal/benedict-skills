"""House style v2 for openpyxl models: classic IB navy/grey, bank layout (see ../references/house-style.md, lessons.md).

Colour code (font), always:  blue 0000FF = hardcode · green 008000 = link to another sheet · black = formula
                             red FF0000 = link to another workbook · fill FFFFCC (+ hair box) = assumption / input cell
Layout (every sheet):        navy header band (rows 1..) with case, sheet, timeline rows · A = "x" section markers ·
                             B labels · C units · D/E constants ("Input") · F onwards = periods · comments "<< ..." right.
Import with:
    import sys; sys.path.insert(0, "<Claude Things>/.claude/skills/deal-desk/scripts"); from xlstyle import *
"""
import re

from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string

STYLE_VERSION = "2.2"
# Timeline label: a real fiscal year shows as "FY2025"; 0 / negatives fall back to "Year 0" / "Year -2" (lessons #30)
YEAR_FMT = '[>=1900]"FY"0;[<0]"Year -"0;"Year "0'

NAVY, STEEL, GREY, BAR, SUBBAR, NAFILL = "1F3864", "8497B0", "808080", "D9D9D9", "DCE6F1", "BFBFBF"
BLUE, BLACK, GREEN, RED, WHITE = "0000FF", "000000", "008000", "FF0000", "FFFFFF"
INPUT_FILL = "FFFFCC"
HEAT_RED, HEAT_MID, HEAT_GREEN = "F8696B", "FFFFFF", "63BE7B"
FONT = "Arial"
SIZE = 10

NUM = '#,##0_);(#,##0);"–"_)'
NUM1 = '#,##0.0_);(#,##0.0);"–"_)'
NUM2 = '#,##0.00_);(#,##0.00);"–"_)'
PCT = '0.0%_);(0.0%);"–"_)'
PCT0 = '0%_);(0%);"–"_)'
MULT = '0.00"x"_);(0.00"x");"–"_)'
MULT1 = '0.0"x"_);(0.0"x");"–"_)'
YEAR = '0'
DATE = 'mmm-yy'
DATE_D = 'dd-mmm-yy'
FLAG = '0;-0;"–"'
CHECK = '[>=0.0005]0.000;[<=-0.0005](0.000);"OK"'  # |x| < 0.0005 shows OK
TEXT = "@"

FIRST_PERIOD = 6  # column F

hair = Side(style="hair", color="808080")
thin = Side(style="thin", color=BLACK)
thin_grey = Side(style="thin", color="A6A6A6")
double = Side(style="double", color=BLACK)
BOX_HAIR = Border(left=hair, right=hair, top=hair, bottom=hair)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)
TOP = Border(top=thin)
TOTAL = Border(top=thin)
GRAND = Border(top=thin, bottom=double)


def col(n):
    return get_column_letter(n)


def cidx(letter):
    return column_index_from_string(letter)


def _font(color=BLACK, bold=False, italic=False, size=SIZE, underline=None):
    return Font(name=FONT, size=size, color=color, bold=bold, italic=italic, underline=underline)


def _fill(rgb):
    return PatternFill("solid", fgColor=rgb)


# ---------------------------------------------------------------------------------------------------- workbook / sheet
def workbook_defaults(wb):
    """Arial 10 as the Normal style, so empty cells and Excel-filled data-table cells match."""
    from openpyxl.utils.indexed_list import IndexedList
    normal = wb._named_styles["Normal"]
    normal.font = Font(name=FONT, size=SIZE)
    wb._fonts = IndexedList([Font(name=FONT, size=SIZE)])  # font 0 = the default for every unstyled cell


def setup(ws, label_w=46, period_w=11, n_periods=40, first_period=FIRST_PERIOD, zoom=85, widths=None):
    """Gridlines off; A = markers (3), B = labels, C = units (9), D/E = constants (11), periods from F."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = label_w
    ws.column_dimensions["C"].width = 9
    for c in range(4, first_period):
        ws.column_dimensions[col(c)].width = 11
    for c in range(first_period, first_period + n_periods):
        ws.column_dimensions[col(c)].width = period_w
    for k, w in (widths or {}).items():
        ws.column_dimensions[k].width = w
    ps = ws.page_setup
    ps.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ps.fitToWidth, ps.fitToHeight = 1, 0


def sheet_header(ws, case, desc, last_col, rows=None, units_label=True, const_labels=("Input", ""),
                 first_period=FIRST_PERIOD, freeze=True):
    """Navy band at the top of every sheet (bank style, no big title).

    rows: timeline rows inside the band, each (label, [values/formulas from first_period], number_format, italic).
          e.g. [("Financial year", [0, "=F5+1", ...], '"Year "0', False), ("Period", ["Actual", "Forecast", ...], "@", True)]
    Returns the first free row below the band (band + 1 blank row).
    """
    rows = rows or []
    band_end = 4 + max(len(rows), 1) + 1  # rows 1..band_end
    for r in range(1, band_end + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.fill = _fill(NAVY)
            cell.font = _font(WHITE)
    ws.cell(2, 2, case).font = _font(WHITE, bold=True, size=11)
    ws.cell(3, 2, desc).font = _font(WHITE)
    r = 5
    if units_label:
        ws.cell(r, 3, "Units").font = _font(WHITE, bold=True)
        ws.cell(r, 3).alignment = Alignment(horizontal="center")
        for i, t in enumerate(const_labels):
            if t:
                ws.cell(r, 4 + i, t).font = _font(WHITE, bold=True)
                ws.cell(r, 4 + i).alignment = Alignment(horizontal="right")
    for lab, vals, fmt, italic in rows:
        ws.cell(r, 2, lab).font = _font(WHITE, italic=italic)
        for i, v in enumerate(vals):
            cell = ws.cell(r, first_period + i, v)
            cell.font = _font(WHITE, bold=not italic, italic=italic)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
        r += 1
    ws.row_dimensions[band_end].height = 6
    if freeze:
        ws.freeze_panes = f"{col(first_period)}{band_end + 1}" if rows else f"A{band_end + 1}"
    return band_end + 2


def title(ws, text, sub=None, last_col=12):
    """Compatibility: a header band without a timeline."""
    return sheet_header(ws, text, sub or "", last_col)


def _bar(ws, row, text, last_col, fill, font_color, marker=True, first_col=2):
    if marker:
        m = ws.cell(row, 1, "x")
        m.font = _font(GREY)
        m.alignment = Alignment(horizontal="right")
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row, c)
        cell.fill = _fill(fill)
        cell.font = _font(font_color, bold=True)
    ws.cell(row, first_col, text)


def part(ws, row, text, last_col):
    """Main part of the case ('Part I – LBO'): navy bar, white bold."""
    _bar(ws, row, text, last_col, NAVY, WHITE)


def section(ws, row, text, first_col=2, last_col=12):
    """Section: light-grey bar, bold black, 'x' marker in column A (Ctrl+down jumps between sections)."""
    _bar(ws, row, text, last_col, BAR, BLACK, first_col=first_col)


def subsection(ws, row, text, last_col):
    _bar(ws, row, text, last_col, SUBBAR, BLACK, marker=False)


def end_marker(ws, row, last_col):
    section(ws, row, "End", last_col=last_col)


def header_row(ws, row, labels, first_col, bold=True):
    """Column headers inside the body (e.g. a small table), right aligned with a bottom rule."""
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=first_col + i, value=lab)
        c.font = _font(BLACK, bold=bold)
        c.alignment = Alignment(horizontal="right")
        c.border = BOTTOM


def label(ws, row, text, unit=None, bold=False, indent=0, italic=False):
    c = ws.cell(row=row, column=2, value=text)
    c.font = _font(BLACK, bold=bold, italic=italic)
    c.alignment = Alignment(indent=indent)
    if unit is not None:
        u = ws.cell(row=row, column=3, value=unit)
        u.font = _font(GREY, italic=True)
        u.alignment = Alignment(horizontal="center")


# ---------------------------------------------------------------------------------------------------- cells
_LINK = re.compile(r"^=-?(?:'[^']+'|[A-Za-z0-9_.]+)!\$?[A-Z]{1,3}\$?\d+$")


def _kind(value):
    """input (typed number/text) · link (pure link to another sheet) · external (other workbook) · calc (formula)."""
    if isinstance(value, str) and value.startswith("="):
        v = value.replace(" ", "")
        if "[" in v and "]" in v:
            return "external"
        return "link" if _LINK.match(v) else "calc"
    return "input"


COLOURS = {"input": BLUE, "calc": BLACK, "link": GREEN, "external": RED}


def put(ws, ref, value, fmt=NUM, bold=False, fill=None, kind=None, italic=False, align=None):
    """Write a value/formula with the colour code applied automatically."""
    c = ws[ref]
    c.value = value
    k = kind or _kind(value)
    c.font = _font(COLOURS[k], bold=bold, italic=italic)
    c.number_format = fmt
    if fill:
        c.fill = _fill(fill)
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def input_cell(ws, ref, value, fmt=NUM, bold=False):
    """An assumption: blue on light yellow with a hair box. value may be None (an optional override cell)."""
    c = ws[ref]
    c.value = value
    c.font = _font(BLUE, bold=bold)
    c.fill = _fill(INPUT_FILL)
    c.border = BOX_HAIR
    c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    return c


def output_box(ws, ref, value, fmt):
    """A key answer: bold with a thin box. No fill (a fill would read as 'input' or 'disabled')."""
    c = put(ws, ref, value, fmt, bold=True)
    c.border = BOX
    return c


key_output = output_box  # v1 name


def na_cell(ws, ref):
    """Not applicable (e.g. Year 0 of a forecast driver): grey fill, no value."""
    ws[ref].fill = _fill(NAFILL)


def row_of(ws, row, first_col, values, fmt=NUM, bold=False, total=False, fill=None, italic=False, grand=False):
    """Write a list of values/formulas across a row starting at first_col (None = skip)."""
    for i, v in enumerate(values):
        if v is None:
            continue
        c = put(ws, f"{col(first_col + i)}{row}", v, fmt, bold, fill, italic=italic)
        if grand:
            c.border = GRAND
        elif total:
            c.border = TOTAL


def comment(ws, ref, text):
    c = ws[ref]
    c.value = f"<< {text}"
    c.font = _font(GREY, italic=True, size=9)


def note(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = _font(GREY, italic=True, size=9)


def check_cell(ws, ref, formula):
    """A check that must evaluate to 0: shows OK, turns red and bold if not."""
    put(ws, ref, formula, fmt=CHECK, align="right")
    ws.conditional_formatting.add(
        ref, CellIsRule(operator="notBetween", formula=["-0.001", "0.001"], font=Font(color=RED, bold=True)))


def freeze(ws, ref):
    ws.freeze_panes = ref


# ---------------------------------------------------------------------------------------------------- blocks
def heat_map(ws, rng, better="high"):
    """3-colour scale: green = better for the reader (high IRR; low cost), white at the midpoint, red = worse."""
    lo, hi = (HEAT_RED, HEAT_GREEN) if better == "high" else (HEAT_GREEN, HEAT_RED)
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile",
                                                      mid_value=50, mid_color=HEAT_MID, end_type="max", end_color=hi))


def sensitivity(ws, top, left, out_formula, side_vals, top_vals, side_title, top_title, side_fmt, top_fmt, body_fmt,
                base_side=None, base_top=None, better="high", centre=False):
    """Two-way sensitivity block (bank layout):

        row top      :                         [   top_title centred over the value columns   ]
        row top+1    :               corner    t1      t2      t3      t4      t5      <- rule below
        rows top+2.. : side_title    s1 |      .       .       .       .       .
                                     s2 |      ...  (heat map, base case bold + boxed)
    left = column index of the corner cell (side axis sits under it). side_title is merged over the two columns
    to its left. Axis values are typed (blue); the corner links to the output. centre=True (formula axes built as
    base ± steps, e.g. a blank template): the middle cell is the base case by construction and is boxed without the
    numeric check. Returns the data-table range
    ("B5:G10" style, corner to bottom-right) and the cell of the base case.
    """
    n_top, n_side = len(top_vals), len(side_vals)
    c0, c1 = left + 1, left + n_top
    ws.merge_cells(start_row=top, start_column=c0, end_row=top, end_column=c1)
    t = ws.cell(top, c0, top_title)
    t.font = _font(BLACK, bold=True)
    t.alignment = Alignment(horizontal="center")
    corner = put(ws, f"{col(left)}{top + 1}", out_formula, body_fmt, bold=True, align="right")
    corner.border = Border(bottom=thin, right=thin)
    for j, v in enumerate(top_vals):
        cell = put(ws, f"{col(c0 + j)}{top + 1}", v, top_fmt, align="right")
        cell.border = BOTTOM
    for i, v in enumerate(side_vals):
        cell = put(ws, f"{col(left)}{top + 2 + i}", v, side_fmt, align="right")
        cell.border = Border(right=thin)
        for j in range(n_top):
            b = ws.cell(top + 2 + i, c0 + j)
            b.number_format = body_fmt
            b.font = _font(BLACK)
    r_first, r_last = top + 2, top + 1 + n_side
    ws.merge_cells(start_row=r_first, start_column=left - 2, end_row=r_last, end_column=left - 1)
    s = ws.cell(r_first, left - 2, side_title)
    s.font = _font(BLACK, bold=True)
    s.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
    body = f"{col(c0)}{r_first}:{col(c1)}{r_last}"
    heat_map(ws, body, better)
    base_ref = None
    if centre:
        assert n_side % 2 and n_top % 2, "formula axes need odd dimensions so there is a true centre"
        i, j = n_side // 2, n_top // 2
        base_ref = f"{col(c0 + j)}{r_first + i}"
    elif base_side is not None and base_top is not None:
        i = min(range(n_side), key=lambda k: abs(side_vals[k] - base_side))
        j = min(range(n_top), key=lambda k: abs(top_vals[k] - base_top))
        assert abs(side_vals[i] - base_side) < 1e-9 and abs(top_vals[j] - base_top) < 1e-9, \
            f"sensitivity axes must contain the base case exactly ({base_side}, {base_top})"
        assert i == n_side // 2 and j == n_top // 2, "base case must sit in the centre of the grid"
        base_ref = f"{col(c0 + j)}{r_first + i}"
    if base_ref:
        b = ws[base_ref]
        b.font = _font(BLACK, bold=True)
        b.border = BOX
    return f"{col(left)}{top + 1}:{col(c1)}{r_last}", base_ref


def sources_uses(ws, row, units, sources, uses, ebitda_ref=None, amt_col=4, guard=False):
    """Stacked Sources then Uses tables (navy header row, % of total, x EBITDA, double-ruled totals).

    sources / uses: [(label, formula), ...] (use the string "PLUG" as the formula for sponsor equity = uses - other
    sources). guard=True wraps every ratio in IF(denominator=0,0,...) so a blank template shows "–", not #DIV/0!.
    Returns ({label: amount ref}, total_sources_ref, total_uses_ref, next free row).
    """
    A, P, X = col(amt_col), col(amt_col + 1), col(amt_col + 2)

    def ratio(num, den):
        return f"=IF({den}=0,0,{num}/{den})" if guard else f"={num}/{den}"
    last = amt_col + (2 if ebitda_ref else 1)
    refs = {}

    def head(r, text):
        for c in range(2, last + 1):
            cell = ws.cell(r, c)
            cell.fill = _fill(NAVY)
            cell.font = _font(WHITE, bold=True)
        ws.cell(r, 2, text)
        for c, t in ((amt_col, units), (amt_col + 1, "% of total")) + (((amt_col + 2, "x EBITDA"),) if ebitda_ref else ()):
            ws.cell(r, c, t).alignment = Alignment(horizontal="right")

    def block(r, text, items, total_label):
        head(r, text)
        first = r + 1
        tot_row = first + len(items)
        for i, (lab, f) in enumerate(items):
            rr = first + i
            ws.cell(rr, 2, lab).font = _font(BLACK)
            refs[lab] = f"{A}{rr}"
            if f != "PLUG":
                put(ws, f"{A}{rr}", f)
            put(ws, f"{P}{rr}", ratio(f"{A}{rr}", f"{A}${tot_row}"), PCT)
            if ebitda_ref:
                put(ws, f"{X}{rr}", ratio(f"{A}{rr}", ebitda_ref), MULT)
        ws.cell(tot_row, 2, total_label).font = _font(BLACK, bold=True)
        for c in range(2, last + 1):
            ws.cell(tot_row, c).border = GRAND
        return first, tot_row

    s_first, s_tot = block(row, "Sources of funds", sources, "Total sources")
    u_row = s_tot + 2
    u_first, u_tot = block(u_row, "Uses of funds", uses, "Total uses")
    put(ws, f"{A}{u_tot}", f"=SUM({A}{u_first}:{A}{u_tot - 1})", bold=True).border = GRAND
    put(ws, f"{P}{u_tot}", f"=SUM({P}{u_first}:{P}{u_tot - 1})", PCT, bold=True).border = GRAND
    for i, (lab, f) in enumerate(sources):
        if f == "PLUG":
            others = "-".join(f"{A}{s_first + k}" for k in range(len(sources)) if sources[k][1] != "PLUG")
            put(ws, f"{A}{s_first + i}", f"={A}{u_tot}" + (f"-{others}" if others else ""))
    put(ws, f"{A}{s_tot}", f"=SUM({A}{s_first}:{A}{s_tot - 1})", bold=True).border = GRAND
    put(ws, f"{P}{s_tot}", f"=SUM({P}{s_first}:{P}{s_tot - 1})", PCT, bold=True).border = GRAND
    if ebitda_ref:
        for rr in (s_tot, u_tot):
            put(ws, f"{X}{rr}", ratio(f"{A}{rr}", ebitda_ref), MULT, bold=True).border = GRAND
    return refs, f"{A}{s_tot}", f"{A}{u_tot}", u_tot + 2


def colour_key(ws, row, col_letter="B"):
    """The colour code, shown on the Cover exactly as used in the model. Font colour and fill are two separate
    signals, so each gets its own row: blue font = hardcode; yellow fill + hair box = assumption (an input cell
    carries both)."""
    items = [("Hardcode (blue font)", "input", None), ("Assumption / input cell (yellow fill)", "calc", INPUT_FILL),
             ("Link to another sheet (green font)", "link", None), ("Formula (black font)", "calc", None),
             ("Link to another workbook (red font)", "external", None)]
    for i, (text, kind, fill) in enumerate(items):
        c = ws[f"{col_letter}{row + i}"]
        c.value = text
        c.font = _font(COLOURS[kind])
        if fill:
            c.fill = _fill(fill)
            c.border = BOX_HAIR
    return row + len(items)
