"""Format linter: checks a workbook against the house rules (references/house-style.md, lessons.md).

    python model_lint.py model.xlsx [--summary]      -> list of violations; exit 1 if any

Rules (each maps to a lesson in lessons.md):
  L1 band       every sheet starts with the navy header band (no big title + subtitle)
  L2 font       Arial only
  L3 colour     typed number = blue 0000FF · pure link to another sheet = green 008000 · formula = black ·
                link to another workbook = red FF0000
  L4 fill       FFFFCC only on input (blue) cells; no fill on formula/output cells (bars, band and N/A grey excepted)
  L5 heatmap    every 'sensitivity' section has a colour-scale rule
  L6 hardcode   no numbers typed inside formulas (true constants excepted: 0, 1, 3, 4, 12, 100, 1000, 365, 8760, 0.5)
  L7 compat     no functions missing from Excel 2016 (MINIFS, MAXIFS, XLOOKUP, LET, IFS, SWITCH, ...)
"""
import re
import sys
from collections import Counter

from openpyxl import load_workbook

NAVY, BAR, SUBBAR, NAFILL, INPUT_FILL = "1F3864", "D9D9D9", "DCE6F1", "BFBFBF", "FFFFCC"
BLUE, GREEN, BLACK, RED = "0000FF", "008000", "000000", "FF0000"
ALLOWED_CONSTANTS = {0.0, 1.0, 3.0, 4.0, 12.0, 100.0, 1000.0, 365.0, 8760.0, 0.5}  # 3 = months per quarter
BANNED = ("MINIFS", "MAXIFS", "XLOOKUP", "XMATCH", "LET(", "IFS(", "SWITCH(", "FILTER(", "UNIQUE(", "SORT(",
          "SORTBY(", "SEQUENCE(", "TEXTJOIN(", "CONCAT(", "LAMBDA(")
LINK = re.compile(r"^=-?(?:'[^']+'|[A-Za-z0-9_.]+)!\$?[A-Z]{1,3}\$?\d+$")


def rgb(color):
    """'FF0000FF' -> '0000FF'; theme / None -> None."""
    try:
        v = color.rgb if color is not None else None
    except Exception:
        return None
    if not isinstance(v, str):
        return None
    return v[-6:].upper()


def kind(v):
    s = v.replace(" ", "")
    if "[" in s and "]" in s:
        return "external"
    return "link" if LINK.match(s) else "calc"


def literals(formula):
    s = re.sub(r'"[^"]*"', "", formula)                      # strings
    s = re.sub(r"'[^']+'!", "", s)                            # quoted sheet names
    s = re.sub(r"[A-Za-z_][A-Za-z0-9_.]*!", "", s)            # sheet names
    s = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", s)                 # cell refs
    s = re.sub(r"[A-Za-z_][A-Za-z0-9_.]*\(", "(", s)          # function names
    s = re.sub(r"[A-Za-z_][A-Za-z0-9_.]*", "", s)             # defined names
    return [float(x) for x in re.findall(r"(?<![\w.])\d+(?:\.\d+)?(?:[eE][-+]?\d+)?", s)]


def band_rows(ws):
    n = 0
    for r in range(1, 30):
        if rgb(ws.cell(r, 2).fill.fgColor) == NAVY and ws.cell(r, 2).fill.fill_type == "solid":
            n = r
        else:
            break
    return n


def lint(path):
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        band = band_rows(ws)
        if band < 3:
            out.append(("L1 band", f"{ws.title}: no navy header band at the top"))
        sens_sections = 0
        dt_cells = set()  # cells filled by Excel data tables hold plain numbers: not typed inputs
        for row in ws.iter_rows():
            for c in row:
                if type(c.value).__name__ == "DataTableFormula":
                    for rr in ws[c.value.ref]:
                        for x in rr:
                            dt_cells.add(x.coordinate)
        for row in ws.iter_rows(min_row=band + 1):
            for c in row:
                v = c.value
                if c.coordinate in dt_cells:
                    continue
                fill = rgb(c.fill.fgColor) if c.fill.fill_type == "solid" else None
                fc = rgb(c.font.color) or BLACK
                where = f"{ws.title}!{c.coordinate}"
                if isinstance(v, str) and "sensitiv" in v.lower() and fill in (BAR, NAVY) and c.column == 2:
                    sens_sections += 1
                if v is None:
                    if fill == INPUT_FILL and fc != BLUE:
                        out.append(("L4 fill", f"{where}: yellow input fill on a non-blue cell"))
                    continue
                if c.font.name and c.font.name != "Arial":
                    out.append(("L2 font", f"{where}: font {c.font.name}"))
                if fill in (NAVY, BAR, SUBBAR):
                    continue
                if isinstance(v, bool):
                    continue
                if isinstance(v, (int, float)):
                    if fc != BLUE:
                        out.append(("L3 colour", f"{where}: typed number {v} is not blue"))
                    if fill not in (None, INPUT_FILL):
                        out.append(("L4 fill", f"{where}: fill {fill} on a number"))
                elif isinstance(v, str) and v.startswith("="):
                    k = kind(v)
                    want = {"link": GREEN, "calc": BLACK, "external": RED}[k]
                    if fc != want:
                        out.append(("L3 colour", f"{where}: {k} should be {want}, is {fc}  ({v[:40]})"))
                    if fill is not None:
                        out.append(("L4 fill", f"{where}: fill {fill} on a formula/output (use a box border)"))
                    bad = [x for x in literals(v) if x not in ALLOWED_CONSTANTS]
                    if bad:
                        out.append(("L6 hardcode", f"{where}: number(s) {bad} typed inside {v[:50]}"))
                    up = v.upper()
                    for b in BANNED:
                        if re.search(r"(?<![A-Z0-9_.])" + re.escape(b.rstrip("(")) + r"\(", up):
                            out.append(("L7 compat", f"{where}: {b.rstrip('(')} is not in Excel 2016"))
                elif fill == NAFILL:
                    out.append(("L4 fill", f"{where}: N/A grey on a cell with a value"))
        scales = sum(1 for rng in ws.conditional_formatting for rule in rng.rules if rule.type == "colorScale")
        if sens_sections > scales:
            out.append(("L5 heatmap", f"{ws.title}: {sens_sections} sensitivity section(s) but {scales} colour scale(s)"))
    return out


if __name__ == "__main__":
    res = lint(sys.argv[1])
    counts = Counter(k for k, _ in res)
    print(f"model_lint: {len(res)} violation(s)" + (f"  {dict(counts)}" if res else ""))
    if "--summary" not in sys.argv:
        for k, msg in res[:60]:
            print(f"  [{k}] {msg}")
        if len(res) > 60:
            print(f"  ... {len(res) - 60} more")
    sys.exit(1 if res else 0)
