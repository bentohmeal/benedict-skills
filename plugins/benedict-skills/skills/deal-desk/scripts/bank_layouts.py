"""Bank page designs (bk##): recurring layouts from the 389 real IB / PE pages in "Real IB-PE Decks", rebuilt in the
house style so deal-desk can fill them with its own numbers.

    d.bank(s, "bk13", box=(x, y, w, h), data={...})     # box defaults to the whole slide body

Each design is a function bkNN(d, s, x, y, w, h, **data) built from Deck primitives (tables, native charts, shapes).
bank_specs.SPECS holds what layout_tag.py build puts in assets/layout_library.json (use when / don't use when, data
type, items, the real pages it was taken from); bank_samples.SAMPLE holds one filled example per design: the data
shape to copy, also used for the thumbnails (python bank_layouts.py thumbs) and gallery (... gallery out.pptx).
Numbers in tables are pre-formatted strings; chart values are numbers. Ids follow Benedict's review sheet
(bank_designs_review.xlsx, 2 Oct 2026): cut designs leave gaps, bk02 was merged into bk01 (panel=...).
"""
import calendar
import os
import statistics
import sys

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from ibdeck import (BLACK, DEEP, DGREY, F_BODY, F_HEAD, GOLD, GREY, LGREY, LIGHT, MID, NAVY, PALE, PALEGREY,  # noqa
                    STEEL, WHITE, Deck, rgb)

DASH = MSO_LINE_DASH_STYLE.DASH
C, LFT, RGT = PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.RIGHT
MID_A, TOP_A = MSO_ANCHOR.MIDDLE, MSO_ANCHOR.TOP
RAMP = [NAVY, MID, STEEL, LIGHT, LGREY]


# ---------------------------------------------------------------- helpers
def _t(d, s, x, y, w, h, text, **kw):
    return d._text(s, x, y, w, h, text, **kw)


def _head(d, s, x, y, w, text, style="rule"):
    """Panel title. 'rule' (navy bold + thin rule) inside dense pages; 'bar' = the filled navy header bar."""
    d.header(s, x, y, w, text, style=style)
    return y + 0.38


def _vtab(d, s, x, y, w, h, text, fill=NAVY, color=WHITE):
    """Vertical tab: filled box with text rotated to read bottom-up (bank side tabs)."""
    b = d._rect(s, x, y, w, h, fill=fill)
    d._shape_text(b, text, size=F_BODY, bold=True, color=color)
    b.text_frame._txBody.find(qn("a:bodyPr")).set("vert", "vert270")
    return b


def _circle_no(d, s, x, y, n, dia=0.22, fill=NAVY):
    c = d._shape(s, MSO_SHAPE.OVAL, x, y, dia, dia, fill=fill)
    d._shape_text(c, str(n), size=F_BODY, bold=True, color=WHITE, margin=0)
    return c


def _box(d, s, x, y, w, h, text="", line=NAVY, fill=None, dash=None, size=F_BODY, color=BLACK, align=LFT,
         anchor=TOP_A, bold=False):
    b = d._shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, line=line, line_w=1.0, dash=dash)
    if text:
        d._shape_text(b, text, size=size, color=color, align=align, anchor=anchor, margin=0.08, bold=bold)
    return b


def _frame(ch, vmin, vmax, cat_labels=False):
    """Fix the value axis and stretch the plot area over the whole chart frame, so category k sits at
    x + (k + 0.5) * w / n and value v at y + h - (v - vmin) / (vmax - vmin) * h (for drawing on top of the chart)."""
    ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = vmin, vmax
    if not cat_labels:
        ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    pa = ch._chartSpace.chart.plotArea
    lay = pa.find(qn("c:layout"))
    if lay is None:
        lay = pa.makeelement(qn("c:layout"), {})
        pa.insert(0, lay)
    for old in list(lay):
        lay.remove(old)
    ml = lay.makeelement(qn("c:manualLayout"), {})
    for tag, val in (("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                     ("c:x", "0"), ("c:y", "0"), ("c:w", "1"), ("c:h", "1")):
        ml.append(ml.makeelement(qn(tag), {"val": val}))
    lay.append(ml)
    va = ch._chartSpace.chart.plotArea.find(qn("c:valAx"))
    cb = va.find(qn("c:crossBetween")) if va is not None else None
    if cb is not None:
        cb.set("val", "between")


def _map(x, y, w, h, n, vmin, vmax):
    fx = lambda k: x + (k + 0.5) * w / n                       # noqa: E731
    fy = lambda v: y + h - (v - vmin) / (vmax - vmin) * h       # noqa: E731
    return fx, fy


def _scale(vals, pad_top=0.25, zero=True):
    lo, hi = min(vals), max(vals)
    lo = min(0, lo) if zero else lo - (hi - lo) * 0.08
    return lo, hi + (hi - lo) * pad_top


def _xlabels(d, s, x, y, w, labels, every=1, h=0.22):
    n = len(labels)
    for k, lab in enumerate(labels):
        if k % every == 0 or k == n - 1:
            _t(d, s, x + k * w / n - w / n * 1.5, y, w / n * 4, h, str(lab), align=C, wrap=False)


def _line_chart(d, s, x, y, w, h, cats, series, colors=None, vmin=None, vmax=None, every=None, width=1.75,
                extra=()):
    """Price / index line without markers or labels; date labels drawn underneath. extra = values the scale must
    also cover (offer price, band). Returns (fx, fy)."""
    allv = [v for _, vs in series for v in vs if v is not None] + [v for v in extra if v is not None]
    lo, hi = _scale(allv, 0.12, zero=False)
    vmin = lo if vmin is None else vmin
    vmax = hi if vmax is None else vmax
    ch = d._chart(s, x, y, w, h - 0.24, XL_CHART_TYPE.LINE, list(cats), series, labels=False,
                  legend=len(series) > 1, colors=colors or [NAVY, GOLD, STEEL, LGREY, MID])
    for ser in ch.series:
        ser.format.line.width = Pt(width)
    _frame(ch, vmin, vmax)
    d._line(s, x, y + h - 0.24, x + w, y + h - 0.24, LGREY, 0.75)
    _xlabels(d, s, x, y + h - 0.22, w, cats, every or max(1, len(cats) // 6))
    _t(d, s, x - 0.55, y - 0.02, 0.5, 0.2, f"{vmax:,.0f}", color=GREY, align=RGT, wrap=False)
    _t(d, s, x - 0.55, y + h - 0.46, 0.5, 0.2, f"{vmin:,.0f}", color=GREY, align=RGT, wrap=False)
    return _map(x, y, w, h - 0.24, len(cats), vmin, vmax)


def _cols(d, s, x, y, w, h, cats, vals, fmt="#,##0.0", highlight=None, color=NAVY, vmin=None, vmax=None,
          cat_labels=True, gap=60):
    """Single-series column chart on the fixed frame; category labels drawn underneath. Returns (fx, fy)."""
    lo, hi = _scale(vals)
    vmin = lo if vmin is None else vmin
    vmax = hi if vmax is None else vmax
    lab_h = 0.3 if cat_labels else 0
    ch = d._chart(s, x, y, w, h - lab_h, XL_CHART_TYPE.COLUMN_CLUSTERED, list(cats), [("v", vals)],
                  number_format=fmt, legend=False, colors=[color], gap=gap)
    _frame(ch, vmin, vmax)
    for k in range(len(vals)):
        pt = ch.series[0].points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(GOLD if k == highlight else color)
    if cat_labels:
        n = len(cats)
        for k, c in enumerate(cats):
            _t(d, s, x + k * w / n, y + h - lab_h + 0.02, w / n, lab_h, str(c), align=C, anchor=TOP_A)
    return _map(x, y, w, h - lab_h, len(cats), vmin, vmax)


def _median_line(d, s, x, w, yv, text, color=GREY):
    d._line(s, x, yv, x + w, yv, color, 1.0, dash=DASH)
    _t(d, s, x + w - 1.6, yv - 0.22, 1.6, 0.2, text, color=DGREY, align=RGT, bold=True, wrap=False)


def _donut(d, s, x, y, dia, labels, values, centre=None, colors=None, hole=60):
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series("s", values)
    ch = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(dia), Inches(dia), cd).chart
    ch.has_legend = ch.has_title = False
    plot = ch.plots[0]
    plot.has_data_labels = False
    for k in range(len(values)):
        pt = plot.series[0].points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb((colors or RAMP)[k % len(colors or RAMP)])
        pt.format.line.color.rgb = rgb(WHITE)
    dn = plot._element
    hs = dn.find(qn("c:holeSize"))
    if hs is None:
        hs = dn.makeelement(qn("c:holeSize"), {})
        dn.append(hs)
    hs.set("val", str(hole))
    if centre:
        _t(d, s, x, y + dia / 2 - 0.2, dia, 0.4, centre, bold=True, align=C, anchor=MID_A)


def _legend(d, s, x, y, items, w=1.6, size=0.12, row=0.24, colors=None):
    for k, lab in enumerate(items):
        d._rect(s, x, y + k * row + 0.05, size, size, fill=(colors or RAMP)[k % len(colors or RAMP)])
        _t(d, s, x + size + 0.06, y + k * row, w, row, lab, anchor=MID_A, wrap=False)


def _shade_cells(tbl, rows, cols, fill):
    for i in rows:
        for j in cols:
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(fill)


def _bold_cells(tbl, rows, cols):
    for i in rows:
        for j in cols:
            for r in tbl.cell(i, j).text_frame.paragraphs[0].runs:
                r.font.bold = True


def _num(v):
    try:
        return float(str(v).replace("$", "").replace(",", "").replace("x", "").replace("%", "").replace("(", "-")
                     .replace(")", ""))
    except ValueError:
        return None


def _fmtv(v, kind="line"):
    """fin_table number style: brackets for negatives, $ on key lines, % rows to 1dp, None -> en dash."""
    if v is None or v == "":
        return "–"
    if isinstance(v, str):
        return v
    if kind == "pct":
        return f"({abs(v):.1%})" if v < 0 else f"{v:.1%}"
    n = f"({abs(v):,.1f})" if v < 0 else f"{v:,.1f}"
    return "$ " + n if kind == "key" else n


def _notes_grid(d, s, x, y, w, h, notes, cols=2):
    """Numbered commentary boxes (number circle + text) in `cols` columns."""
    n = len(notes)
    rows = -(-n // cols)
    cw, rh = (w - 0.2 * (cols - 1)) / cols, min(0.55, h / max(rows, 1))
    for k, txt in enumerate(notes):
        cx, cy = x + (k // rows) * (cw + 0.2), y + (k % rows) * rh
        _circle_no(d, s, cx, cy + 0.03, k + 1)
        _t(d, s, cx + 0.3, cy, cw - 0.3, rh, txt)


# ---------------------------------------------------------------- deal terms
def bk01(d, s, x, y, w, h, prices, rows, highlight=None, price_label="Offer price per share", panel=None,
         panel_title="Key assumptions"):
    """Offer summary: one column per price; rows = (label, [values]) or (label, [values], "key"/"group"/"pct").
    panel = [(label, value)] adds the proposal / key assumptions box on the right (old bk02)."""
    tw = w * 0.64 if panel else w
    _t(d, s, x + tw * 0.34, y, tw * 0.66, 0.25, price_label, bold=True, align=C)
    d._line(s, x + tw * 0.34 + 0.05, y + 0.27, x + tw - 0.05, y + 0.27, NAVY, 1.0)
    head = [""] + list(prices)
    body, bold, shade, groups = [], [], [], []
    for r in rows:
        kind = r[2] if len(r) > 2 else "line"
        if kind == "group":
            body.append([r[0]] + [""] * len(prices))
            groups.append(len(body))
        else:
            body.append([("    " if kind == "pct" else "") + r[0]] + list(r[1]))
            if kind == "key":
                bold.append(len(body))
    rh = min(0.3, (h - 0.35) / (len(body) + 1))
    lw = tw * 0.34
    shp = d.table(s, x, y + 0.32, tw, [head] + body, col_widths=[lw] + [(tw - lw) / len(prices)] * len(prices),
                  row_h=rh, bold_rows=bold, header_style="rule")
    tbl = shp.table
    for i in groups:
        _bold_cells(tbl, [i], [0])
        for r in tbl.cell(i, 0).text_frame.paragraphs[0].runs:
            r.font.color.rgb = rgb(NAVY)
    if highlight is not None:
        _shade_cells(tbl, range(1, len(body) + 1), [highlight + 1], PALE)
        _bold_cells(tbl, [0], [highlight + 1])
    if panel:
        px, pw = x + tw + 0.3, w - tw - 0.3
        yy = _head(d, s, px, y, pw, panel_title, "bar")
        d.pill_rows(s, px, yy + 0.05, pw, min(h - 0.45, 0.62 * len(panel)), panel, label_w=pw * 0.38)
    return shp


def bk05(d, s, x, y, w, h, rows, heads=("Concern", "Our response")):
    """Issue -> response pairs: rows = [(issue, response)]; issue box grey outline, arrow, response box navy."""
    aw = 0.45
    bw = (w - aw - 0.2) / 2
    _t(d, s, x, y, bw, 0.28, heads[0], bold=True, color=NAVY, size=F_HEAD)
    _t(d, s, x + bw + aw + 0.2, y, bw, 0.28, heads[1], bold=True, color=NAVY, size=F_HEAD)
    n = len(rows)
    rh = (h - 0.4 - 0.15 * (n - 1)) / n
    for k, (a, b) in enumerate(rows):
        ry = y + 0.4 + k * (rh + 0.15)
        _box(d, s, x, ry, bw, rh, a, line=GREY, anchor=MID_A)
        ar = d._shape(s, MSO_SHAPE.RIGHT_ARROW, x + bw + 0.1, ry + rh / 2 - 0.18, aw, 0.36, fill=NAVY)
        ar.adjustments[0] = 0.5
        _box(d, s, x + bw + aw + 0.2, ry, bw, rh, b, line=NAVY, anchor=MID_A)


# ---------------------------------------------------------------- financials
def bk06(d, s, x, y, w, h, periods, lines, notes, split=None, marks=None):
    """Projections table + numbered commentary below. lines as fin_table; marks = {line index: note number} puts
    the note's number beside that table row; split = index of the first projected period (dashed divider)."""
    n_rows = len(lines) + 1
    note_h = max(0.9, min(1.7, 0.32 * -(-len(notes) // 2)))
    rh = min(0.25, (h - note_h - 0.25) / n_rows)
    lw = w * 0.3
    d.fin_table(s, x, y, w, periods, lines, label_w=lw, row_h=rh)
    if split is not None:
        sx = x + lw + split * (w - lw) / len(periods)
        d._line(s, sx, y, sx, y + rh * n_rows, GREY, 1.0, dash=DASH)
    for li, no in (marks or {}).items():
        _circle_no(d, s, x + lw - 0.27, y + rh * (li + 1) + (rh - 0.2) / 2, no, dia=0.2)
    ny = y + rh * n_rows + 0.2
    d._line(s, x, ny, x + w, ny, NAVY, 1.0)
    _notes_grid(d, s, x, ny + 0.1, w, y + h - ny - 0.1, notes)


def bk07(d, s, x, y, w, h, periods, segments, commentary):
    """Segment financials: segments = [(name, lines)] each with a vertical navy tab; commentary = one text per
    segment in a numbered column on the right."""
    tab_w, cw = 0.32, w * 0.32
    tw = w - tab_w - cw - 0.25
    nrows = sum(len(ls) for _, ls in segments) + 1
    rh = min(0.25, (h - 0.1 * len(segments)) / nrows)
    lw = tw * 0.36
    # period header once
    for k, p in enumerate(periods):
        _t(d, s, x + tab_w + lw + k * (tw - lw) / len(periods), y, (tw - lw) / len(periods), rh, p, bold=True,
           align=RGT, anchor=MID_A)
    d._line(s, x + tab_w, y + rh, x + tab_w + tw, y + rh, NAVY, 1.0)
    _t(d, s, x + tab_w + tw + 0.25, y, cw, rh, "Commentary", bold=True, anchor=MID_A)
    d._line(s, x + tab_w + tw + 0.25, y + rh, x + w, y + rh, NAVY, 1.0)
    yy = y + rh + 0.06
    for k, (name, ls) in enumerate(segments):
        sh = rh * len(ls)
        _vtab(d, s, x, yy, tab_w - 0.04, sh, name, fill=RAMP[k % 3])
        rows = [[ln[0]] + [_fmtv(v, ln[2] if len(ln) > 2 else "line") for v in ln[1]] for ln in ls]
        shp = d.table(s, x + tab_w, yy, tw, rows, col_widths=[lw] + [(tw - lw) / len(periods)] * len(periods),
                      row_h=rh, header=False, bold_rows=[i for i, ln in enumerate(ls) if len(ln) > 2 and ln[2] == "key"])
        for i, ln in enumerate(ls):
            if len(ln) > 2 and ln[2] == "pct":
                for j in range(len(periods) + 1):
                    for r in shp.table.cell(i, j).text_frame.paragraphs[0].runs:
                        r.font.italic = True
                        r.font.color.rgb = rgb(GREY)
        cx = x + tab_w + tw + 0.25
        _circle_no(d, s, cx, yy + 0.03, k + 1, dia=0.2)
        _t(d, s, cx + 0.28, yy, cw - 0.28, sh, commentary[k] if k < len(commentary) else "")
        yy += sh + 0.1
        d._line(s, x, yy - 0.05, x + w, yy - 0.05, LGREY, 0.75)


def bk08(d, s, x, y, w, h, periods, lines, split, extra=None, labels=("Historical", "Projected")):
    """Historical | projected table: span labels over the two period groups, dashed divider, extra columns at the
    end (extra = (heading, [value per line]) e.g. CAGR)."""
    lw = w * 0.28
    ew = w * 0.1 if extra else 0
    pw = (w - lw - ew) / len(periods)
    for k, (a, b) in enumerate(((0, split), (split, len(periods)))):
        x0, x1 = x + lw + a * pw + 0.05, x + lw + b * pw - 0.05
        _t(d, s, x0, y, x1 - x0, 0.24, labels[k], bold=True, align=C, color=NAVY if k else BLACK)
        d._line(s, x0, y + 0.26, x1, y + 0.26, NAVY if k else GREY, 1.25)
    rh = min(0.26, (h - 0.35) / (len(lines) + 1))
    per = list(periods) + ([extra[0]] if extra else [])
    ls = [(ln[0], list(ln[1]) + ([extra[1][i]] if extra else []), ln[2] if len(ln) > 2 else "line")
          for i, ln in enumerate(lines)]
    shp = d.fin_table(s, x, y + 0.32, w, per, ls, label_w=lw, row_h=rh)
    tbl = shp.table
    for j in range(1, len(per) + 1):
        tbl.columns[j].width = Inches(pw if j <= len(periods) else ew)
    _shade_cells(tbl, range(1, len(lines) + 1), range(split + 1, len(periods) + 1), PALE)
    sx = x + lw + split * pw
    d._line(s, sx, y + 0.3, sx, y + 0.32 + rh * (len(lines) + 1), GREY, 1.0, dash=DASH)
    return shp


def bk09(d, s, x, y, w, h, rows, names=("Management", "Consensus"), note=None):
    """Management vs consensus: rows = [(metric, mgmt, consensus, fmt)] (numbers). Table on the left, dot rows on
    the right with the gap as a chip (management navy, consensus grey)."""
    tw = w * 0.45
    hdr = ["Metric", names[0], names[1], "Gap"]
    body = []
    for m, a, b, f in rows:
        gap = (a - b) / abs(b) if b else 0
        body.append([m, f.format(a), f.format(b), f"{gap:+.1%}"])
    rh = min(0.42, (h - 0.6) / (len(rows) + 1))
    d.table(s, x, y, tw, [hdr] + body, col_widths=[tw * 0.37] + [tw * 0.21] * 3, row_h=rh)
    px, pw = x + tw + 0.4, w - tw - 0.4
    _legend(d, s, px, y + 0.02, list(names), colors=[NAVY, LGREY], w=1.4)
    for k, (m, a, b, f) in enumerate(rows):
        ry = y + rh * (k + 1) + rh / 2
        lo, hi = min(a, b), max(a, b)
        span = (hi - lo) or abs(hi) or 1
        x0, x1 = lo - span * 1.2, hi + span * 1.2
        fx = lambda v: px + 0.1 + (v - x0) / (x1 - x0) * (pw - 1.0)   # noqa: E731
        d._line(s, px + 0.1, ry, px + pw - 0.9, ry, LGREY, 1.0)
        d._line(s, fx(a), ry, fx(b), ry, NAVY, 2.0)
        for v, col in ((b, LGREY), (a, NAVY)):
            d._shape(s, MSO_SHAPE.OVAL, fx(v) - 0.08, ry - 0.08, 0.16, 0.16, fill=col)
        gap = (a - b) / abs(b) if b else 0
        chip = d._shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px + pw - 0.8, ry - 0.13, 0.78, 0.26, fill=None,
                        line=GOLD if abs(gap) >= 0.05 else NAVY)
        d._shape_text(chip, f"{gap:+.1%}", size=F_BODY, bold=True, color=BLACK)
    if note:
        d.takeaway(s, x, y + h - 0.5, w, 0.5, note)


def bk10(d, s, x, y, w, h, rows, heads=("", "Prior", "Revised", "Change", "Comment"), revised=()):
    """Revision / variance table: rows = [label, prior, revised, change, comment]; rows in `revised` (0-based)
    are boxed with a dashed gold outline."""
    rh = min(0.34, h / (len(rows) + 1))
    cw = [w * 0.24, w * 0.12, w * 0.12, w * 0.12, w * 0.4]
    shp = d.table(s, x, y, w, [list(heads)] + [list(r) for r in rows], col_widths=cw, row_h=rh, num_cols={1, 2, 3})
    for i in revised:
        _box(d, s, x - 0.04, y + rh * (i + 1) + 0.01, w + 0.08, rh - 0.02, line=GOLD, dash=DASH)
    return shp


def bk11(d, s, x, y, w, h, panels, names=("Management plan", "Street case")):
    """Plan vs street 2x2: panels = [{"title", "periods", "a": [...], "b": [...], "fmt", "callout"}]; each a
    clustered column of the two cases with a dashed callout."""
    n = len(panels)
    cols = 2 if n > 1 else 1
    rows = -(-n // cols)
    pw, ph = (w - 0.3 * (cols - 1)) / cols, (h - 0.25 * (rows - 1)) / rows
    for k, p in enumerate(panels):
        px, py = x + (k % cols) * (pw + 0.3), y + (k // cols) * (ph + 0.25)
        yy = _head(d, s, px, py, pw, p["title"])
        fmt = p.get("fmt", "#,##0")
        ch = d._chart(s, px, yy, pw, ph - (yy - py) - (0.45 if p.get("callout") else 0), XL_CHART_TYPE.COLUMN_CLUSTERED,
                      p["periods"], [(names[0], p["a"]), (names[1], p["b"])], number_format=fmt, legend=k == 0,
                      colors=[NAVY, LIGHT], gap=80)
        ch.value_axis.minimum_scale = 0
        if p.get("callout"):
            d.takeaway(s, px, py + ph - 0.42, pw, 0.4, p["callout"])


# ---------------------------------------------------------------- valuation
def bk13(d, s, x, y, w, h, periods, lines, sens, assumptions=None, split=None):
    """DCF page: FCF build (fin_table) on top, 1-2 sensitivity grids below, optional key-assumptions column.
    sens = [{"title", "corner", "cols", "rows", "values", "fmt", "base"}]."""
    aw = w * 0.26 if assumptions else 0
    tw = w - aw - (0.3 if assumptions else 0)
    rh = 0.22
    yy = _head(d, s, x, y, tw, "Unlevered free cash flow build")
    d.fin_table(s, x, yy, tw, periods, lines, label_w=tw * 0.3, row_h=rh)
    if split is not None:
        sx = x + tw * 0.3 + split * tw * 0.7 / len(periods)
        d._line(s, sx, yy, sx, yy + rh * (len(lines) + 1), GREY, 1.0, dash=DASH)
    sy = yy + rh * (len(lines) + 1) + 0.25
    gw = (tw - 0.3 * (len(sens) - 1)) / len(sens)
    for k, g in enumerate(sens):
        gx = x + k * (gw + 0.3)
        gy = _head(d, s, gx, sy, gw, g["title"])
        d.heatmap(s, gx, gy, gw, g.get("corner", ""), g["cols"], g["rows"], g["values"], fmt=g.get("fmt", "{:,.1f}"),
                  base=g.get("base"), row_h=min(0.26, (y + h - gy) / (len(g["rows"]) + 1)))
    if assumptions:
        ax = x + tw + 0.3
        ay = _head(d, s, ax, y, aw, "Key assumptions", "bar")
        d.bullets(s, ax, ay + 0.05, aw, h - (ay - y), assumptions)


def bk14(d, s, x, y, w, h, equity, debt, wacc, peers=None):
    """WACC build: cost of equity and cost of debt tables (label, value), the WACC result row, peer betas table.
    peers = rows with a header row, e.g. [["Company", "Levered beta", "D / E", "Unlevered beta"], ...]."""
    lw = w * 0.42 if peers else w
    rh = 0.26
    yy = _head(d, s, x, y, lw, "Cost of equity", "bar")
    d.table(s, x, yy, lw, [["Input", "Value"]] + [list(r) for r in equity], col_widths=[lw * 0.7, lw * 0.3],
            row_h=rh, bold_rows=(len(equity),))
    yy += rh * (len(equity) + 1) + 0.2
    yy = _head(d, s, x, yy, lw, "Cost of debt", "bar")
    d.table(s, x, yy, lw, [["Input", "Value"]] + [list(r) for r in debt], col_widths=[lw * 0.7, lw * 0.3],
            row_h=rh, bold_rows=(len(debt),))
    yy += rh * (len(debt) + 1) + 0.25
    d.stat_strip(s, x, yy, lw, min(0.75, y + h - yy), wacc)
    if peers:
        px, pw = x + lw + 0.35, w - lw - 0.35
        py = _head(d, s, px, y, pw, "Peer betas", "bar")
        rhp = min(0.3, (h - 0.4) / len(peers))
        shp = d.table(s, px, py, pw, peers, row_h=rhp)
        _shade_cells(shp.table, [len(peers) - 2, len(peers) - 1], range(len(peers[0])), PALE)
        _bold_cells(shp.table, [len(peers) - 2, len(peers) - 1], range(len(peers[0])))


def bk15(d, s, x, y, w, h, rows, refs=(), fmt="${:,.0f}", unit="", highlight=None, comment_head="Commentary"):
    """Football field + commentary column: rows = [(method, low, high, comment)]; refs = [(label, value)] dashed
    reference lines (offer price, current price). Bars are drawn shapes so each row lines up with its comment."""
    mw, cw = w * 0.2, w * 0.34
    fx0, fw = x + mw + 0.1, w - mw - cw - 0.45
    vals = [v for r in rows for v in r[1:3]] + [v for _, v in refs]
    lo, hi = min(vals), max(vals)
    pad = (hi - lo) * 0.18
    v0, v1 = lo - pad, hi + pad
    fx = lambda v: fx0 + (v - v0) / (v1 - v0) * fw    # noqa: E731
    _t(d, s, x, y, mw, 0.26, "Methodology", bold=True, anchor=MID_A)
    _t(d, s, fx0, y, fw, 0.26, f"Implied value{(' (' + unit + ')') if unit else ''}", bold=True, align=C,
       anchor=MID_A)
    _t(d, s, x + w - cw, y, cw, 0.26, comment_head, bold=True, anchor=MID_A)
    d._line(s, x, y + 0.3, x + w, y + 0.3, NAVY, 1.0)
    n = len(rows)
    rh = (h - 0.4) / n
    for k, (m, a, b, *cm) in enumerate(rows):
        ry = y + 0.35 + k * rh
        _t(d, s, x, ry, mw, rh, m, bold=True, anchor=MID_A)
        bh = min(0.34, rh * 0.5)
        bar = d._rect(s, fx(a), ry + (rh - bh) / 2, fx(b) - fx(a), bh, fill=GOLD if k == highlight else NAVY)
        _t(d, s, fx(a) - 0.75, ry, 0.72, rh, fmt.format(a), align=RGT, anchor=MID_A, wrap=False)
        _t(d, s, fx(b) + 0.04, ry, 0.75, rh, fmt.format(b), anchor=MID_A, wrap=False)
        _t(d, s, x + w - cw, ry, cw, rh, cm[0] if cm else "", anchor=MID_A)
        d._line(s, x, ry + rh, x + w, ry + rh, LGREY, 0.75)
        del bar
    for k, (lab, v) in enumerate(refs):
        col = GOLD if k == 0 else GREY
        d._line(s, fx(v), y + 0.3, fx(v), y + h, col, 1.5, dash=DASH)
        _t(d, s, fx(v) - 0.8, y + h + 0.0, 1.6, 0.22, f"{lab}: {fmt.format(v)}", bold=True, align=C, color=BLACK,
           wrap=False)


def bk17(d, s, x, y, w, h, rows, value_head=None):
    """Methodology overview: rows = [(method, description)] or (method, description, value); navy label boxes."""
    n = len(rows)
    gap = 0.12
    rh = min(0.95, (h - gap * (n - 1)) / n)
    lw = w * 0.24
    vw = w * 0.17 if any(len(r) > 2 for r in rows) else 0
    if vw and value_head:
        _t(d, s, x + w - vw, y, vw, 0.26, value_head, bold=True, align=C)
        d._line(s, x + w - vw, y + 0.28, x + w, y + 0.28, NAVY, 1.0)
        y, h = y + 0.38, h - 0.38
        rh = min(0.95, (h - gap * (n - 1)) / n)
    for k, r in enumerate(rows):
        ry = y + k * (rh + gap)
        b = d._rect(s, x, ry, lw, rh, fill=NAVY if k % 2 == 0 else MID)
        d._shape_text(b, r[0], size=F_BODY, bold=True)
        _t(d, s, x + lw + 0.15, ry, w - lw - vw - 0.3, rh, r[1], anchor=MID_A)
        if vw:
            _t(d, s, x + w - vw, ry, vw, rh, r[2] if len(r) > 2 else "", bold=True, align=C, anchor=MID_A, size=F_HEAD)
        d._line(s, x + lw + 0.15, ry + rh + gap / 2, x + w, ry + rh + gap / 2, LGREY, 0.75)


def bk18(d, s, x, y, w, h, steps, fmt="#,##0", side=None):
    """EV bridge waterfall: steps = [(label, value, kind)] with kind "total" (bar from 0) or "delta" (+ / -).
    side = optional bullets on the right."""
    sw = w * 0.3 if side else 0
    cw = w - sw - (0.3 if side else 0)
    base, up, run = [], [], 0.0
    for lab, v, kind in steps:
        if kind == "total":
            base.append(0)
            up.append(v)
            run = v
        else:
            lo = run + min(v, 0)
            base.append(lo)
            up.append(abs(v))
            run += v
    cats = [st[0] for st in steps]
    ch = d._chart(s, x, y, cw, h - 0.35, XL_CHART_TYPE.COLUMN_STACKED, cats, [("b", base), ("v", up)],
                  number_format=fmt, legend=False, gap=40)
    vmax = max(b + u for b, u in zip(base, up)) * 1.18
    _frame(ch, 0, vmax)
    ch.series[0].format.fill.background()
    ch.plots[0].has_data_labels = False
    fx, fy = _map(x, y, cw, h - 0.35, len(steps), 0, vmax)
    for k, (lab, v, kind) in enumerate(steps):
        pt = ch.series[1].points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(NAVY if kind == "total" else (STEEL if v >= 0 else LGREY))
        top = fy(base[k] + up[k])
        txt = f"{v:,.0f}" if kind == "total" else (f"+{v:,.0f}" if v >= 0 else f"({abs(v):,.0f})")
        _t(d, s, fx(k) - 0.6, top - 0.24, 1.2, 0.22, txt, bold=kind == "total", align=C, wrap=False)
        _t(d, s, fx(k) - cw / len(steps) / 2, y + h - 0.33, cw / len(steps), 0.33, lab, align=C)
    d._line(s, x, y + h - 0.35, x + cw, y + h - 0.35, LGREY, 0.75)
    if side:
        d.bullets(s, x + cw + 0.3, y, sw, h, side)


# ---------------------------------------------------------------- comps & deals
def _stats(vals):
    nums = [v for v in (_num(a) for a in vals) if v is not None]
    return (statistics.mean(nums), statistics.median(nums)) if nums else (None, None)


def bk19(d, s, x, y, w, h, columns, groups, fmts, target=None, overall=True, widths=None):
    """Multiples table with mean / median band: columns = header row; groups = [(group, [(name, [values])])]
    with numeric values (formatted by fmts per column, None -> "n.m."); mean / median rows per group shaded."""
    body, shade, bold, grp_rows, tgt = [], [], [], [], None
    every = []
    for gname, members in groups:
        if gname:
            body.append([gname] + [""] * (len(columns) - 1))
            grp_rows.append(len(body))
        for nm, vals in members:
            body.append([nm] + [("n.m." if v is None else f.format(v)) if f else str(v) for v, f in zip(vals, fmts)])
            if nm == target:
                tgt = len(body)
        cols_v = list(zip(*[vals for _, vals in members]))
        every.append(cols_v)
        for lab, fn in (("Mean", statistics.mean), ("Median", statistics.median)):
            row = [f"{gname} {lab.lower()}" if gname else lab]
            for cv, f in zip(cols_v, fmts):
                nums = [v for v in cv if isinstance(v, (int, float))]
                row.append(f.format(fn(nums)) if f and nums else "")
            body.append(row)
            shade.append(len(body))
            bold.append(len(body))
    if overall and len(groups) > 1:
        allv = [list(c) for c in zip(*[vals for _, ms in groups for _, vals in ms])]
        for lab, fn in (("Overall mean", statistics.mean), ("Overall median", statistics.median)):
            row = [lab]
            for cv, f in zip(allv, fmts):
                nums = [v for v in cv if isinstance(v, (int, float))]
                row.append(f.format(fn(nums)) if f and nums else "")
            body.append(row)
            shade.append(len(body))
            bold.append(len(body))
    rh = min(0.27, h / (len(body) + 1))
    cw = widths or [w * 0.26] + [w * 0.74 / (len(columns) - 1)] * (len(columns) - 1)
    shp = d.table(s, x, y, w, [list(columns)] + body, col_widths=cw, row_h=rh, shade_rows=shade, bold_rows=bold,
                  num_cols={j for j, f in enumerate(fmts, 1) if f})
    tbl = shp.table
    for i in grp_rows:
        _bold_cells(tbl, [i], [0])
        for r in tbl.cell(i, 0).text_frame.paragraphs[0].runs:
            r.font.color.rgb = rgb(NAVY)
    if tgt:
        _box(d, s, x, y + rh * tgt, w, rh, line=GOLD)
    return rh * (len(body) + 1)


def bk20(d, s, x, y, w, h, comps, implied, implied_head=None):
    """Comps + implied value range: comps = bk19 kwargs (columns, groups, fmts); implied = rows under a header,
    e.g. [["Metric", "Target metric", "Multiple range", "Implied EV", "Implied price"], ...]."""
    ih = 0.3 * (len(implied) + 0.5) + 0.4
    used = bk19(d, s, x, y, w, h - ih - 0.2, **comps)
    yy = y + used + 0.25
    yy = _head(d, s, x, yy, w, implied_head or "Implied valuation", "bar")
    shp = d.table(s, x, yy, w, implied, row_h=0.28)
    _bold_cells(shp.table, range(1, len(implied)), [len(implied[0]) - 1])


def bk21(d, s, x, y, w, h, columns, rows, bar_col, fmt="{:.1f}x", mean=True):
    """Precedent table with inline bars: rows = lists; rows[i][bar_col] is a number drawn as a bar in that
    column (with mean / median dashed lines)."""
    nc = len(columns)
    cw = [w * 0.1] + [w * 0.6 / (nc - 2)] * (nc - 2) + [w * 0.3]
    cw[bar_col] = w * 0.3 if bar_col == nc - 1 else cw[bar_col]
    rh = min(0.3, (h - 0.3) / (len(rows) + 1))
    body = [[("" if j == bar_col else c) for j, c in enumerate(r)] for r in rows]
    d.table(s, x, y, w, [list(columns)] + body, col_widths=cw, row_h=rh)
    bx = x + sum(cw[:bar_col]) + 0.08
    bw = cw[bar_col] - 0.7
    vals = [r[bar_col] for r in rows]
    vmax = max(vals) * 1.05
    for k, v in enumerate(vals):
        ry = y + rh * (k + 1)
        d._rect(s, bx, ry + rh * 0.22, bw * v / vmax, rh * 0.56, fill=NAVY)
        _t(d, s, bx + bw * v / vmax + 0.04, ry, 0.6, rh, fmt.format(v), anchor=MID_A, wrap=False)
    if mean:
        bottom = y + rh * (len(rows) + 1)
        for lab, v, col in (("Mean", statistics.mean(vals), GOLD), ("Median", statistics.median(vals), GREY)):
            lx = bx + bw * v / vmax
            d._line(s, lx, y + rh, lx, bottom, col, 1.25, dash=DASH)
        mv, md = statistics.mean(vals), statistics.median(vals)
        _t(d, s, bx, bottom + 0.03, cw[bar_col], 0.22, f"Mean {fmt.format(mv)} (gold)  ·  Median {fmt.format(md)} (grey)",
           bold=True, wrap=False)


def bk22(d, s, x, y, w, h, panels, target=None):
    """Ranked columns with median line: panels = [{"title", "items": [(name, value)], "fmt", "callout"}]
    stacked vertically; each sorted high -> low, target gold, dashed median."""
    n = len(panels)
    ph = (h - 0.2 * (n - 1)) / n
    for k, p in enumerate(panels):
        py = y + k * (ph + 0.2)
        yy = _head(d, s, x, py, w, p["title"])
        items = sorted(p["items"], key=lambda t: -t[1])
        names, vals = [i[0] for i in items], [i[1] for i in items]
        hi = names.index(target) if target in names else None
        cw = w - (2.0 if p.get("callout") else 0)
        fx, fy = _cols(d, s, x, yy, cw, py + ph - yy, names, vals, fmt=p.get("fmt", "0.0"), highlight=hi)
        med = statistics.median(vals)
        _median_line(d, s, x, fy(med), cw, f"Median: {p.get('label_fmt', '{:.1f}').format(med)}")
        if p.get("callout"):
            d.takeaway(s, x + cw + 0.15, yy + 0.2, 1.85, 1.2, p["callout"])


def bk23(d, s, x, y, w, h, rows, fmt="{:.1f}x", logos=None, stats_head="Summary statistics"):
    """Logo bars + side stats: rows = [(acquirer, target, value)] as horizontal bars with the two names (or logos)
    on the left; mean / median / high / low table on the right."""
    sw = w * 0.27
    bw = w - sw - 0.35
    nw = bw * 0.36
    n = len(rows)
    rh = min(0.42, h / n)
    vals = [r[2] for r in rows]
    vmax = max(vals) * 1.12
    logos = logos or {}
    for k, (acq, tgt, v) in enumerate(rows):
        ry = y + k * rh
        for j, nm in enumerate((acq, tgt)):
            lx = x + j * nw / 2
            if nm in logos:
                d._fit_picture(s, logos[nm], lx + 0.05, ry + 0.05, nw / 2 - 0.1, rh - 0.1)
            else:
                _t(d, s, lx, ry, nw / 2 - 0.05, rh, nm, bold=j == 1, anchor=MID_A)
        bx = x + nw + 0.1
        d._rect(s, bx, ry + rh * 0.2, (bw - nw - 0.7) * v / vmax, rh * 0.6, fill=NAVY)
        _t(d, s, bx + (bw - nw - 0.7) * v / vmax + 0.05, ry, 0.6, rh, fmt.format(v), anchor=MID_A, wrap=False)
        d._line(s, x, ry + rh, x + bw, ry + rh, PALEGREY, 0.5)
    mv = statistics.mean(vals)
    mx = x + nw + 0.1 + (bw - nw - 0.7) * mv / vmax
    d._line(s, mx, y, mx, y + n * rh, GOLD, 1.25, dash=DASH)
    _t(d, s, mx - 0.8, y + n * rh + 0.02, 1.6, 0.22, f"Mean: {fmt.format(mv)}", bold=True, align=C, wrap=False)
    px = x + bw + 0.35
    py = _head(d, s, px, y, sw, stats_head, "bar")
    st = [["Statistic", "Value"], ["Mean", fmt.format(mv)], ["Median", fmt.format(statistics.median(vals))],
          ["High", fmt.format(max(vals))], ["Low", fmt.format(min(vals))], ["Deals", str(n)]]
    d.table(s, px, py, sw, st, row_h=0.3)


def bk24(d, s, x, y, w, h, deals, fmt="0.0x", label_fmt="{:.1f}x", metric="EV / LTM EBITDA"):
    """Column chart + deal table: deals = [(date, acquirer, target, ev, value)]; chart of value by deal on top with
    the mean dashed, the deal table underneath in the same order."""
    ch_h = h * 0.45
    yy = _head(d, s, x, y, w, metric)
    vals = [dl[4] for dl in deals]
    fx, fy = _cols(d, s, x, yy, w, ch_h, [dl[2] for dl in deals], vals, fmt=fmt)
    mv = statistics.mean(vals)
    _median_line(d, s, x, fy(mv), w, f"Mean: {label_fmt.format(mv)}", color=GOLD)
    ty = yy + ch_h + 0.15
    rows = [["Date", "Acquirer", "Target", "EV", metric]] + [[a, b, c, e, label_fmt.format(v)] for a, b, c, e, v in deals]
    d.table(s, x, ty, w, rows, col_widths=[w * 0.12, w * 0.3, w * 0.3, w * 0.13, w * 0.15],
            row_h=min(0.26, (y + h - ty) / len(rows)), num_cols={3, 4})


def bk25(d, s, x, y, w, h, periods, series, offer=None, table=None, fmt="0%", offer_label="Implied by offer"):
    """Premiums paid: series = [(name, [premium per period])] as clustered columns (e.g. 25th pct / median / 75th
    pct); offer = [premium per period] drawn as gold markers; table = stats rows with a header on the right."""
    tw = w * 0.34 if table else 0
    cw = w - tw - (0.3 if table else 0)
    allv = [v for _, vs in series for v in vs] + list(offer or [])
    vmax = max(allv) * 1.3
    ch = d._chart(s, x, y, cw, h - 0.3, XL_CHART_TYPE.COLUMN_CLUSTERED, list(periods), series, number_format=fmt,
                  legend=True, colors=[LIGHT, NAVY, STEEL], gap=70)
    _frame(ch, 0, vmax)
    fx, fy = _map(x, y, cw, h - 0.3, len(periods), 0, vmax)
    for k, p in enumerate(periods):
        _t(d, s, fx(k) - cw / len(periods) / 2, y + h - 0.28, cw / len(periods), 0.26, p, align=C, bold=True)
    if offer:
        for k, v in enumerate(offer):
            hw = cw / len(periods) * 0.42
            d._line(s, fx(k) - hw, fy(v), fx(k) + hw, fy(v), GOLD, 2.25)
            _t(d, s, fx(k) - hw, fy(v) - 0.24, hw * 2, 0.22, f"{v:.0%}", bold=True, align=C, color=BLACK, wrap=False)
        _legend(d, s, x, y + h + 0.0, [offer_label], colors=[GOLD], w=2.5)
    if table:
        px = x + cw + 0.3
        py = _head(d, s, px, y, tw, "Premiums paid statistics", "bar")
        d.table(s, px, py, tw, table, row_h=0.3)


# ---------------------------------------------------------------- share price
def bk26(d, s, x, y, w, h, dates, prices, events, every=None, unit="$"):
    """Share price with lettered events: events = [(index into dates, date text, description)]; letters A, B, C on
    the line, the event list keyed to them on the right."""
    lw = w * 0.36
    cw = w - lw - 0.3
    fx, fy = _line_chart(d, s, x + 0.5, y, cw - 0.5, h, dates, [("Price", prices)], every=every)
    for k, (i, dt, txt) in enumerate(events):
        lt = chr(65 + k)
        cx, cy = fx(i), fy(prices[i])
        d._line(s, cx, cy, cx, cy - 0.3, GREY, 0.75)
        _circle_no(d, s, cx - 0.12, cy - 0.54, lt, dia=0.24, fill=GOLD if k == len(events) - 1 else NAVY)
    ex = x + cw + 0.3
    rh = min(0.62, h / len(events))
    for k, (i, dt, txt) in enumerate(events):
        ry = y + k * rh
        _circle_no(d, s, ex, ry + 0.02, chr(65 + k), dia=0.24, fill=GOLD if k == len(events) - 1 else NAVY)
        _t(d, s, ex + 0.32, ry, lw - 0.32, rh, f"**{dt}:** {txt}")


def bk27(d, s, x, y, w, h, dates, prices, callouts, volume=None, ref=None, every=None):
    """Price chart with callout boxes: callouts = [(index, text)] in boxes along the top joined to their point;
    volume = optional bars underneath; ref = (label, value) dashed line (offer price)."""
    top = 0.95
    vh = h * 0.18 if volume else 0
    fx, fy = _line_chart(d, s, x + 0.5, y + top, w - 0.5, h - top - vh, dates, [("Price", prices)], every=every,
                         extra=[ref[1]] if ref else ())
    if ref:
        d._line(s, x + 0.5, fy(ref[1]), x + w, fy(ref[1]), GOLD, 1.25, dash=DASH)
        _t(d, s, x + w - 2.2, fy(ref[1]) - 0.24, 2.2, 0.22, f"{ref[0]}: {ref[1]:,.2f}", bold=True, align=RGT, wrap=False)
    n = len(callouts)
    bw = min(2.1, (w - 0.5) / n - 0.1)
    for k, (i, txt) in enumerate(sorted(callouts)):
        bx = x + 0.5 + k * (w - 0.5) / n
        _box(d, s, bx, y, bw, top - 0.15, txt, line=NAVY, size=F_BODY)
        px, py = fx(i), fy(prices[i])
        d._line(s, bx + bw / 2, y + top - 0.15, px, py, GREY, 0.75)
        d._shape(s, MSO_SHAPE.OVAL, px - 0.05, py - 0.05, 0.1, 0.1, fill=NAVY)
    if volume:
        vy = y + h - vh
        ch = d._chart(s, x + 0.5, vy, w - 0.5, vh, XL_CHART_TYPE.COLUMN_CLUSTERED, list(dates), [("Volume", volume)],
                      labels=False, legend=False, colors=[LIGHT], gap=20)
        _frame(ch, 0, max(volume) * 1.05)
        _t(d, s, x - 0.1, vy, 0.6, 0.3, "Volume", color=GREY, wrap=False)


def bk28(d, s, x, y, w, h, dates, prices, stats, band=None, refs=(), every=None, stats_head="Price statistics"):
    """Price chart + stats table: stats = rows with a header (e.g. ["Measure", "Price", "Offer premium"]);
    band = (low, high, label) shaded offer range; refs = [(label, value)] dashed lines."""
    tw = w * 0.36
    cw = w - tw - 0.3
    fx, fy = _line_chart(d, s, x + 0.5, y + 0.3, cw - 0.5, h - 0.3, dates, [("Price", prices)], every=every,
                         extra=list(band[:2] if band else ()) + [v for _, v in refs])
    if band:
        lo, hi, lab = band
        r = d._rect(s, x + 0.5, fy(hi), cw - 0.5, fy(lo) - fy(hi), fill=GOLD)
        d._alpha(r.fill._xPr.find(qn("a:solidFill")), 25)
        _t(d, s, x + 0.55, fy(hi) - 0.22, 3, 0.2, lab, bold=True, wrap=False)
    for k, (lab, v) in enumerate(refs):
        d._line(s, x + 0.5, fy(v), x + cw, fy(v), GREY, 1.0, dash=DASH)
        _t(d, s, x + cw - 2.0, fy(v) + 0.01, 2.0, 0.2, f"{lab}: {v:,.2f}", align=RGT, wrap=False)
    px = x + cw + 0.3
    py = _head(d, s, px, y, tw, stats_head, "bar")
    d.table(s, px, py, tw, stats, row_h=min(0.3, (h - 0.4) / len(stats)))


def bk29(d, s, x, y, w, h, dates, series, every=None, periods=None):
    """Indexed performance + table: series = [(name, prices)] rebased to 100 on the first date; table of total
    return per series (periods = [(label, start index)] adds columns, default since start)."""
    tw = w * 0.33
    cw = w - tw - 0.3
    reb = [(n, [v / vs[0] * 100 for v in vs]) for n, vs in series]
    _line_chart(d, s, x + 0.5, y + 0.3, cw - 0.5, h - 0.3, dates, reb, every=every)
    periods = periods or [("Since " + str(dates[0]), 0)]
    rows = [["Company"] + [p for p, _ in periods]]
    for n, vs in series:
        rows.append([n] + [f"{vs[-1] / vs[i] - 1:+.1%}" for _, i in periods])
    px = x + cw + 0.3
    py = _head(d, s, px, y, tw, "Share price performance", "bar")
    shp = d.table(s, px, py, tw, rows, row_h=0.3)
    _bold_cells(shp.table, [1], range(len(rows[0])))
    _shade_cells(shp.table, [1], range(len(rows[0])), PALE)


def bk30(d, s, x, y, w, h, bands, volumes, current=None, unit="Volume (m)"):
    """Volume at price: bands = price band labels, volumes = volume traded in each; current band boxed; a table
    under the chart with volume, % of total and cumulative %, columns aligned to the bars."""
    lw = 1.1
    cw = w - lw
    tot = sum(volumes)
    ch_h = h - 0.85
    fx, fy = _cols(d, s, x + lw, y, cw, ch_h, bands, volumes, fmt="#,##0.0", highlight=current, cat_labels=False)
    n = len(bands)
    if current is not None:
        _box(d, s, x + lw + current * cw / n, y, cw / n, ch_h, line=GOLD, dash=DASH)
        _t(d, s, x + lw + current * cw / n - 0.6, y - 0.24, cw / n + 1.2, 0.22, "Current trading range", bold=True,
           align=C, wrap=False)
    cum = 0
    rows = [["Price band"] + list(bands), [unit] + [f"{v:,.1f}" for v in volumes], ["% of total"], ["Cumulative %"]]
    for v in volumes:
        cum += v
        rows[2].append(f"{v / tot:.0%}")
        rows[3].append(f"{cum / tot:.0%}")
    d.table(s, x, y + ch_h, w, rows, col_widths=[lw] + [cw / n] * n, row_h=0.21, size=8 if n > 10 else F_BODY,
            num_cols=set(range(1, n + 1)))


def bk31(d, s, x, y, w, h, market, holders, dates, prices, every=None):
    """Public market overview: market data (label, value) and top shareholders table on top, price chart below."""
    th = min(h * 0.5, 0.26 * (max(len(market), len(holders) - 1) + 1) + 0.4)
    lw = w * 0.4
    yy = _head(d, s, x, y, lw, "Public market overview", "bar")
    d.table(s, x, yy, lw, [["", ""]] + [list(r) for r in market], col_widths=[lw * 0.62, lw * 0.38], row_h=0.25,
            header=False)
    rx, rw = x + lw + 0.3, w - lw - 0.3
    yy2 = _head(d, s, rx, y, rw, "Top shareholders", "bar")
    d.table(s, rx, yy2, rw, holders, row_h=0.25)
    py = y + th + 0.2
    py = _head(d, s, x, py, w, "Share price performance")
    _line_chart(d, s, x + 0.5, py, w - 0.5, y + h - py, dates, [("Price", prices)], every=every)


# ---------------------------------------------------------------- benchmarking
def bk32(d, s, x, y, w, h, companies, metrics, target=None):
    """Benchmarking bar matrix: metrics = [(name, [value per company], fmt)]; one column of horizontal bars per
    metric, companies as rows, target gold."""
    nw = w * 0.17
    m = len(metrics)
    mw = (w - nw) / m
    n = len(companies)
    rh = (h - 0.4) / n
    for j, (name, vals, f) in enumerate(metrics):
        mx = x + nw + j * mw
        _t(d, s, mx, y, mw - 0.1, 0.34, name, bold=True, align=C, anchor=MSO_ANCHOR.BOTTOM)
        vmax = max(abs(v) for v in vals) or 1
        for i, v in enumerate(vals):
            ry = y + 0.4 + i * rh
            bl = (mw - 0.75) * max(v, 0) / vmax
            d._rect(s, mx + 0.08, ry + rh * 0.2, max(bl, 0.02), rh * 0.6,
                    fill=GOLD if companies[i] == target else NAVY)
            _t(d, s, mx + 0.1 + bl, ry, 0.65, rh, f.format(v), anchor=MID_A, wrap=False)
        if j:
            d._line(s, mx, y + 0.4, mx, y + h, LGREY, 0.75)
    d._line(s, x, y + 0.38, x + w, y + 0.38, NAVY, 1.0)
    for i, c in enumerate(companies):
        _t(d, s, x, y + 0.4 + i * rh, nw, rh, c, bold=c == target, anchor=MID_A)


def bk33(d, s, x, y, w, h, peers, metrics, target=None):
    """Peer small multiples: metrics = [(name, [value per peer], fmt, label_fmt)] stacked as column-chart rows with
    a navy side tab, the target gold and a dashed median; peer names under the last row only."""
    m = len(metrics)
    tab = 0.32
    rh = (h - 0.35 - 0.1 * (m - 1)) / m
    for k, (name, vals, f, lf) in enumerate(metrics):
        ry = y + k * (rh + 0.1)
        _vtab(d, s, x, ry, tab, rh, name, fill=RAMP[k % 3])
        hi = peers.index(target) if target in peers else None
        fx, fy = _cols(d, s, x + tab + 0.1, ry, w - tab - 0.1, rh, peers, vals, fmt=f, highlight=hi, cat_labels=False,
                       gap=80)
        _median_line(d, s, x + tab + 0.1, fy(statistics.median(vals)), w - tab - 0.1,
                     f"Median: {lf.format(statistics.median(vals))}")
    n = len(peers)
    for i, p in enumerate(peers):
        _t(d, s, x + tab + 0.1 + i * (w - tab - 0.1) / n, y + h - 0.33, (w - tab - 0.1) / n, 0.33, p, align=C,
           bold=p == target)


def bk35(d, s, x, y, w, h, entities, rows, highlight=None):
    """Comparison matrix: one column per entity (header box, highlight filled navy), rows = [(attribute, [values])];
    a value 0-4 (int) draws a Harvey ball, True a tick, str as text."""
    lw = w * 0.2
    n = len(entities)
    cw = (w - lw) / n
    hh = 0.45
    for j, e in enumerate(entities):
        b = d._shape(s, MSO_SHAPE.RECTANGLE, x + lw + j * cw + 0.04, y, cw - 0.08, hh, fill=NAVY if j == highlight else None,
                     line=NAVY)
        d._shape_text(b, e, size=F_BODY, bold=True, color=WHITE if j == highlight else NAVY)
    rh = (h - hh - 0.1) / len(rows)
    for i, (attr, vals) in enumerate(rows):
        ry = y + hh + 0.1 + i * rh
        _t(d, s, x, ry, lw, rh, attr, bold=True, anchor=MID_A)
        for j, v in enumerate(vals):
            cx = x + lw + j * cw
            if isinstance(v, bool):
                _t(d, s, cx, ry, cw, rh, "✓" if v else "–", bold=True, align=C, anchor=MID_A, size=F_HEAD,
                   color=NAVY if v else GREY)
            elif isinstance(v, int):
                d._harvey_ball(s, cx + cw / 2 - 0.13, ry + rh / 2 - 0.13, 0.26, v)
            else:
                _t(d, s, cx + 0.06, ry, cw - 0.12, rh, str(v), align=C, anchor=MID_A)
        d._line(s, x, ry + rh, x + w, ry + rh, LGREY, 0.75)
    if highlight is not None:
        _box(d, s, x + lw + highlight * cw + 0.04, y + hh, cw - 0.08, h - hh, line=NAVY)


# ---------------------------------------------------------------- dashboards
def _panel_chart(d, s, x, y, w, h, p):
    kind = p.get("kind", "column")
    fmt = p.get("fmt", "#,##0")
    if kind == "fin_chart":
        d.fin_chart(s, x, y, w, h, p["cats"], p["series"][0][1], strip=p.get("strip"), cagr=p.get("cagr"),
                    split=p.get("split"), number_format=fmt, label_w=0.6 if p.get("strip") else 0.0)
        return
    if kind == "donut":
        dia = min(w * 0.55, h)
        _donut(d, s, x, y + (h - dia) / 2, dia, p["cats"], p["series"][0][1], centre=p.get("centre"))
        _legend(d, s, x + dia + 0.15, y + (h - dia) / 2 + 0.1,
                [f"{c} ({v:.0%})" if v <= 1 else f"{c} ({v:,.0f})" for c, v in zip(p["cats"], p["series"][0][1])],
                w=w - dia - 0.3)
        return
    fn = {"column": d.column, "bar": d.bar, "stacked": d.stacked, "line": d.line}[kind]
    kw = {"number_format": fmt}
    if p.get("highlight") is not None:
        kw["highlight"] = p["highlight"]
    if kind == "line":
        kw["labels"] = p.get("labels", False)
    fn(s, x, y, w, h, p["cats"], p["series"], **kw)


def bk36(d, s, x, y, w, h, panels, cols=None):
    """Chart dashboard: panels = [{"title", "kind" (column / bar / stacked / line / fin_chart / donut), "cats",
    "series": [(name, values)], "fmt", "highlight", "strip", "cagr", "split"}] in a 2x2 (or cols) grid."""
    n = len(panels)
    cols = cols or (2 if n in (2, 4) else 3 if n in (3, 5, 6) else 2)
    rows = -(-n // cols)
    pw, ph = (w - 0.3 * (cols - 1)) / cols, (h - 0.2 * (rows - 1)) / rows
    for k, p in enumerate(panels):
        px, py = x + (k % cols) * (pw + 0.3), y + (k // cols) * (ph + 0.2)
        yy = _head(d, s, px, py, pw, p["title"])
        _panel_chart(d, s, px, yy, pw, py + ph - yy, p)


def bk37(d, s, x, y, w, h, description, periods, lines, chart, mix=None):
    """Company at a glance: description bullets (top left), financial highlights fin_table (top right), a chart
    panel (bottom left, bk36 panel dict) and mix donuts (bottom right, [(title, labels, values)])."""
    lw = w * 0.48
    rw = w - lw - 0.3
    th = h * 0.5
    yy = _head(d, s, x, y, lw, "Business description", "bar")
    d.bullets(s, x, yy + 0.05, lw, th - 0.45, description)
    rx = x + lw + 0.3
    yy = _head(d, s, rx, y, rw, "Financial highlights", "bar")
    d.fin_table(s, rx, yy, rw, periods, lines, label_w=rw * 0.34, row_h=min(0.24, (th - 0.45) / (len(lines) + 1)))
    by = y + th + 0.1
    yy = _head(d, s, x, by, lw, chart["title"], "bar")
    _panel_chart(d, s, x, yy + 0.05, lw, y + h - yy - 0.05, chart)
    if mix:
        yy = _head(d, s, rx, by, rw, "Business mix", "bar")
        mw = rw / len(mix)
        for k, (title, labs, vals) in enumerate(mix):
            mx = rx + k * mw
            _t(d, s, mx, yy, mw, 0.24, title, bold=True, align=C)
            dia = min(mw - 0.3, (y + h - yy) * 0.55)
            _donut(d, s, mx + (mw - dia) / 2, yy + 0.28, dia, labs, vals)
            _legend(d, s, mx + 0.1, yy + 0.35 + dia, [f"{a} {b:.0%}" for a, b in zip(labs, vals)], w=mw - 0.3, row=0.2)


def bk38(d, s, x, y, w, h, periods, lines, kpis, split=None):
    """Wide table + KPI mini charts: fin_table over ~65% width, 2-3 small fin_charts stacked on the right.
    kpis = [{"title", "cats", "values", "fmt"}]."""
    tw = w * 0.63
    rh = min(0.26, h / (len(lines) + 1))
    d.fin_table(s, x, y, tw, periods, lines, label_w=tw * 0.32, row_h=rh)
    if split is not None:
        sx = x + tw * 0.32 + split * tw * 0.68 / len(periods)
        d._line(s, sx, y, sx, y + rh * (len(lines) + 1), GREY, 1.0, dash=DASH)
    kx, kw = x + tw + 0.35, w - tw - 0.35
    n = len(kpis)
    kh = (h - 0.15 * (n - 1)) / n
    for k, p in enumerate(kpis):
        ky = y + k * (kh + 0.15)
        yy = _head(d, s, kx, ky, kw, p["title"])
        d.fin_chart(s, kx, yy, kw, ky + kh - yy, p["cats"], p["values"], number_format=p.get("fmt", "#,##0"),
                    label_w=0.0, split=p.get("split"))


def bk39(d, s, x, y, w, h, capital, maturities, mix=None, notes=None):
    """Debt / liquidity profile: capital structure table (rows with a header) top left, maturity columns
    (maturities = (years, [(instrument, values)])) bottom left, fixed / floating or currency donuts and notes on
    the right (mix = [(title, labels, values)])."""
    lw = w * 0.62 if (mix or notes) else w
    rh = 0.25
    yy = _head(d, s, x, y, lw, "Capital structure", "bar")
    shp = d.table(s, x, yy, lw, capital, row_h=rh)
    _bold_cells(shp.table, [len(capital) - 1], range(len(capital[0])))
    my = yy + rh * len(capital) + 0.2
    my = _head(d, s, x, my, lw, "Debt maturity profile")
    years, ser = maturities
    fn = d.stacked if len(ser) > 1 else d.column
    ch = fn(s, x, my, lw, y + h - my, years, ser, number_format="#,##0;-#,##0;;", colors=[NAVY, STEEL, LIGHT, LGREY])
    ch.value_axis.minimum_scale = 0
    if mix or notes:
        rx, rw = x + lw + 0.3, w - lw - 0.3
        ry = y
        for title, labs, vals in (mix or []):
            ry = _head(d, s, rx, ry, rw, title)
            dia = min(1.5, rw * 0.45)
            _donut(d, s, rx, ry, dia, labs, vals)
            _legend(d, s, rx + dia + 0.15, ry + 0.2, [f"{a} {b:.0%}" for a, b in zip(labs, vals)], w=rw - dia - 0.2)
            ry += dia + 0.15
        if notes:
            d.bullets(s, rx, ry, rw, y + h - ry, notes)


def bk40(d, s, x, y, w, h, columns, rows, labels):
    """Mix donuts before / after: columns = entity names (Target, Acquirer, Pro forma); rows = [(metric,
    [[shares per label] per column])]; labels = the segment names (one legend per page)."""
    lw = 1.3
    n = len(columns)
    cw = (w - lw - 1.8) / n
    rh = (h - 0.35) / len(rows)
    for j, c in enumerate(columns):
        _t(d, s, x + lw + j * cw, y, cw, 0.3, c, bold=True, align=C, color=NAVY if j == n - 1 else BLACK)
    d._line(s, x, y + 0.32, x + w - 1.8, y + 0.32, NAVY, 1.0)
    for i, (metric, vals) in enumerate(rows):
        ry = y + 0.38 + i * rh
        _t(d, s, x, ry, lw, rh, metric, bold=True, anchor=MID_A)
        dia = min(cw - 0.2, rh - 0.1)
        for j, v in enumerate(vals):
            _donut(d, s, x + lw + j * cw + (cw - dia) / 2, ry + (rh - dia) / 2, dia, labels, v,
                   centre=f"{max(v):.0%}")
        if i < len(rows) - 1:
            d._line(s, x, ry + rh, x + w - 1.8, ry + rh, LGREY, 0.75)
    _legend(d, s, x + w - 1.65, y + 0.45, labels, w=1.6)


def bk41(d, s, x, y, w, h, panels):
    """Market dynamics: three panels across, each = header bar, bold finding, bullets and a chart (bk36 panel dict),
    with chevrons between the headers (cause -> effect)."""
    n = len(panels)
    g = 0.4
    pw = (w - g * (n - 1)) / n
    for k, p in enumerate(panels):
        px = x + k * (pw + g)
        yy = _head(d, s, px, y, pw, p["title"], "bar")
        if k < n - 1:
            ar = d._shape(s, MSO_SHAPE.CHEVRON, px + pw + 0.1, y + 0.02, 0.2, 0.26, fill=GOLD)
            del ar
        d.exhibit_takeaway(s, px, yy, pw, p["finding"], h=0.5)
        by = yy + 0.55
        bh = min(1.4, 0.36 * len(p["bullets"]) + 0.2)
        d.bullets(s, px, by, pw, bh, p["bullets"])
        _panel_chart(d, s, px, by + bh + 0.05, pw, y + h - by - bh - 0.05, p["chart"])


# ---------------------------------------------------------------- structure
def _su_rows(items, unit):
    tot = sum(v for _, v in items)
    rows = [[lab, f"{v:,.0f}", f"{v / tot:.1%}"] for lab, v in items]
    return rows + [["Total", f"{tot:,.0f}", "100.0%"]]


def bk42(d, s, x, y, w, h, sources, uses, unit="$m", captable=None, notes=None):
    """Sources & uses: sources / uses = [(label, amount)] (% and totals added); captable = rows with a header for a
    pro forma capitalisation table below (or notes = bullets below)."""
    gw = (w - 0.3) / 2
    n = max(len(sources), len(uses)) + 2
    rh = 0.28
    for k, (title, items) in enumerate((("Sources", sources), ("Uses", uses))):
        gx = x + k * (gw + 0.3)
        yy = _head(d, s, gx, y, gw, f"{title} ({unit})", "bar")
        rows = [["", "Amount", "%"]] + _su_rows(items, unit)
        d.table(s, gx, yy, gw, rows, col_widths=[gw * 0.6, gw * 0.2, gw * 0.2], row_h=rh, bold_rows=(len(rows) - 1,))
    by = y + 0.38 + rh * n + 0.25
    if captable:
        by = _head(d, s, x, by, w, "Pro forma capitalisation", "bar")
        d.table(s, x, by, w, captable, row_h=min(0.26, (y + h - by) / len(captable)))
    elif notes:
        d.bullets(s, x, by, w, y + h - by, notes)


def bk43(d, s, x, y, w, h, summary, sources, uses, price, sens, unit="$m"):
    """LBO summary page: transaction summary bullets | sources & uses (top); purchase price table | returns
    sensitivity grid (bottom). price = rows with a header; sens = bk44 grid dict."""
    gw = (w - 0.3) / 2
    th = h * 0.5
    yy = _head(d, s, x, y, gw, "Transaction summary", "bar")
    d.bullets(s, x, yy + 0.05, gw, th - 0.45, summary)
    rx = x + gw + 0.3
    yy = _head(d, s, rx, y, gw, f"Sources & uses ({unit})", "bar")
    hw = (gw - 0.15) / 2
    for k, (title, items) in enumerate((("Sources", sources), ("Uses", uses))):
        rows = [[title, "Amount"]] + [[a, b, ] for a, b, _ in _su_rows(items, unit)]
        d.table(s, rx + k * (hw + 0.15), yy, hw, rows, col_widths=[hw * 0.66, hw * 0.34], row_h=0.24,
                bold_rows=(len(rows) - 1,), header_style="rule")
    by = y + th + 0.1
    yy = _head(d, s, x, by, gw, "Purchase price", "bar")
    d.table(s, x, yy, gw, price, row_h=min(0.26, (y + h - yy) / len(price)), header_style="rule")
    yy = _head(d, s, rx, by, gw, sens["title"], "bar")
    d.heatmap(s, rx, yy, gw, sens.get("corner", ""), sens["cols"], sens["rows"], sens["values"],
              fmt=sens.get("fmt", "{:.1%}"), base=sens.get("base"),
              row_h=min(0.27, (y + h - yy) / (len(sens["rows"]) + 1)))


def bk44(d, s, x, y, w, h, grids, note=None):
    """Twin sensitivity tables: grids = [{"title", "corner", "cols", "rows", "values", "fmt", "base", "better",
    "xlabel", "ylabel"}] side by side, each under a header bar with its driver labels; note = takeaway below."""
    n = len(grids)
    gw = (w - 0.4 * (n - 1)) / n
    nh = 0.55 if note else 0
    for k, g in enumerate(grids):
        gx = x + k * (gw + 0.4)
        yy = _head(d, s, gx, y, gw, g["title"], "bar")
        lab = 0.3 if g.get("ylabel") else 0
        if g.get("xlabel"):
            _t(d, s, gx + lab, yy, gw - lab, 0.24, g["xlabel"], bold=True, align=C, color=NAVY)
            yy += 0.26
        rh = min(0.32, (y + h - nh - yy) / (len(g["rows"]) + 1))
        if lab:
            _vtab(d, s, gx, yy + rh, lab - 0.05, rh * len(g["rows"]), g["ylabel"], fill=PALE, color=NAVY)
        d.heatmap(s, gx + lab, yy, gw - lab, g.get("corner", ""), g["cols"], g["rows"], g["values"],
                  fmt=g.get("fmt", "{:.1%}"), base=g.get("base"), better=g.get("better", "high"), row_h=rh)
    if note:
        d.takeaway(s, x, y + h - 0.48, w, 0.48, note)


def bk45(d, s, x, y, w, h, assumptions, periods, standalone, proforma, metric="EPS", callout=None, fmt="${:.2f}"):
    """Accretion / dilution: assumptions bullets left; table right with standalone and pro forma per period, the
    $ and % accretion / (dilution) computed; the first year's % boxed in gold with a callout underneath."""
    lw = w * 0.34
    yy = _head(d, s, x, y, lw, "Key assumptions", "bar")
    d.bullets(s, x, yy + 0.05, lw, h - 0.45, assumptions)
    rx, rw = x + lw + 0.35, w - lw - 0.35
    yy = _head(d, s, rx, y, rw, f"Pro forma {metric} impact", "bar")
    acc = [p - a for a, p in zip(standalone, proforma)]
    pct = [(p - a) / a for a, p in zip(standalone, proforma)]
    pc = lambda v: f"({abs(v):.1%})" if v < 0 else f"{v:.1%}"   # noqa: E731
    am = lambda v: f"({fmt.format(abs(v))})" if v < 0 else fmt.format(v)   # noqa: E731
    rows = [[""] + list(periods), [f"Standalone {metric}"] + [fmt.format(v) for v in standalone],
            [f"Pro forma {metric}"] + [fmt.format(v) for v in proforma], ["Accretion / (dilution) ($)"] + [am(v) for v in acc],
            ["Accretion / (dilution) (%)"] + [pc(v) for v in pct]]
    rh = 0.34
    cw = [rw * 0.36] + [rw * 0.64 / len(periods)] * len(periods)
    d.table(s, rx, yy, rw, rows, col_widths=cw, row_h=rh, bold_rows=(4,), header_style="rule")
    _box(d, s, rx + cw[0], yy + rh * 4, cw[1], rh, line=GOLD)
    if callout:
        cy = yy + rh * 5 + 0.3
        d.takeaway(s, rx, cy, rw, min(0.8, y + h - cy), callout)


def bk46(d, s, x, y, w, h, periods, lines, highlight=(), split=None, notes=None, labels=("Historical", "Projected")):
    """Credit statistics: historical | projected table (bk08) with the key leverage rows boxed in gold, bullets
    underneath. highlight = 0-based line indexes."""
    nh = min(1.4, 0.38 * len(notes or [])) if notes else 0
    rh = min(0.26, (h - nh - 0.45) / (len(lines) + 1))
    bk08(d, s, x, y, w, h - nh, periods, lines, split if split is not None else len(periods), labels=labels)
    for i in highlight:
        _box(d, s, x - 0.03, y + 0.32 + rh * (i + 1), w + 0.06, rh, line=GOLD)
    if notes:
        d.bullets(s, x, y + 0.32 + rh * (len(lines) + 1) + 0.25, w, nh, notes)


def bk47(d, s, x, y, w, h, tranches):
    """Financing term sheet: tranches = [(name, [(term, value)])] each with a vertical navy tab and term rows."""
    tab = 0.4
    nrows = sum(len(t) for _, t in tranches)
    rh = min(0.3, (h - 0.08 * len(tranches)) / nrows)
    yy = y
    for k, (name, terms) in enumerate(tranches):
        th = rh * len(terms)
        _vtab(d, s, x, yy, tab - 0.05, th, name, fill=RAMP[k % 3])
        d.table(s, x + tab, yy, w - tab, [[a, b] for a, b in terms], col_widths=[(w - tab) * 0.22, (w - tab) * 0.78],
                row_h=rh, header=False, num_cols=set())
        tbl = s.shapes[-1].table
        _bold_cells(tbl, range(len(terms)), [0])
        yy += th + 0.08


def bk48(d, s, x, y, w, h, options, rows, xlabel="Share price at closing", ylabel="Value received"):
    """Structure options matrix: options = [{"name", "x": [...], "y": [...]}] each a column with a mini payoff
    line chart; rows = [(attribute, [text per option])] underneath."""
    lw = w * 0.17
    n = len(options)
    cw = (w - lw) / n
    hh = 0.32
    chh = 1.35
    for j, o in enumerate(options):
        cx = x + lw + j * cw
        b = d._rect(s, cx + 0.05, y, cw - 0.1, hh, fill=NAVY)
        d._shape_text(b, o["name"], size=F_BODY, bold=True)
        ch = d._chart(s, cx + 0.1, y + hh + 0.05, cw - 0.2, chh, XL_CHART_TYPE.LINE, [str(v) for v in o["x"]],
                      [("v", o["y"])], labels=False, legend=False, colors=[NAVY])
        ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        ch.category_axis.format.line.color.rgb = rgb(LGREY)
    _t(d, s, x, y + hh + 0.05, lw - 0.1, chh, f"Illustrative payoff\n({ylabel} vs {xlabel.lower()})", color=GREY,
       anchor=MID_A, italic=True)
    ry0 = y + hh + chh + 0.15
    d._line(s, x, ry0, x + w, ry0, NAVY, 1.0)
    rh = (y + h - ry0) / len(rows)
    for i, (attr, vals) in enumerate(rows):
        ry = ry0 + i * rh
        _t(d, s, x, ry, lw, rh, attr, bold=True, anchor=MID_A)
        for j, v in enumerate(vals):
            _t(d, s, x + lw + j * cw + 0.08, ry, cw - 0.16, rh, v, align=C, anchor=MID_A)
        d._line(s, x, ry + rh, x + w, ry + rh, LGREY, 0.75)


# ---------------------------------------------------------------- process
def bk49(d, s, x, y, w, h, stages, cols=("Strategic", "Financial"), note=None):
    """Buyer funnel: stages = [(label, total, [count per col])] as a narrowing funnel (counts inside) with the stage
    labels and a strategic / financial split table on the right."""
    fw = w * 0.42
    tw = w * 0.22
    lw = w - fw - tw - 0.2
    n = len(stages)
    nh = 0.5 if note else 0
    rh = (h - 0.35 - nh) / n
    tx = x + fw + 0.1 + lw + 0.1
    for j, c in enumerate(cols):
        _t(d, s, tx + j * tw / len(cols), y, tw / len(cols), 0.3, c, bold=True, align=C)
    d._line(s, tx, y + 0.32, tx + tw, y + 0.32, NAVY, 1.0)
    for k, (lab, tot, split) in enumerate(stages):
        ry = y + 0.35 + k * rh
        top_w = fw * (1 - k / n * 0.75)
        bot_w = fw * (1 - (k + 1) / n * 0.75)
        tr = d._shape(s, MSO_SHAPE.TRAPEZOID, x + (fw - top_w) / 2, ry + 0.03, top_w, rh - 0.06,
                      fill=RAMP[min(k, 3)] if k < n - 1 else GOLD)
        tr.rotation = 180
        tr.adjustments[0] = (top_w - bot_w) / 2 / (rh - 0.06) if rh > 0.06 else 0.25
        _t(d, s, x, ry, fw, rh, str(tot), bold=True, size=F_HEAD, align=C, anchor=MID_A,
           color=WHITE if k < 3 or k == n - 1 else BLACK)
        _t(d, s, x + fw + 0.1, ry, lw, rh, lab, bold=True, anchor=MID_A)
        d._line(s, x + fw / 2 + top_w / 2, ry + rh / 2, x + fw + 0.08, ry + rh / 2, LGREY, 0.75, dash=DASH)
        for j, v in enumerate(split):
            _t(d, s, tx + j * tw / len(cols), ry, tw / len(cols), rh, str(v), align=C, anchor=MID_A)
        d._line(s, tx, ry + rh, tx + tw, ry + rh, LGREY, 0.75)
    if note:
        d.takeaway(s, x, y + h - 0.45, w, 0.45, note)


def bk50(d, s, x, y, w, h, months, marks, legend, events=None):
    """Calendar: months = [(year, month)] as small month grids; marks = {"YYYY-MM-DD": legend index};
    legend = [(label, colour)]; events = [(date text, event)] table on the right."""
    ew = w * 0.36 if events else 0
    cw_all = w - ew - (0.3 if events else 0)
    n = len(months)
    per_row = min(n, 3)
    nrow = -(-n // per_row)
    mw = (cw_all - 0.25 * (per_row - 1)) / per_row
    mh = (h - 0.4 - 0.2 * (nrow - 1)) / nrow
    calendar.setfirstweekday(calendar.SUNDAY)
    for k, (yr, mo) in enumerate(months):
        mx, my = x + (k % per_row) * (mw + 0.25), y + (k // per_row) * (mh + 0.2)
        _t(d, s, mx, my, mw, 0.26, f"{calendar.month_name[mo]} {yr}", bold=True, color=NAVY)
        d._line(s, mx, my + 0.27, mx + mw, my + 0.27, NAVY, 1.0)
        cell = mw / 7
        weeks = calendar.monthcalendar(yr, mo)
        rh = (mh - 0.55) / len(weeks)
        for j, dn in enumerate("SMTWTFS"):
            _t(d, s, mx + j * cell, my + 0.3, cell, 0.22, dn, bold=True, align=C)
        for i, wk in enumerate(weeks):
            for j, day in enumerate(wk):
                if not day:
                    continue
                cx, cy = mx + j * cell, my + 0.55 + i * rh
                key = f"{yr}-{mo:02d}-{day:02d}"
                if key in marks:
                    col = legend[marks[key]][1]
                    b = d._rect(s, cx + 0.02, cy + 0.01, cell - 0.04, rh - 0.02, fill=col)
                    d._shape_text(b, str(day), size=F_BODY, bold=True, color=WHITE if col != LIGHT else BLACK)
                else:
                    _t(d, s, cx, cy, cell, rh, str(day), align=C, anchor=MID_A, color=GREY if j in (0, 6) else BLACK)
    lx = x
    for lab, col in legend:
        d._rect(s, lx, y + h - 0.22, 0.16, 0.16, fill=col)
        _t(d, s, lx + 0.22, y + h - 0.28, 2.2, 0.26, lab, anchor=MID_A, wrap=False)
        lx += 0.4 + len(lab) * 0.075
    if events:
        ex = x + cw_all + 0.3
        yy = _head(d, s, ex, y, ew, "Key dates", "bar")
        d.table(s, ex, yy, ew, [["Date", "Event"]] + [list(e) for e in events], col_widths=[ew * 0.28, ew * 0.72],
                row_h=min(0.4, (h - 0.4) / (len(events) + 1)), num_cols=set())


def bk51(d, s, x, y, w, h, periods, rows, milestones=(), today=None):
    """Gantt: periods = column labels (weeks / months); rows = [(workstream, start, end)] bar from start to end
    (period indexes, end inclusive; a 4th item "group" makes a bold group row); milestones = [(row, at, label)]
    gold diamonds; today = period index for a dashed line."""
    lw = w * 0.26
    n = len(periods)
    pw = (w - lw) / n
    rh = min(0.36, (h - 0.35) / len(rows))
    for k, p in enumerate(periods):
        _t(d, s, x + lw + k * pw, y, pw, 0.3, p, bold=True, align=C, anchor=MID_A)
        d._line(s, x + lw + k * pw, y + 0.32, x + lw + k * pw, y + 0.35 + rh * len(rows), PALEGREY, 0.75)
    d._line(s, x, y + 0.32, x + w, y + 0.32, NAVY, 1.0)
    for i, r in enumerate(rows):
        ry = y + 0.35 + i * rh
        grp = len(r) > 3 and r[3] == "group"
        _t(d, s, x + (0 if grp else 0.12), ry, lw, rh, r[0], bold=grp, anchor=MID_A, color=NAVY if grp else BLACK)
        if r[1] is not None:
            d._rect(s, x + lw + r[1] * pw + 0.03, ry + rh * (0.2 if grp else 0.28), (r[2] - r[1] + 1) * pw - 0.06,
                    rh * (0.6 if grp else 0.44), fill=NAVY if grp else STEEL)
        d._line(s, x, ry + rh, x + w, ry + rh, PALEGREY, 0.5)
    for i, at, lab in milestones:
        ry = y + 0.35 + i * rh
        cx = x + lw + (at + 0.5) * pw
        d._shape(s, MSO_SHAPE.DIAMOND, cx - 0.1, ry + rh / 2 - 0.1, 0.2, 0.2, fill=GOLD)
        _t(d, s, cx + 0.12, ry, 2.0, rh, lab, bold=True, anchor=MID_A, wrap=False)
    if today is not None:
        tx = x + lw + today * pw
        d._line(s, tx, y + 0.3, tx, y + 0.35 + rh * len(rows), GOLD, 1.25, dash=DASH)
        _t(d, s, tx - 0.5, y + 0.35 + rh * len(rows), 1.0, 0.22, "Today", bold=True, align=C)


def bk52(d, s, x, y, w, h, phases, current=None, recommendation=None):
    """Phase arrows with columns: phases = [(title, timing, [bullets])] as arrows across the top (current phase
    navy, others steel), bullets under each, optional recommendation takeaway at the bottom."""
    n = len(phases)
    pw = w / n
    ah = 0.62
    rh = 0.55 if recommendation else 0
    for k, (title, timing, bl) in enumerate(phases):
        px = x + k * pw
        a = d._shape(s, MSO_SHAPE.PENTAGON if k == 0 else MSO_SHAPE.CHEVRON, px, y, pw + 0.12, ah,
                     fill=NAVY if current is None or k == current else STEEL)
        d._shape_text(a, [f"**{title}**", timing], size=F_BODY, color=WHITE)
        d.bullets(s, px + 0.12, y + ah + 0.15, pw - 0.25, h - ah - 0.25 - rh, bl)
        if k:
            d._line(s, px, y + ah + 0.15, px, y + h - rh - 0.15, LGREY, 0.75)
    if recommendation:
        d.takeaway(s, x, y + h - 0.5, w, 0.5, recommendation)


def bk55(d, s, x, y, w, h, buyers, heads=("Buyer", "Key financials", "Strategic rationale", "Considerations",
                                         "Likelihood"), logos=None):
    """Buyer profiles: buyers = [(name, [(label, value)], [rationale], [considerations], likelihood 0-4)]."""
    cw = [w * 0.14, w * 0.2, w * 0.3, w * 0.28, w * 0.08]
    xs = [x + sum(cw[:j]) for j in range(5)]
    for j, hd in enumerate(heads):
        _t(d, s, xs[j] + 0.05, y, cw[j] - 0.1, 0.3, hd, bold=True, align=C if j == 4 else LFT, anchor=MID_A)
    d._line(s, x, y + 0.32, x + w, y + 0.32, NAVY, 1.0)
    rh = (h - 0.35) / len(buyers)
    logos = logos or {}
    for i, (nm, fin, rat, con, lk) in enumerate(buyers):
        ry = y + 0.35 + i * rh
        if nm in logos:
            d._fit_picture(s, logos[nm], xs[0] + 0.05, ry + 0.1, cw[0] - 0.1, min(0.6, rh - 0.2))
        else:
            _t(d, s, xs[0] + 0.05, ry, cw[0] - 0.1, rh, nm, bold=True, size=F_HEAD, anchor=MID_A)
        fl = [f"{a}: **{b}**" for a, b in fin]
        _t(d, s, xs[1] + 0.05, ry + 0.05, cw[1] - 0.1, rh - 0.1, fl, anchor=MID_A)
        d.bullets(s, xs[2] + 0.05, ry + 0.05, cw[2] - 0.1, rh - 0.1, rat, gap=2)
        d.bullets(s, xs[3] + 0.05, ry + 0.05, cw[3] - 0.1, rh - 0.1, con, gap=2)
        d._harvey_ball(s, xs[4] + cw[4] / 2 - 0.14, ry + rh / 2 - 0.14, 0.28, lk)
        d._line(s, x, ry + rh, x + w, ry + rh, LGREY, 0.75)
    _t(d, s, x + w - 3.0, y + h + 0.02, 3.0, 0.2, "Likelihood: full ball = most likely", italic=True, color=GREY,
       align=RGT, wrap=False)


def bk57(d, s, x, y, w, h, cases, takeaway=None):
    """Two-case comparison: cases = [(title, subtitle, [(label, text)])] as two cards side by side with aligned
    rows; takeaway = the implication underneath."""
    n = len(cases)
    g = 0.3
    cw = (w - g * (n - 1)) / n
    th = 0.55 if takeaway else 0
    nrow = max(len(c[2]) for c in cases)
    for k, (title, sub, rows) in enumerate(cases):
        cx = x + k * (cw + g)
        yy = _head(d, s, cx, y, cw, title, "bar")
        if sub:
            _t(d, s, cx, yy - 0.04, cw, 0.28, sub, italic=True, color=GREY, align=C)
            yy += 0.26
        rh = (y + h - th - yy) / nrow
        for i, (lab, txt) in enumerate(rows):
            ry = yy + i * rh
            _t(d, s, cx, ry, cw * 0.28, rh, lab, bold=True, color=NAVY, anchor=MID_A)
            _t(d, s, cx + cw * 0.3, ry, cw * 0.7, rh, txt, anchor=MID_A)
            d._line(s, cx, ry + rh, cx + cw, ry + rh, LGREY, 0.75)
    if takeaway:
        d.takeaway(s, x, y + h - 0.48, w, 0.48, takeaway)


def bk58(d, s, x, y, w, h, columns):
    """Headed box columns: columns = [[(head, [items])]] where an item is a bullet string or a (left, right) pair
    (name, role); each box has a navy outline and a navy rule title."""
    n = len(columns)
    g = 0.3
    cw = (w - g * (n - 1)) / n
    for k, boxes in enumerate(columns):
        cx = x + k * (cw + g)
        weights = [len(items) + 1.2 for _, items in boxes]
        tot = sum(weights)
        avail = h - 0.15 * (len(boxes) - 1)
        yy = y
        for head, items in boxes:
            bh = avail * (len(items) + 1.2) / tot
            _box(d, s, cx, yy, cw, bh, line=NAVY)
            _t(d, s, cx + 0.1, yy + 0.05, cw - 0.2, 0.26, head, bold=True, color=NAVY)
            d._line(s, cx + 0.1, yy + 0.33, cx + cw - 0.1, yy + 0.33, NAVY, 0.75)
            if items and isinstance(items[0], (list, tuple)):
                ih = (bh - 0.45) / len(items)
                for i, (a, b) in enumerate(items):
                    _t(d, s, cx + 0.1, yy + 0.4 + i * ih, cw * 0.45, ih, a, bold=True, anchor=MID_A)
                    _t(d, s, cx + cw * 0.47, yy + 0.4 + i * ih, cw * 0.5, ih, b, anchor=MID_A)
            else:
                d.bullets(s, cx + 0.1, yy + 0.4, cw - 0.2, bh - 0.45, items, gap=2)
            yy += bh + 0.15


# ---------------------------------------------------------------- gallery / thumbnails
def gallery(out, ids=None):
    """One slide per design, filled with bank_samples.SAMPLE."""
    from bank_samples import SAMPLE
    from bank_specs import SPECS
    d = Deck()
    for bid in ids or sorted(SPECS):
        s = d.slide(f"{bid}: {SPECS[bid]['name']}", SPECS[bid]["use_when"])
        d.bank(s, bid, data=SAMPLE[bid])
    d.save(out)
    return d


def thumbs(tmp=None, ids=None):
    """Render the gallery and save assets/layout_thumbs/bk-NN.png (480 px wide)."""
    import tempfile
    import pymupdf
    from bank_specs import SPECS
    from office import render
    ids = ids or sorted(SPECS)
    tmp = tmp or tempfile.mkdtemp(prefix="bk_")
    deck = os.path.join(tmp, "bank_gallery.pptx")
    gallery(deck, ids)
    out = os.path.join(HERE, "..", "assets", "layout_thumbs")
    os.makedirs(out, exist_ok=True)
    for bid, f in zip(ids, render(deck, os.path.join(tmp, "png"))):
        pix = pymupdf.Pixmap(f)
        fac = 480 / pix.width
        pymupdf.Pixmap(pix, int(pix.width * fac), int(pix.height * fac), None).save(
            os.path.join(out, f"bk-{bid[2:]}.png"))
    print(f"{len(ids)} thumbnails -> {os.path.abspath(out)} (full size in {tmp})")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "gallery":
        gallery(sys.argv[2], sys.argv[3:] or None)
    elif cmd == "thumbs":
        thumbs(None, sys.argv[2:] or None)
    else:
        sys.exit("usage: bank_layouts.py gallery out.pptx [bk01 ...] | thumbs [bk01 ...]")
